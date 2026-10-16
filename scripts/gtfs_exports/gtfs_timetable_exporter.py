"""Generates Excel-formatted transit timetables from GTFS data.

This script processes GTFS files (e.g., trips.txt, stop_times.txt) to export
Excel timetable sheets by route, direction, and service type. Output mimics
print-style public timetables: one row per trip, one column per timepoint.

Columns cover every timepoint served by any trip in the direction, so branches
and short-turns keep all their times. Trips are matched to columns by the order
of the stops they visit, not by stop_sequence numbers, which GTFS only defines
within a single trip. If trips in one direction visit stops in conflicting
orders, the extra patterns go to additional sheets (e.g. ``Direction_0_2``).
Each cell shows the departure time, except at a trip's final stop, which shows
the arrival time. Rows are ordered by times at shared timepoints, so trips
starting partway along a route appear in the correct position. First displayed
times break ties between trips whose relative order is unconstrained. If no
single row order works at every timepoint, the first displayed time order is
retained and the conflicts are reported for review.

Outputs:
    - One Excel file per route × service type with timetable sheets by direction.
      Files with problems to review end in ``WARNING_FILENAME_SUFFIX``
      (e.g. ``route_101_timetable_weekday_sid_1_warning.xlsx``).
    - ``RUN_LOG_FILENAME`` in ``BASE_OUTPUT_PATH``: the problems found in each
      flagged workbook, any failed exports, and the configuration used.

Typical usage:
    Configuration is manual: edit the constants in the CONFIGURATION section
    near the top before execution, then run within ArcGIS Pro or a Jupyter
    notebook.
"""

from __future__ import annotations

import heapq
import json
import logging
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# =============================================================================
# CONFIGURATION
# =============================================================================
# === BEGIN CONFIG ===

GTFS_FOLDER_PATH = r"Path\To\Your\GTFS_Folder"  # Folder contains GTFS .txt files

BASE_OUTPUT_PATH = r"Path\To\Your\Output_Folder"

# Filter service id's can be found in calendar.txt or calendar_dates.txt
FILTER_SERVICE_IDS: list[str] = ["1", "2", "3", "4"]
FILTER_IN_ROUTES: list[str] = []  # routes to keep
FILTER_OUT_ROUTES: list[str] = []  # routes to exclude

# Explicit schedule label overrides for the selected service_ids
# Assignments can be found in calendar.txt file
SERVICE_LABEL_OVERRIDES: dict[str, str] = {
    "1": "Saturday",
    "2": "Sunday",
    "3": "Holiday",
    "4": "Weekday",
}

TIME_FORMAT_OPTION = "24"  # "12" or "24"
MISSING_TIME = "---"
MAX_COLUMN_WIDTH = 30

# Workbooks with problems to review (e.g. times out of order, skipped trips) get
# this suffix: route_101_timetable_weekday_sid_1_warning.xlsx. The run log lists
# each problem. Set to "" to keep plain file names.
WARNING_FILENAME_SUFFIX = "_warning"

RUN_LOG_FILENAME = "gtfs_timetable_exporter_runlog.txt"  # in BASE_OUTPUT_PATH

# Every output must be traceable: if the run log cannot be written, the script
# exits with an error. Set to False only when writing to a genuinely read-only
# location.
REQUIRE_RUN_LOG: bool = True

# Required and optional GTFS files. At least one of calendar.txt and
# calendar_dates.txt must be present.
REQUIRED_GTFS_FILES: tuple[str, ...] = (
    "agency.txt",
    "stops.txt",
    "routes.txt",
    "trips.txt",
    "stop_times.txt",
)

OPTIONAL_GTFS_FILES: tuple[str, ...] = (
    "calendar.txt",
    "calendar_dates.txt",
    "feed_info.txt",
    "shapes.txt",
    "frequencies.txt",
    "transfers.txt",
    "fare_attributes.txt",
    "fare_rules.txt",
)

LOG_LEVEL: int = logging.INFO  # DEBUG / INFO / WARNING / ERROR

# === END CONFIG ===

# =============================================================================
# FUNCTIONS
# =============================================================================


def format_output_folder_name(service_id: str, schedule_type: str) -> str:
    """Return an output folder name reflecting the schedule label and service_id.

    Example:
        Weekday (service_id=15) → 'weekday_sid_15'
    """
    return f"{_slugify(schedule_type)}_sid_{_safe_filename_part(service_id)}"


def format_workbook_name(
    route_short_name: str, schedule_type: str, service_id: str, flagged: bool = False
) -> str:
    """Return the workbook file name for one route and service_id.

    A *flagged* workbook (one with problems to review) ends in
    ``WARNING_FILENAME_SUFFIX``.

    Example:
        Route 101, Weekday (service_id=1) → 'route_101_timetable_weekday_sid_1.xlsx'
    """
    suffix = WARNING_FILENAME_SUFFIX if flagged else ""
    return (
        f"route_{_safe_filename_part(route_short_name)}_timetable_"
        f"{_slugify(schedule_type)}_sid_{_safe_filename_part(service_id)}{suffix}.xlsx"
    )


def _in_ipython() -> bool:
    """Return True if running under IPython/Jupyter."""
    return "IPYKERNEL" in os.environ or hasattr(sys, "ps1")


def _list_feed_files(gtfs_path: str) -> set[str]:
    """Return the file names in a GTFS folder or ``.zip`` (empty set if unreadable)."""
    try:
        if os.path.isdir(gtfs_path):
            return set(os.listdir(gtfs_path))
        if zipfile.is_zipfile(gtfs_path):
            with zipfile.ZipFile(gtfs_path) as archive:
                return {os.path.basename(name) for name in archive.namelist()}
    except OSError:
        pass
    return set()


def _safe_filename_part(text: str) -> str:
    """Replace characters that Windows or POSIX reject in file names with ``_``."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(text)).strip(" .")
    return cleaned or "_"


def _slugify(label: str) -> str:
    """Return a filesystem-friendly slug from a schedule label."""
    slug = str(label).strip().lower().replace(" ", "_").replace("-", "_").replace("/", "_")
    return _safe_filename_part(slug)


def _minutes_to_hhmm(minutes: int) -> str:
    """Format minutes after midnight as *HH:MM* (hours may exceed 23)."""
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def time_to_minutes(time_str: str) -> Optional[int]:
    """Convert a time string to minutes after midnight.

    Supports *HH:MM*, GTFS-style *HH:MM:SS* (seconds are ignored), and
    *HH:MM AM/PM* tokens.  Times such as ``'25:15'`` (for past-midnight
    service) are accepted.  ``MISSING_TIME`` or invalid inputs return ``None``.

    Args:
        time_str: Raw time string.

    Returns:
        Minutes after midnight, or ``None`` if the string is invalid.
    """
    if not isinstance(time_str, str):
        return None
    if time_str.strip() == MISSING_TIME:
        return None

    result = None
    try:
        match = re.match(
            r"^(\d{1,2}):(\d{2})(?::\d{2})?(?:\s*(AM|PM))?$", time_str.strip(), re.IGNORECASE
        )
        if match:
            hour_str, minute_str, period = match.groups()
            hour = int(hour_str)
            minute = int(minute_str)
            if period:
                period = period.upper()
                if period == "PM" and hour != 12:
                    hour += 12
                elif period == "AM" and hour == 12:
                    hour = 0
            result = hour * 60 + minute
    except (ValueError, TypeError, re.error):
        result = None

    return result


def adjust_time(time_str: str, time_format: str = "24") -> Optional[str]:
    """Re-format a GTFS time string.

    Args:
        time_str: Raw *arrival_time* or *departure_time* field.
        time_format: Either ``"12"`` for 12-hour *h:MM AM/PM* output or
            ``"24"`` for zero-padded 24-hour output. In 12-hour output,
            past-midnight GTFS times wrap to the clock time (``25:15`` →
            ``1:15 AM``); 24-hour output keeps them as-is (``25:15``).

    Returns:
        The re-formatted time, ``MISSING_TIME`` if no time should be
        shown, or ``None`` for an unparsable input.
    """
    if not isinstance(time_str, str):
        return None
    if time_str.strip() == MISSING_TIME:
        return MISSING_TIME

    parts = time_str.strip().split(":")
    if len(parts) < 2:
        return None

    try:
        hours = int(parts[0])
        minutes = int(parts[1])
        if time_format == "12":
            # Past-midnight service (e.g. 25:15) shows the next day's clock time.
            clock_hours = hours % 24
            period = "AM" if clock_hours < 12 else "PM"
            adjusted_hour = clock_hours % 12
            if adjusted_hour == 0:
                adjusted_hour = 12
            return f"{adjusted_hour}:{minutes:02d} {period}"
        # 24-hour format:
        return f"{hours:02d}:{minutes:02d}"
    except ValueError:
        return None


def prepare_timepoints(stop_times: pd.DataFrame) -> pd.DataFrame:
    """Return only the rows in *stop_times* where ``timepoint == 1``.

    If the ``timepoint`` column is missing or entirely blank, every row is
    returned: GTFS treats all times as exact when no timepoint values are
    provided. Each row also gets an ``is_final_stop`` flag marking the last
    stop of its trip (computed before filtering).

    Args:
        stop_times: Raw *stop_times.txt* DataFrame.

    Returns:
        A DataFrame containing just the timepoint rows.
    """
    stop_times = stop_times.copy()
    stop_times["stop_sequence"] = pd.to_numeric(stop_times["stop_sequence"], errors="coerce")
    if stop_times["stop_sequence"].isna().any():
        # Using logging instead of print for consistency, though this could be print too.
        logging.warning("Some 'stop_sequence' values could not be converted to numeric.")

    blank_stop_ids = int(stop_times["stop_id"].isna().sum())
    if blank_stop_ids:
        logging.warning(
            "%d stop_times rows have a blank stop_id and will be left out of the timetables "
            "(pandas also reads IDs such as 'NA' or 'N/A' as blank).",
            blank_stop_ids,
        )

    last_sequence = stop_times.groupby("trip_id")["stop_sequence"].transform("max")
    stop_times["is_final_stop"] = stop_times["stop_sequence"] == last_sequence

    if "timepoint" not in stop_times.columns:
        logging.warning("Warning: 'timepoint' column not found. Using all stops as timepoints.")
        return stop_times

    flags = stop_times["timepoint"].fillna("").astype(str).str.strip()
    if (flags == "").all():
        logging.info(
            "'timepoint' column is blank; GTFS treats all times as exact, so using all stops "
            "as timepoints."
        )
        return stop_times

    blank_count = int((flags == "").sum())
    if blank_count:
        logging.warning(
            "%d stop_times rows have a blank 'timepoint' value and are excluded.", blank_count
        )
    timepoints = stop_times[flags == "1"].copy()
    logging.info("Filtered STOP_TIMES to rows with timepoint=1.")
    return timepoints


def remove_empty_schedule_columns(input_df: pd.DataFrame) -> pd.DataFrame:
    """Drop schedule columns whose entire contents equal ``MISSING_TIME``.

    Args:
        input_df: Schedule grid whose columns end with ``" Schedule"``.

    Returns:
        The same DataFrame, with all-blank schedule columns removed.
    """
    schedule_cols = [col for col in input_df.columns if col.endswith("Schedule")]
    all_blank_cols = [col for col in schedule_cols if (input_df[col] == MISSING_TIME).all()]
    input_df = input_df.drop(columns=all_blank_cols)
    return input_df


def flag_issue(issues: Optional[list[str]], message: str, *args: Any) -> None:
    """Log a warning about the workbook being built and record it in *issues*.

    A workbook with recorded issues is flagged for review: its file name ends
    in ``WARNING_FILENAME_SUFFIX`` and the run log lists the issues. Recording
    does not depend on ``LOG_LEVEL``.
    """
    logging.warning(message, *args)
    if issues is not None:
        issues.append((message % args if args else message).strip())


def check_schedule_order(
    minutes_df: pd.DataFrame,
    route_short_name: str,
    schedule_type: str,
    dir_id: str,
    issues: Optional[list[str]] = None,
) -> int:
    """Validate that times increase by stop (row) and by trip (column).

    A time that decreases down a column means a trip serves a stop earlier
    than the preceding trip with a time in that column. This can reflect row
    ordering or legitimate overtaking, so it is reported for review rather
    than treated as proof of a GTFS error.

    Args:
        minutes_df (pandas.DataFrame): Minutes after midnight, one row per trip
            (indexed by ``trip_id``) and one column per timepoint, in display
            order. Missing times are ``NaN``.
        route_short_name (str): Human-readable route label.
        schedule_type (str): *Weekday*, *Saturday*, *Holiday*, etc.
        dir_id (str): ``direction_id`` value.
        issues (list[str] | None): Receives one entry per violation; see
            :func:`flag_issue`.

    Returns:
        The number of order violations found (0 means the check passed).
    """
    violations = 0

    # Row-wise check
    for trip_id, row in minutes_df.iterrows():
        last_time = None
        for col_name, current_time in row.items():
            if pd.isna(current_time):
                continue
            if last_time is not None and current_time < last_time:
                violations += 1
                flag_issue(
                    issues,
                    "⚠️ Time order violation in Route '%s', Schedule '%s', Direction '%s', "
                    "Trip '%s': '%s' time %s is earlier than the previous stop's time.",
                    route_short_name,
                    schedule_type,
                    dir_id,
                    trip_id,
                    col_name,
                    _minutes_to_hhmm(int(current_time)),
                )
                break
            last_time = current_time

    # Column-wise check
    for col_name in minutes_df.columns:
        last_time = None
        last_trip_id = None
        for trip_id, current_time in minutes_df[col_name].items():
            if pd.isna(current_time):
                continue
            if last_time is not None and current_time < last_time:
                violations += 1
                flag_issue(
                    issues,
                    "⚠️ Time order violation in Route '%s', Schedule '%s', Direction '%s', "
                    "Stop '%s': trip '%s' at %s is earlier than preceding nonblank trip "
                    "'%s' at %s (row ordering, overtaking, or a data error).",
                    route_short_name,
                    schedule_type,
                    dir_id,
                    col_name,
                    trip_id,
                    _minutes_to_hhmm(int(current_time)),
                    last_trip_id,
                    _minutes_to_hhmm(int(last_time)),
                )
            last_time = current_time
            last_trip_id = trip_id

    if violations:
        logging.warning(
            "Schedule order check found %d violation(s) for Route '%s', Schedule '%s', "
            "Direction '%s'.",
            violations,
            route_short_name,
            schedule_type,
            dir_id,
        )
    else:
        logging.info(
            "✅ Schedule order check passed for Route '%s', Schedule '%s', Direction '%s'.",
            route_short_name,
            schedule_type,
            dir_id,
        )
    return violations


def safe_check_schedule_order(
    minutes_df: pd.DataFrame,
    route_short_name: str,
    schedule_type: str,
    dir_id: str,
    issues: Optional[list[str]] = None,
) -> Optional[int]:
    """Run :func:`check_schedule_order` but never propagate its errors.

    A check that could not run is recorded in *issues*, since the workbook's
    order is then unverified.
    """
    try:
        return check_schedule_order(minutes_df, route_short_name, schedule_type, dir_id, issues)
    except (ValueError, TypeError, KeyError) as error:
        logging.error(  # Changed print to logging.error
            f"❌ Skipping schedule order check for route '{route_short_name}', "
            f"schedule '{schedule_type}', direction '{dir_id}' due to error:\n  {error}"
        )
        if issues is not None:
            issues.append(
                f"Schedule order check skipped for direction '{dir_id}' due to error: {error}"
            )
        return None


def map_service_id_to_schedule(service_row_local: pd.Series) -> str:
    """Translate one **calendar.txt** row into a schedule label.

    Args:
        service_row_local: A single record from *calendar.txt*.

    Returns:
        A label such as ``"Weekday"``, ``"Saturday"``, ``"Daily"``, etc.
    """
    days = [
        "monday",
        "tuesday",
        "wednesday",
        "thursday",
        "friday",
        "saturday",
        "sunday",
    ]
    served_days = [day for day in days if service_row_local.get(day, "0") == "1"]

    weekday = {"monday", "tuesday", "wednesday", "thursday", "friday"}
    weekday_except_friday = {"monday", "tuesday", "wednesday", "thursday"}
    saturday = {"saturday"}
    sunday = {"sunday"}
    weekend = {"saturday", "sunday"}
    daily = set(days)

    schedule_label = "Holiday"
    if served_days:
        served_set = set(served_days)
        if served_set == weekday:
            schedule_label = "Weekday"
        elif served_set == weekday_except_friday:
            schedule_label = "Weekday_except_Friday"
        elif served_set == saturday:
            schedule_label = "Saturday"
        elif served_set == sunday:
            schedule_label = "Sunday"
        elif served_set == weekend:
            schedule_label = "Weekend"
        elif served_set == {"friday", "saturday"}:
            schedule_label = "Friday-Saturday"
        elif served_set == daily:
            schedule_label = "Daily"
        else:
            schedule_label = "Special"  # Could be a custom name based on days
    return schedule_label


def build_service_calendar(
    calendar_df: Optional[pd.DataFrame],
    calendar_dates_df: Optional[pd.DataFrame],
) -> pd.DataFrame:
    """Return one row per service_id with weekday flags, from both calendar files.

    Services defined in *calendar.txt* keep their own flags. Services that
    appear only in *calendar_dates.txt* (as added dates, ``exception_type`` 1)
    get a flag for each weekday their added dates fall on, so they can be
    labelled by :func:`map_service_id_to_schedule` like any other service.

    Args:
        calendar_df: Parsed *calendar.txt*, or ``None`` if the feed has none.
        calendar_dates_df: Parsed *calendar_dates.txt*, or ``None``.

    Returns:
        DataFrame with ``service_id`` and the seven weekday columns (``"1"``/``"0"``).
    """
    days = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
    frames = []
    known_ids: set[str] = set()
    if calendar_df is not None and not calendar_df.empty:
        frames.append(calendar_df)
        known_ids = set(calendar_df["service_id"].dropna())

    if calendar_dates_df is not None and not calendar_dates_df.empty:
        added = calendar_dates_df[
            calendar_dates_df["exception_type"].fillna("").astype(str).str.strip() == "1"
        ]
        added = added[~added["service_id"].isin(known_ids)]
        weekday_names = pd.to_datetime(
            added["date"].astype(str).str.strip(), format="%Y%m%d", errors="coerce"
        ).dt.day_name()
        added = added.assign(weekday=weekday_names.str.lower()).dropna(subset=["weekday"])
        rows = []
        for service_id, group in added.groupby("service_id", sort=True):
            served = set(group["weekday"])
            rows.append(
                {"service_id": service_id, **{d: "1" if d in served else "0" for d in days}}
            )
        if rows:
            frames.append(pd.DataFrame(rows))
            logging.info("Found %d service_id(s) defined only in calendar_dates.txt.", len(rows))

    if not frames:
        return pd.DataFrame(columns=["service_id", *days])
    return pd.concat(frames, ignore_index=True)


def build_service_id_schedule_map(calendar_df: pd.DataFrame) -> dict[str, str]:
    """Create service_id → schedule_type lookup with explicit overrides.

    Args:
        calendar_df: The parsed calendar.txt dataframe filtered to relevant services.

    Returns:
        Mapping from service_id to a human-friendly schedule label (e.g., "Weekday").
        Explicit overrides from SERVICE_LABEL_OVERRIDES take precedence; otherwise
        labels are inferred from the weekday/weekend bit flags.
    """
    service_id_schedule_map: dict[str, str] = {}
    for _, service_row_local in calendar_df.iterrows():
        sid_val = service_row_local["service_id"]
        stype_var = SERVICE_LABEL_OVERRIDES.get(
            sid_val, map_service_id_to_schedule(service_row_local)
        )
        service_id_schedule_map[sid_val] = stype_var
    return service_id_schedule_map


def get_all_route_short_names(routes_df: pd.DataFrame) -> list[str]:
    """Return every unique ``route_short_name`` in ascending order."""
    return sorted(routes_df["route_short_name"].dropna().unique().tolist())


def apply_in_out_filters(route_list: list[str]) -> list[str]:
    """Apply ``FILTER_IN_ROUTES`` and ``FILTER_OUT_ROUTES`` to *route_list*."""
    route_set = set(route_list)

    if FILTER_IN_ROUTES:
        route_set = route_set.intersection(set(FILTER_IN_ROUTES))

    if FILTER_OUT_ROUTES:
        route_set = route_set.difference(set(FILTER_OUT_ROUTES))

    return sorted(list(route_set))


def _clean_time(value: Any) -> str:
    """Return a stripped GTFS time string, or ``""`` for a blank or missing value."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    return str(value).strip()


def select_trip_times(trip_stop_times: pd.DataFrame, time_fmt: str) -> list[tuple[str, int, str]]:
    """Return a trip's displayable timepoints in trip order.

    Each cell shows the departure time, except at the trip's final stop
    (``is_final_stop``), which shows the arrival time. If the preferred field
    is blank the other one is used; rows with neither time are skipped.

    Args:
        trip_stop_times: Timepoint rows for one trip.
        time_fmt: ``"12"`` or ``"24"``; see :func:`adjust_time`.

    Returns:
        ``(stop_id, minutes_after_midnight, display_time)`` tuples ordered by
        ``stop_sequence``.
    """
    selected = []
    for _, row in trip_stop_times.sort_values("stop_sequence").iterrows():
        stop_id = row["stop_id"]
        if pd.isna(stop_id):
            continue
        arr_val = _clean_time(row.get("arrival_time"))
        dep_val = _clean_time(row.get("departure_time"))
        if bool(row.get("is_final_stop", False)):
            time_val = arr_val or dep_val
        else:
            time_val = dep_val or arr_val

        minutes = time_to_minutes(time_val)
        display = adjust_time(time_val, time_fmt)
        if minutes is None or not display:
            continue
        selected.append((str(stop_id), minutes, display))
    return selected


def merge_stop_patterns(base: Sequence[str], pattern: Sequence[str]) -> list[str]:
    """Return the shortest stop order that contains both *base* and *pattern* in order.

    This is the shortest common supersequence of the two stop lists. Where
    the two diverge (e.g. branches A-B-D and A-C-D), *base*'s stops come first.

    Args:
        base: Current column order, as ``stop_id`` values.
        pattern: One trip pattern's ``stop_id`` values in visit order.

    Returns:
        The merged column order.
    """
    n, m = len(base), len(pattern)
    # common[i][j] = length of the longest common subsequence of base[i:] and pattern[j:]
    common = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            if base[i] == pattern[j]:
                common[i][j] = common[i + 1][j + 1] + 1
            else:
                common[i][j] = max(common[i + 1][j], common[i][j + 1])

    merged: list[str] = []
    i = j = 0
    while i < n and j < m:
        if base[i] == pattern[j]:
            merged.append(base[i])
            i += 1
            j += 1
        elif common[i + 1][j] >= common[i][j + 1]:
            merged.append(base[i])
            i += 1
        else:
            merged.append(pattern[j])
            j += 1
    merged.extend(base[i:])
    merged.extend(pattern[j:])
    return merged


def build_column_layouts(
    pattern_counts: Mapping[tuple[str, ...], int],
) -> tuple[list[list[str]], dict[tuple[str, ...], int]]:
    """Group a direction's stop patterns into as few column layouts as possible.

    Patterns are merged, most common first, into a single column order. A
    pattern that can only be merged by repeating a stop more often than any
    trip visits it (e.g. A-B-C and C-B-A) visits stops in a conflicting order;
    it starts or joins another layout, which is exported as a separate sheet.

    Args:
        pattern_counts: Number of trips for each pattern (tuple of ``stop_id``).

    Returns:
        The column layouts (each a list of ``stop_id``), and the index of the
        layout each pattern was assigned to.
    """
    ordered = sorted(pattern_counts, key=lambda p: (-pattern_counts[p], -len(p), p))
    layouts: list[list[str]] = []
    assignment: dict[tuple[str, ...], int] = {}
    for pattern in ordered:
        for idx, columns in enumerate(layouts):
            merged = merge_stop_patterns(columns, pattern)
            # Counter union keeps, per stop, the larger visit count of the two.
            if len(merged) == sum((Counter(columns) | Counter(pattern)).values()):
                layouts[idx] = merged
                assignment[pattern] = idx
                break
        else:
            assignment[pattern] = len(layouts)
            layouts.append(list(pattern))
    return layouts, assignment


def embed_pattern(pattern: Sequence[str], columns: Sequence[str]) -> list[int]:
    """Return the column index for each stop in *pattern*, preserving visit order.

    Only stops that a route visits more than once can match more than one
    column. In that case the most compact placement is used, so a short trip
    around the end of a loop lands on the loop's later columns.

    Args:
        pattern: One trip's ``stop_id`` values in visit order.
        columns: A column layout that contains *pattern* in order.

    Returns:
        One column index per entry in *pattern*.

    Raises:
        ValueError: *pattern* does not fit in *columns* in order.
    """
    if not pattern:
        return []
    best: Optional[list[int]] = None
    for start, stop_id in enumerate(columns):
        if stop_id != pattern[0]:
            continue
        positions = [start]
        for next_stop in pattern[1:]:
            try:
                positions.append(columns.index(next_stop, positions[-1] + 1))
            except ValueError:
                break
        if len(positions) < len(pattern):
            continue
        if best is None or positions[-1] - positions[0] < best[-1] - best[0]:
            best = positions
    if best is None:
        raise ValueError(f"Stop pattern {list(pattern)} does not fit the column layout.")
    return best


def build_column_headers(columns: Sequence[str], stops_df: pd.DataFrame) -> list[str]:
    """Return a unique ``"<stop label> Schedule"`` header for each column.

    The label is the stop name. Distinct stops that share a name are told
    apart by ``stop_code`` (or ``stop_id`` when codes are missing or not
    unique), e.g. ``"Central Station [1234]"``. Repeat visits to the same
    stop are numbered, e.g. ``"Central Station (2)"``.

    Args:
        columns: Column layout as ``stop_id`` values; may repeat a stop.
        stops_df: Parsed *stops.txt* file.

    Returns:
        Header strings, one per column, all distinct.
    """
    lookup = stops_df.drop_duplicates("stop_id").set_index("stop_id")
    names: dict[str, str] = {}
    codes: dict[str, str] = {}
    for stop_id in dict.fromkeys(columns):
        name = lookup.loc[stop_id, "stop_name"] if stop_id in lookup.index else None
        code = (
            lookup.loc[stop_id, "stop_code"]
            if "stop_code" in lookup.columns and stop_id in lookup.index
            else None
        )
        names[stop_id] = (
            name.strip() if isinstance(name, str) and name.strip() else f"Unknown stop {stop_id}"
        )
        codes[stop_id] = code.strip() if isinstance(code, str) else ""

    ids_by_name: dict[str, list[str]] = defaultdict(list)
    for stop_id, name in names.items():
        ids_by_name[name].append(stop_id)

    stop_labels: dict[str, str] = {}
    for name, stop_ids in ids_by_name.items():
        if len(stop_ids) == 1:
            stop_labels[stop_ids[0]] = name
            continue
        shared_codes = [codes[s] for s in stop_ids]
        use_codes = all(shared_codes) and len(set(shared_codes)) == len(shared_codes)
        for stop_id in stop_ids:
            stop_labels[stop_id] = f"{name} [{codes[stop_id] if use_codes else stop_id}]"

    visit_counts: dict[str, int] = defaultdict(int)
    label_counts: dict[str, int] = defaultdict(int)
    headers = []
    for stop_id in columns:
        visit_counts[stop_id] += 1
        visit = visit_counts[stop_id]
        label = stop_labels[stop_id] if visit == 1 else f"{stop_labels[stop_id]} ({visit})"
        # Guard against a stop whose own name collides with a numbered label.
        label_counts[label] += 1
        if label_counts[label] > 1:
            label = f"{label} #{label_counts[label]}"
        headers.append(f"{label} Schedule")
    return headers


def process_single_trip(
    trip_id: str,
    trip_times: Sequence[tuple[str, int, str]],
    columns: Sequence[str],
) -> tuple[list[str], list[Optional[int]]]:
    """Place one trip's times into the columns of its layout.

    Args:
        trip_id: Trip being placed (used in error messages).
        trip_times: Output of :func:`select_trip_times` for the trip.
        columns: The column layout the trip's stop pattern belongs to.

    Returns:
        The display times (``MISSING_TIME`` where the trip has no time) and
        the matching minutes after midnight (``None`` where missing).
    """
    display = [MISSING_TIME] * len(columns)
    minutes: list[Optional[int]] = [None] * len(columns)
    try:
        positions = embed_pattern([stop_id for stop_id, _, _ in trip_times], columns)
    except ValueError as error:
        raise ValueError(f"Trip {trip_id}: {error}") from error
    for col_idx, (_, trip_minutes, trip_display) in zip(positions, trip_times):
        display[col_idx] = trip_display
        minutes[col_idx] = trip_minutes
    return display, minutes


def get_timetable_row_order(minutes_df: pd.DataFrame) -> tuple[list[int], bool]:
    """Find a row order that makes every timepoint column nondecreasing.

    Earlier times at a shared column impose precedence between trips. Equal
    times impose no order. Connecting adjacent distinct time groups suffices
    because their ordering also orders every earlier group before later ones.
    A stable topological sort combines those constraints across all columns,
    including layouts with no single timepoint served by every trip.

    Args:
        minutes_df: Numeric timetable in first displayed time order. Missing
            values are NaN; times retain GTFS hours beyond midnight. Repeated
            visits to a stop must occupy separate columns.

    Returns:
        Positional row indices and whether all column constraints can be met.
        If constraints form a cycle, return the original order and False;
        overtaking or conflicting patterns require review, not hidden times.
    """
    row_count = len(minutes_df)
    successors: list[set[int]] = [set() for _ in range(row_count)]
    indegrees = [0] * row_count

    for column_idx in range(minutes_df.shape[1]):
        time_groups: dict[float, list[int]] = defaultdict(list)
        for row_idx, value in enumerate(minutes_df.iloc[:, column_idx]):
            if pd.notna(value):
                time_groups[float(value)].append(row_idx)

        previous_group: list[int] = []
        for time_value in sorted(time_groups):
            current_group = time_groups[time_value]
            for earlier in previous_group:
                for later in current_group:
                    if later not in successors[earlier]:
                        successors[earlier].add(later)
                        indegrees[later] += 1
            previous_group = current_group

    # Original row positions break ties deterministically, without comparing
    # departure times from different origins when a shared stop supplies order.
    ready = [idx for idx, degree in enumerate(indegrees) if degree == 0]
    heapq.heapify(ready)
    ordered = []
    while ready:
        earlier = heapq.heappop(ready)
        ordered.append(earlier)
        for later in successors[earlier]:
            indegrees[later] -= 1
            if indegrees[later] == 0:
                heapq.heappush(ready, later)

    if len(ordered) != row_count:
        return list(range(row_count)), False
    return ordered, True


def process_trips_for_direction(params: dict[str, Any]) -> list[pd.DataFrame]:
    """Generate the timetable grid(s) for a single ``direction_id``.

    Problems that need review (skipped trips, conflicting stop orders, times
    out of order) are recorded in ``params["issues"]``; see :func:`flag_issue`.

    Returns:
        One DataFrame per column layout (see :func:`build_column_layouts`);
        usually one. Empty list if the direction has no usable trips.
    """
    trips_dir = params["trips_dir"]
    dir_id = params["dir_id"]
    timepoints = params["timepoints"]
    ctx = params["ctx"]
    route_short = params["route_short"]
    sched_type = params["sched_type"]
    issues: Optional[list[str]] = params.get("issues")

    timepoints_dir = timepoints[timepoints["trip_id"].isin(trips_dir["trip_id"])]
    if timepoints_dir.empty:
        logging.info(
            "No stop times in TIMEPOINTS for route '%s', schedule '%s', direction '%s'. Skipping.",
            route_short,
            sched_type,
            dir_id,
        )
        return []

    if "trip_headsign" in trips_dir.columns:
        headsigns = trips_dir.set_index("trip_id")["trip_headsign"].fillna("").to_dict()
    else:
        headsigns = {}

    trip_times: dict[str, list[tuple[str, int, str]]] = {}
    for trip_id, grp in timepoints_dir.groupby("trip_id", sort=True):
        times = select_trip_times(grp, ctx["time_fmt"])
        if times:
            trip_times[str(trip_id)] = times
        else:
            flag_issue(issues, "Trip %s has no usable timepoint times. Skipping.", trip_id)
    if not trip_times:
        return []

    patterns = {trip_id: tuple(s for s, _, _ in times) for trip_id, times in trip_times.items()}
    layouts, assignment = build_column_layouts(Counter(patterns.values()))
    if len(layouts) > 1:
        flag_issue(
            issues,
            "Route '%s', schedule '%s', direction '%s': trips visit stops in conflicting "
            "orders, so the timetable is split across %d sheets.",
            route_short,
            sched_type,
            dir_id,
            len(layouts),
        )

    frames = []
    for layout_idx, columns in enumerate(layouts):
        headers = build_column_headers(columns, ctx["stops"])
        rows = []
        for trip_id, times in trip_times.items():
            if assignment[patterns[trip_id]] != layout_idx:
                continue
            display, minutes = process_single_trip(trip_id, times, columns)
            # Establish the tie-break/fallback order before shared-stop sorting.
            sort_key = (times[0][1], times[-1][1], trip_id)
            rows.append((sort_key, trip_id, display, minutes))
        rows.sort(key=lambda r: r[0])

        out_df = pd.DataFrame(
            [
                [route_short, dir_id, headsigns.get(trip_id, "")] + display
                for _, trip_id, display, _ in rows
            ],
            columns=["Route Name", "Direction ID", "Trip Headsign"] + headers,
        )
        minutes_df = pd.DataFrame(
            [minutes for _, _, _, minutes in rows],
            index=[trip_id for _, trip_id, _, _ in rows],
            columns=headers,
            dtype="float",
        )
        row_order, consistent_order = get_timetable_row_order(minutes_df)
        if not consistent_order:
            flag_issue(
                issues,
                "Route '%s', schedule '%s', direction '%s', layout %d: no single trip row "
                "order makes every timepoint chronological. Possible overtaking or "
                "conflicting patterns; retaining first displayed time order for review.",
                route_short,
                sched_type,
                dir_id,
                layout_idx + 1,
            )
        out_df = out_df.iloc[row_order].reset_index(drop=True)
        minutes_df = minutes_df.iloc[row_order]
        safe_check_schedule_order(minutes_df, route_short, sched_type, dir_id, issues)
        frames.append(remove_empty_schedule_columns(out_df))
    return frames


def export_to_excel_multiple_sheets(
    df_dict: dict[str, pd.DataFrame],
    out_file: str,
) -> bool:
    """Write each DataFrame in *df_dict* to its own sheet in *out_file*.

    Returns:
        ``False`` if the workbook could not be written (e.g. it is open in
        Excel), otherwise ``True`` (including when there is nothing to write).
    """
    if not df_dict:
        logging.info(f"No data to export to {out_file}.")  # Changed to logging
        return True

    try:
        with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
            for sheet_name, input_df in df_dict.items():
                if input_df.empty:
                    logging.info(
                        f"No data for sheet '{sheet_name}'. Skipping."
                    )  # Changed to logging
                    continue

                input_df.to_excel(writer, index=False, sheet_name=sheet_name)
                worksheet = writer.sheets[sheet_name]

                # Adjust column widths & alignment
                for col_num, _ in enumerate(input_df.columns, 1):
                    col_letter = get_column_letter(col_num)
                    header_cell = worksheet[f"{col_letter}1"]
                    header_cell.alignment = Alignment(
                        horizontal="left", vertical="top", wrap_text=True
                    )
                    for row_num in range(2, worksheet.max_row + 1):
                        cell = worksheet[f"{col_letter}{row_num}"]
                        cell.alignment = Alignment(horizontal="left")

                    # Calculate column width
                    column_cells = worksheet[col_letter]
                    try:
                        max_length = max(
                            len(str(cell.value)) for cell in column_cells if cell.value
                        )
                    except (
                        ValueError,
                        TypeError,
                    ):  # Handles case where column might be all None or empty
                        max_length = 10  # Default width
                    adjusted_width = min(max_length + 2, MAX_COLUMN_WIDTH)
                    worksheet.column_dimensions[col_letter].width = adjusted_width
        logging.info(f"Data exported to {out_file}")  # Changed to logging
    except Exception as e:
        logging.error(f"Failed to export to Excel file {out_file}. Error: {e}")
        return False
    return True


# -----------------------------------------------------------------------------
# SUB-STEPS
# -----------------------------------------------------------------------------


def filter_calendar_df(
    calendar_df: pd.DataFrame,
) -> Optional[pd.DataFrame]:
    """Filter *calendar_df* by ``FILTER_SERVICE_IDS``; return *None* if empty."""
    if FILTER_SERVICE_IDS:
        missing = sorted(set(FILTER_SERVICE_IDS) - set(calendar_df["service_id"]))
        if missing:
            logging.warning(
                "FILTER_SERVICE_IDS not found in calendar.txt or calendar_dates.txt: %s",
                ", ".join(missing),
            )
        calendar_df = calendar_df[calendar_df["service_id"].isin(FILTER_SERVICE_IDS)]
    if calendar_df.empty:
        logging.warning(
            "No service_ids found after applying FILTER_SERVICE_IDS. Exiting."
        )  # Changed to logging
        return None
    return calendar_df


def export_route_service(
    route_short_name: str,
    route_ids: pd.Series,
    service_id: str,
    ctx: dict[str, Any],
    flagged: Optional[dict[str, list[str]]] = None,
) -> bool:
    """Build and export the timetable workbook for one route and service_id.

    Trips without a ``direction_id`` (the column is optional in GTFS) are
    exported together on a ``No_Direction`` sheet rather than dropped.

    A workbook with problems to review is saved with ``WARNING_FILENAME_SUFFIX``
    and recorded in *flagged*. The previous run's copy under the other name
    (flagged or not) is removed, so the folder never holds both.

    Args:
        route_short_name: Route to export.
        route_ids: ``route_id`` values sharing that short name.
        service_id: Service to export.
        ctx: Run context; see :func:`process_route_service_combinations`.
        flagged: Receives ``{workbook path: issues}`` for a flagged workbook.

    Returns:
        ``False`` if the workbook could not be written, or an outdated copy
        could not be removed, otherwise ``True``.
    """
    trips_df = ctx["trips"]
    schedule_type = ctx["service_id_schedule_map"].get(service_id, "Unknown")
    issues: list[str] = []

    # Filter trips for this route + this service_id
    relevant_trips = trips_df[
        (trips_df["route_id"].isin(route_ids)) & (trips_df["service_id"] == service_id)
    ]
    if relevant_trips.empty:
        logging.info(
            "  No trips for route='%s' and service_id='%s' (%s).",
            route_short_name,
            service_id,
            schedule_type,
        )
        return True

    if "direction_id" in relevant_trips.columns:
        directions = relevant_trips["direction_id"].fillna("").astype(str).str.strip()
    else:
        directions = pd.Series("", index=relevant_trips.index)
    relevant_trips = relevant_trips.assign(direction_id=directions)
    if (directions == "").any():
        flag_issue(
            issues,
            "  %d trip(s) on route '%s' have no direction_id; exporting them on a "
            "'No_Direction' sheet.",
            int((directions == "").sum()),
            route_short_name,
        )

    df_sheets: dict[str, pd.DataFrame] = {}
    for dir_id in sorted(directions.unique()):
        logging.info(
            "    Building direction_id '%s' for service_id='%s' (%s)...",
            dir_id,
            service_id,
            schedule_type,
        )
        params_dict = {
            "trips_dir": relevant_trips[relevant_trips["direction_id"] == dir_id],
            "dir_id": dir_id,
            "timepoints": ctx["timepoints"],
            "ctx": ctx,
            "route_short": route_short_name,
            "sched_type": schedule_type,
            "issues": issues,
        }
        output_dfs = process_trips_for_direction(params_dict)
        if not output_dfs:
            logging.info("      No output data for direction_id '%s'.", dir_id)
        base_name = f"Direction_{dir_id}" if dir_id else "No_Direction"
        for layout_number, output_df in enumerate(output_dfs, 1):
            sheet_name = base_name if layout_number == 1 else f"{base_name}_{layout_number}"
            df_sheets[sheet_name[:31]] = output_df  # Excel caps sheet names at 31 chars

    if not df_sheets:
        logging.info(
            "  No data to export for service_id '%s' (%s) on route '%s'.",
            service_id,
            schedule_type,
            route_short_name,
        )
        return True

    # Folder reflects the override label (and the service_id)
    service_output_path = os.path.join(
        BASE_OUTPUT_PATH, format_output_folder_name(service_id, schedule_type)
    )
    os.makedirs(service_output_path, exist_ok=True)
    out_file = os.path.join(
        service_output_path,
        format_workbook_name(route_short_name, schedule_type, service_id, flagged=bool(issues)),
    )
    # Where a previous run saved this workbook if it was flagged differently then.
    outdated_file = os.path.join(
        service_output_path,
        format_workbook_name(route_short_name, schedule_type, service_id, flagged=not issues),
    )
    if not export_to_excel_multiple_sheets(df_sheets, out_file):
        return False
    if issues and flagged is not None:
        flagged[out_file] = issues
    if outdated_file != out_file and os.path.isfile(outdated_file):
        try:
            os.remove(outdated_file)
        except OSError as error:
            logging.error(
                "Could not remove the previous run's workbook %s, which now sits beside %s: %s",
                outdated_file,
                out_file,
                error,
            )
            return False
        logging.info("Removed the previous run's workbook %s.", outdated_file)
    return True


def process_route_service_combinations(
    ctx: dict[str, Any],
) -> tuple[list[str], dict[str, list[str]]]:
    """Loop through every route × service_id and export timetable files.

    A failure on one route × service_id is logged and recorded, and the
    remaining combinations are still processed.

    Args:
        ctx: Context with dataframes and configuration, including:
            - routes, trips, stops, calendar, timepoints
            - service_id_schedule_map
            - time_fmt

    Returns:
        The route × service_id combinations that failed, and the flagged
        workbooks (``{workbook path: issues}``).
    """
    routes_df = ctx["routes"]
    calendar_df = ctx["calendar"]

    # 1) Determine final routes to process
    all_routes = get_all_route_short_names(routes_df)
    final_routes = apply_in_out_filters(all_routes)
    logging.info("Final route selection after filters: %s", final_routes)

    # 2) For each route, build timetables by direction & service_id
    failed: list[str] = []
    flagged: dict[str, list[str]] = {}
    for route_short_name in final_routes:
        logging.info("Processing route '%s'...", route_short_name)
        route_ids = routes_df[routes_df["route_short_name"] == route_short_name]["route_id"]
        if route_ids.empty:
            logging.error("Error: Route '%s' not found in routes.txt.", route_short_name)
            continue

        for service_id in calendar_df["service_id"]:
            label = f"route '{route_short_name}', service_id '{service_id}'"
            try:
                if not export_route_service(route_short_name, route_ids, service_id, ctx, flagged):
                    failed.append(label)
            except Exception:
                logging.exception(
                    "Failed to build the timetable for route '%s', service_id '%s'.",
                    route_short_name,
                    service_id,
                )
                failed.append(label)
    return failed, flagged


# -----------------------------------------------------------------------------
# RUN LOG
# -----------------------------------------------------------------------------


# Canonical version lives in utils/run_log.py -- keep this copy in sync.
def extract_config_block(source_file: Path) -> str:
    r"""Return the text between the CONFIG markers in *source_file*.

    Reads ``source_file`` as UTF-8 text and slices out the lines strictly
    *between* the first occurrence of ``# === BEGIN CONFIG ===`` and the first
    subsequent occurrence of ``# === END CONFIG ===``.  The marker lines
    themselves are excluded; whitespace and inline comments inside the block
    are preserved verbatim.

    Args:
        source_file: Path to the Python source file to scan (typically
            ``Path(__file__)`` from the calling script).

    Returns:
        The verbatim text of the configuration block, joined with ``\n``.

    Raises:
        ValueError: If either marker is missing or they appear out of order.
        OSError: If ``source_file`` cannot be read.
    """
    _BEGIN = "# === BEGIN CONFIG ==="
    _END = "# === END CONFIG ==="

    lines: list[str] = source_file.read_text(encoding="utf-8").splitlines()

    begin_idx: int | None = None
    end_idx: int | None = None
    for i, line in enumerate(lines):
        stripped: str = line.strip()
        if begin_idx is None and stripped == _BEGIN:
            begin_idx = i
        elif begin_idx is not None and stripped == _END:
            end_idx = i
            break

    if begin_idx is None or end_idx is None:
        raise ValueError(
            f"Config markers not found in '{source_file}'. Expected '{_BEGIN}' and '{_END}'."
        )

    return "\n".join(lines[begin_idx + 1 : end_idx])


def resolve_source_file() -> Optional[Path]:
    """Path of this script on disk, or ``None`` in an interactive session without ``__file__``."""
    try:
        return Path(__file__).resolve()
    except NameError:
        return None


def write_run_log(
    output_dir: Path,
    summary_lines: Sequence[str],
    flagged: Mapping[str, list[str]],
    failed: Sequence[str],
) -> bool:
    """Write the ``_runlog.txt`` sidecar for this run into *output_dir*.

    The run writes one workbook per route × service_id, so a single log named
    after the script sits above their folders. It lists each flagged workbook
    with the problems found and any failed exports, then captures this
    script's CONFIGURATION block verbatim, between the
    ``# === BEGIN CONFIG ===`` / ``# === END CONFIG ===`` markers, and the
    effective runtime values, including notebook edits.

    Args:
        output_dir: Folder that receives ``RUN_LOG_FILENAME``.
        summary_lines: Human-readable lines describing this run.
        flagged: ``{workbook path: issues}`` for workbooks to review.
        failed: Route × service_id combinations that were not exported.

    Returns:
        ``True`` if the log was written successfully, ``False`` otherwise.
    """
    log_path = output_dir / RUN_LOG_FILENAME

    source_file = resolve_source_file()
    if source_file is None:
        config_text = "(config block unavailable: interactive session, no __file__ on disk)"
        source_display = "<interactive>"
    else:
        try:
            config_text = extract_config_block(source_file)
        except (OSError, ValueError) as exc:
            logging.error("Could not extract config block for run log: %s", exc)
            return False
        source_display = str(source_file)

    flagged_lines: list[str] = []
    for workbook in sorted(flagged):
        flagged_lines.append(os.path.relpath(workbook, output_dir))
        flagged_lines.extend(f"  - {issue}" for issue in flagged[workbook])

    lines: list[str] = [
        "=" * 72,
        "GTFS TIMETABLE EXPORTER RUN LOG",
        "=" * 72,
        f"Run timestamp:    {datetime.now().isoformat(timespec='seconds')}",
        f"Output folder:    {output_dir}",
        f"Source script:    {source_display}",
        "",
        "-" * 72,
        "RUN SUMMARY",
        "-" * 72,
        *summary_lines,
        "",
        "-" * 72,
        f"FLAGGED WORKBOOKS ({len(flagged)}): review before use",
        "-" * 72,
        *(flagged_lines or ["(none)"]),
        "",
        "-" * 72,
        f"FAILED EXPORTS ({len(failed)})",
        "-" * 72,
        *(failed or ["(none)"]),
        "",
        "-" * 72,
        "CONFIGURATION (verbatim from source)",
        "-" * 72,
        config_text,
        "",
        "EFFECTIVE RUNTIME SETTINGS",
        json.dumps(
            {name: value for name, value in globals().items() if name.isupper()},
            indent=2,
            default=str,
        ),
        "=" * 72,
    ]

    try:
        log_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    except OSError as exc:
        logging.error("Error writing run log: %s", exc)
        return False
    logging.info("Run log saved to %s", log_path)
    return True


# -----------------------------------------------------------------------------
# REUSABLE FUNCTIONS
# -----------------------------------------------------------------------------


def load_gtfs_data(
    gtfs_path: str,
    files: Optional[Sequence[str]] = None,
    dtype: str | type[str] | Mapping[str, Any] = str,
    logger: Optional[logging.Logger] = None,
) -> dict[str, pd.DataFrame]:
    """Load one or more GTFS text files into memory.

    Args:
        gtfs_path: Absolute or relative path to the folder containing the
            GTFS feed, or to a ``.zip`` archive of it — the form GTFS
            producers and most open-data portals distribute feeds in. Zip
            members may sit at the archive root or nested one level inside
            a single wrapper folder; both layouts are handled.
        files: Explicit sequence of file names to load. If ``None``,
            the standard 13 GTFS text files are attempted.
        dtype: Value forwarded to :pyfunc:`pandas.read_csv(dtype=…)` to
            control column dtypes. Supply a mapping for per-column dtypes.
        logger: Logger for progress messages. Defaults to this module's
            logger (``logging.getLogger(__name__)``) rather than the root
            logger, so callers keep control of handler configuration.

    Returns:
        Mapping of file stem → :class:`pandas.DataFrame`; for example,
        ``data["trips"]`` holds the parsed *trips.txt* table.

    Raises:
        OSError: Path missing, one of *files* not present in the feed, or
            an OS-level failure while reading a file.
        ValueError: *gtfs_path* is neither a directory nor a valid ``.zip``
            file, a requested file matches more than one location inside
            the zip, a file is empty, or the CSV parser fails.

    Notes:
        All columns default to ``str`` to avoid pandas’ type-inference
        pitfalls (e.g. leading zeros in IDs).
    """
    log = logger if logger is not None else logging.getLogger(__name__)

    if not os.path.exists(gtfs_path):
        raise OSError(f"The path '{gtfs_path}' does not exist.")

    if files is None:
        files = (
            "agency.txt",
            "stops.txt",
            "routes.txt",
            "trips.txt",
            "stop_times.txt",
            "calendar.txt",
            "calendar_dates.txt",
            "fare_attributes.txt",
            "fare_rules.txt",
            "feed_info.txt",
            "frequencies.txt",
            "shapes.txt",
            "transfers.txt",
        )

    is_zip = os.path.isfile(gtfs_path) and gtfs_path.lower().endswith(".zip")
    if not is_zip and not os.path.isdir(gtfs_path):
        raise ValueError(f"'{gtfs_path}' is neither a directory nor a .zip file.")

    archive: zipfile.ZipFile | None = None
    members_by_name: dict[str, list[str]] = {}
    if is_zip:
        try:
            archive = zipfile.ZipFile(gtfs_path)
        except zipfile.BadZipFile as exc:
            raise ValueError(f"'{gtfs_path}' is not a valid zip archive.") from exc
        for name in archive.namelist():
            members_by_name.setdefault(os.path.basename(name), []).append(name)

    try:
        missing: list[str] = []
        ambiguous: list[str] = []
        resolved: dict[str, str] = {}
        for file_name in files:
            if archive is None:
                if not os.path.exists(os.path.join(gtfs_path, file_name)):
                    missing.append(file_name)
                continue
            candidates = members_by_name.get(file_name, [])
            if not candidates:
                missing.append(file_name)
            elif len(candidates) > 1:
                ambiguous.append(file_name)
            else:
                resolved[file_name] = candidates[0]

        if ambiguous:
            raise ValueError(
                f"Ambiguous GTFS files in '{gtfs_path}' (found in multiple "
                f"locations): {', '.join(ambiguous)}"
            )
        if missing:
            raise OSError(f"Missing GTFS files in '{gtfs_path}': {', '.join(missing)}")

        data: dict[str, pd.DataFrame] = {}
        for file_name in files:
            key = file_name.replace(".txt", "")
            try:
                if archive is None:
                    df = pd.read_csv(
                        os.path.join(gtfs_path, file_name), dtype=dtype, low_memory=False
                    )
                else:
                    with archive.open(resolved[file_name]) as handle:
                        df = pd.read_csv(handle, dtype=dtype, low_memory=False)
                data[key] = df
                log.info("Loaded %s (%d records).", file_name, len(df))

            except pd.errors.EmptyDataError as exc:
                raise ValueError(f"File '{file_name}' in '{gtfs_path}' is empty.") from exc

            except pd.errors.ParserError as exc:
                raise ValueError(f"Parser error in '{file_name}' in '{gtfs_path}': {exc}") from exc

        return data
    finally:
        if archive is not None:
            archive.close()


# =============================================================================
# MAIN
# =============================================================================


def main() -> int:
    """Coordinate the end-to-end GTFS → Excel workflow.

    Returns:
        Process exit code: 0 on success, 1 on failure, 2 if required
        CONFIGURATION values are still placeholders.
    """
    logging.basicConfig(
        level=LOG_LEVEL,
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    if (
        GTFS_FOLDER_PATH == r"Path\To\Your\GTFS_Folder"
        or BASE_OUTPUT_PATH == r"Path\To\Your\Output_Folder"
    ):
        logging.warning(
            "GTFS_FOLDER_PATH and/or BASE_OUTPUT_PATH are still set to their default placeholder "
            "values. Please update them in the CONFIGURATION section before running."
        )
        return 2

    if not os.path.exists(BASE_OUTPUT_PATH):
        os.makedirs(BASE_OUTPUT_PATH)

    # Build file list: required + optional that actually exist (folder or .zip)
    available_files = _list_feed_files(GTFS_FOLDER_PATH)
    files_to_load = list(REQUIRED_GTFS_FILES)
    files_to_load += [f for f in OPTIONAL_GTFS_FILES if f in available_files]

    # Load GTFS data
    try:
        data = load_gtfs_data(GTFS_FOLDER_PATH, files=tuple(files_to_load), dtype=str)
        logging.info("Successfully loaded GTFS files overall.")
    except OSError as error:
        logging.error("GTFS data loading error (OS): %s", error)
        if _in_ipython():
            return 1
        raise
    except ValueError as error:
        logging.error("GTFS data loading error (Value): %s", error)
        if _in_ipython():
            return 1
        raise
    except RuntimeError as error:
        logging.error("GTFS data loading error (Runtime): %s", error)
        if _in_ipython():
            return 1
        raise
    except Exception as error:
        logging.error("An unexpected error occurred while loading GTFS files: %s", error)
        if _in_ipython():
            return 1
        raise

    # Ensure required tables are present
    if "calendar" not in data and "calendar_dates" not in data:
        logging.error(
            "Error: neither 'calendar.txt' nor 'calendar_dates.txt' found in loaded GTFS data. "
            "Exiting."
        )
        if _in_ipython():
            return 1
        raise RuntimeError("'calendar.txt' and 'calendar_dates.txt' missing from GTFS data.")

    service_calendar = build_service_calendar(data.get("calendar"), data.get("calendar_dates"))
    calendar_df_filtered = filter_calendar_df(service_calendar)
    if calendar_df_filtered is None:
        logging.info("Exiting due to no service IDs to process.")
        return 0

    if "stop_times" not in data:
        logging.error("Error: 'stop_times.txt' not found in loaded GTFS data. Exiting.")
        if _in_ipython():
            return 1
        raise RuntimeError("'stop_times.txt' missing from loaded GTFS data.")
    timepoints_df = prepare_timepoints(data["stop_times"])

    # Build service_id => schedule_type (with overrides)
    service_id_schedule_map = build_service_id_schedule_map(calendar_df_filtered)
    for sid in FILTER_SERVICE_IDS:
        logging.info("Service %s → %s", sid, service_id_schedule_map.get(sid, "<missing>"))

    # Build context
    required_data_keys = ["stops", "trips", "routes"]
    for key in required_data_keys:
        if key not in data:
            logging.error("Error: '%s.txt' not found in loaded GTFS data. Exiting.", key)
            if _in_ipython():
                return 1
            raise RuntimeError(f"'{key}.txt' missing from loaded GTFS data.")

    ctx = {
        "calendar": calendar_df_filtered,
        "timepoints": timepoints_df,
        "service_id_schedule_map": service_id_schedule_map,
        "time_fmt": TIME_FORMAT_OPTION,
        "stops": data["stops"],
        "trips": data["trips"],
        "routes": data["routes"],
    }

    # Produce timetables
    try:
        failed, flagged = process_route_service_combinations(ctx)
    except Exception as e:
        logging.error("An error occurred during schedule processing: %s", e, exc_info=True)
        if _in_ipython():
            return 1
        raise

    missing_ids = [sid for sid in FILTER_SERVICE_IDS if sid not in service_id_schedule_map]
    summary_lines = [
        f"GTFS feed:            {GTFS_FOLDER_PATH}",
        "Service labels:       "
        + ", ".join(f"{sid} → {label}" for sid, label in service_id_schedule_map.items()),
        f"Service IDs missing:  {', '.join(missing_ids) or '-'}",
        f"Flagged workbooks:    {len(flagged)}",
        f"Failed exports:       {len(failed)}",
    ]
    run_log_path = Path(BASE_OUTPUT_PATH) / RUN_LOG_FILENAME
    if not write_run_log(Path(BASE_OUTPUT_PATH), summary_lines, flagged, failed):
        if REQUIRE_RUN_LOG:
            logging.error(
                "Run log could not be written to %s. Set REQUIRE_RUN_LOG = False to "
                "continue without one.",
                run_log_path,
            )
            return 1
        logging.warning("Continuing without a run log (REQUIRE_RUN_LOG is False).")
    if flagged:
        logging.warning(
            "%d workbook(s) flagged for review; see the run log at %s.",
            len(flagged),
            run_log_path,
        )
    if failed:
        logging.error(
            "Script finished with %d route/service timetable(s) not exported; see errors "
            "above and the run log.",
            len(failed),
        )
        return 1
    logging.info("Script completed successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
