"""List the GTFS stops inside zone polygons as config text for the bay scripts.

ArcPy port of ``cluster_stops_from_zones_gpd.py`` for ArcGIS Pro's bundled
Python (arcpy + pandas, no geopandas or shapely). Reads a polygon layer of
zones (for example, one polygon around each transit center) and a GTFS feed,
finds every stop inside or on the edge of each zone, and logs the stop_id
lists as text to paste into the CONFIGURATION sections of
``block_status_timeline_exporter.py`` (Step 1), ``bay_usage_analyzer.py``
(Step 2) and ``bay_change_sweep.py`` (Step 3). Every stop is listed as a
single bay; add capacity (double or triple bays, overflow spaces) by hand
after pasting. The three scripts are neither read nor changed.

Zones may touch but not overlap, and a stop on the shared edge of two zones
stops the run: a stop may belong to only one cluster. Polygons that share a
name form one zone. Every stop (location_type 0 or blank) in a zone is
listed, whether or not it has scheduled trips.

Zones without stops are reported and left out. Stops outside every zone but
within ``NEAR_MISS_FEET`` of one (near misses, often a zone drawn a little
short) are reported against their nearest zone, so they can be added by
redrawing it.

Inputs
------
- GTFS feed folder or .zip (stops.txt): the feed Step 1 reads, so the
  stop_ids match. The optional workbook also reads routes.txt, trips.txt
  and stop_times.txt.
- Zone polygon layer arcpy reads (shapefile, file geodatabase feature class,
  ...) with a defined coordinate system. A name field is optional; unnamed
  zones are numbered by their position in the layer.

Outputs
-------
- The config text, written to the log. Its header lists zones without stops
  and near misses, which are also logged as warnings.
- Optionally, the same text as ``OUTPUT_FILENAME`` in ``OUTPUT_DIR`` with a
  ``_runlog.txt`` sidecar capturing the verbatim CONFIGURATION block and
  SHA-256 fingerprints of the inputs.
- Optional Excel workbook: an index and one sheet per cluster, with stop
  names and IDs across the top and unique route-direction entries below.
  It includes all services in the feed and all scheduled stop visits,
  including no-pickup visits. It is an inventory, not a boarding guide or
  an optimized assignment. Requires openpyxl only when this export is on.

Notes:
------
Unlike the twin, polygons are checked against Esri's geometry rules (the
Check Geometry tool) rather than OGC validity, the point-in-polygon and
overlap tests use ArcGIS's geometry engine, and stops are projected with no
datum transformation unless GEO_TRANSFORMATION names one.

Typical usage
-------------
Update the paths in the CONFIGURATION section (or pass the matching CLI
flags) and run from a shell, ArcGIS Pro's Python window, or a Jupyter
notebook using ArcGIS Pro's bundled Python. Paste each block over the
matching setting, then add capacities.
"""

from __future__ import annotations

import argparse
import datetime as datetime_module  # the from-import errored in some ArcGIS Pro setups
import hashlib
import json
import logging
import math
import os
import re
import sys
import tempfile
import zipfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any, Dict, List, NamedTuple, Optional, Tuple

import arcpy
import pandas as pd

# =============================================================================
# CONFIGURATION
# =============================================================================
# === BEGIN CONFIG ===

# GTFS feed folder or .zip. Use the feed Step 1 (block_status_timeline_exporter.py)
# reads so the stop_ids match.
GTFS_FOLDER_PATH: str = r"Path\To\Your\GTFS_data"

# Zone polygons, e.g. one around each transit center: a shapefile, a file
# geodatabase feature class (r"C:\Data\Facilities.gdb\Zones") or another polygon
# layer arcpy reads. The layer needs a coordinate system; stops are projected into it.
ZONES_PATH: str = r"Path\To\Your\Zones.shp"

# Optional attribute that names each zone, e.g. "NAME". Polygons that share a
# name form one zone. Leave "" to number zones by position in the layer ("Zone 1"
# is the first polygon); a polygon with a blank name is numbered the same way.
ZONE_NAME_FIELD: str = ""

# Report stops outside every zone but within this many feet of one: a zone drawn a
# little short misses them. 0 turns the check off.
NEAR_MISS_FEET: float = 250.0

# Optional geographic (datum) transformation for projecting the stops (WGS84)
# into the zone layer's coordinate system, e.g. "WGS_1984_(ITRF00)_To_NAD_1983".
# "" applies none; name one if the datum shift matters for stops near a zone edge.
GEO_TRANSFORMATION: str = ""

# Output folder for the config text, workbook and run logs. Leave OUTPUT_DIR
# "" for log-only output when WRITE_BAY_ASSIGNMENT_XLSX is False.
OUTPUT_DIR: str = r""
OUTPUT_FILENAME: str = r"cluster_stops_from_zones.txt"

# Optional scheduled bay-use inventory. Set OUTPUT_DIR to save the workbook.
# All services are combined; no date, day-of-week or pickup filter is applied.
# Set False (or pass --no-bay-assignment-xlsx) for the original text-only workflow.
WRITE_BAY_ASSIGNMENT_XLSX: bool = True
BAY_ASSIGNMENT_FILENAME: str = r"cluster_bay_assignments.xlsx"

# Every output must be traceable: a failed run-log write aborts the script.
# Set to False only when writing to a genuinely read-only location.
REQUIRE_RUN_LOG: bool = True

LOG_LEVEL: int = logging.INFO  # DEBUG / INFO / WARNING / ERROR

# === END CONFIG ===

# Longest generated config line: the repository's ruff line length. A stop's
# comment is shortened to fit.
MAX_LINE_LENGTH = 100

# The scripts the generated blocks are pasted into.
STEP1_SCRIPT = "scripts/gtfs_exports/block_status_timeline_exporter.py"
STEP2_SCRIPT = "scripts/facilities_tools/bay_usage_analyzer.py"
STEP3_SCRIPT = "scripts/facilities_tools/bay_change_sweep.py"

ASSIGNMENT_COLUMNS = ["zone", "stop_id", "stop_code", "stop_name"]
NEAR_MISS_COLUMNS = [*ASSIGNMENT_COLUMNS, "feet"]

METERS_PER_FOOT = 0.3048

# stops.txt coordinates are WGS84 longitude/latitude.
WGS84_WKID = 4326

# =============================================================================
# GTFS STOPS
# =============================================================================


# Based on utils/gtfs_helpers.py; this variant preserves literal IDs such as "NA".
def load_gtfs_data(
    gtfs_path: str,
    files: Optional[Sequence[str]] = None,
    dtype: str | type[str] | Mapping[str, Any] = str,
    logger: Optional[logging.Logger] = None,
) -> dict[str, pd.DataFrame]:
    """Load one or more GTFS text files into memory.

    Args:
        gtfs_path: Absolute or relative path to the folder containing the
            GTFS feed, or to a ``.zip`` archive of it — the form GTFS
            producers and most open-data portals distribute feeds in. Zip
            members may sit at the archive root or nested one level inside
            a single wrapper folder; both layouts are handled.
        files: Explicit sequence of file names to load. If ``None``,
            the standard 13 GTFS text files are attempted.
        dtype: Value forwarded to :pyfunc:`pandas.read_csv(dtype=…)` to
            control column dtypes. Supply a mapping for per-column dtypes.
        logger: Logger for progress messages. Defaults to this module's
            logger (``logging.getLogger(__name__)``) rather than the root
            logger, so callers keep control of handler configuration.

    Returns:
        Mapping of file stem → :class:`pandas.DataFrame`; for example,
        ``data["trips"]`` holds the parsed *trips.txt* table.

    Raises:
        OSError: Path missing, one of *files* not present in the feed, or
            an OS-level failure while reading a file.
        ValueError: *gtfs_path* is neither a directory nor a valid ``.zip``
            file, a requested file matches more than one location inside
            the zip, a file is empty, or the CSV parser fails.

    Notes:
        All columns default to ``str`` to avoid pandas’ type-inference
        pitfalls (e.g. leading zeros in IDs). Default NA tokens are disabled
        so IDs and names such as "NA" and "NULL" remain literal text.
    """
    log = logger if logger is not None else logging.getLogger(__name__)

    if not os.path.exists(gtfs_path):
        raise OSError(f"The path '{gtfs_path}' does not exist.")

    if files is None:
        files = (
            "agency.txt",
            "stops.txt",
            "routes.txt",
            "trips.txt",
            "stop_times.txt",
            "calendar.txt",
            "calendar_dates.txt",
            "fare_attributes.txt",
            "fare_rules.txt",
            "feed_info.txt",
            "frequencies.txt",
            "shapes.txt",
            "transfers.txt",
        )

    is_zip = os.path.isfile(gtfs_path) and gtfs_path.lower().endswith(".zip")
    if not is_zip and not os.path.isdir(gtfs_path):
        raise ValueError(f"'{gtfs_path}' is neither a directory nor a .zip file.")

    archive: zipfile.ZipFile | None = None
    members_by_name: dict[str, list[str]] = {}
    if is_zip:
        try:
            archive = zipfile.ZipFile(gtfs_path)
        except zipfile.BadZipFile as exc:
            raise ValueError(f"'{gtfs_path}' is not a valid zip archive.") from exc
        for name in archive.namelist():
            members_by_name.setdefault(os.path.basename(name), []).append(name)

    try:
        missing: list[str] = []
        ambiguous: list[str] = []
        resolved: dict[str, str] = {}
        for file_name in files:
            if archive is None:
                if not os.path.exists(os.path.join(gtfs_path, file_name)):
                    missing.append(file_name)
                continue
            candidates = members_by_name.get(file_name, [])
            if not candidates:
                missing.append(file_name)
            elif len(candidates) > 1:
                ambiguous.append(file_name)
            else:
                resolved[file_name] = candidates[0]

        if ambiguous:
            raise ValueError(
                f"Ambiguous GTFS files in '{gtfs_path}' (found in multiple "
                f"locations): {', '.join(ambiguous)}"
            )
        if missing:
            raise OSError(f"Missing GTFS files in '{gtfs_path}': {', '.join(missing)}")

        data: dict[str, pd.DataFrame] = {}
        for file_name in files:
            key = file_name.replace(".txt", "")
            try:
                if archive is None:
                    df = pd.read_csv(
                        os.path.join(gtfs_path, file_name),
                        dtype=dtype,
                        low_memory=False,
                        keep_default_na=False,
                    )
                else:
                    with archive.open(resolved[file_name]) as handle:
                        df = pd.read_csv(
                            handle, dtype=dtype, low_memory=False, keep_default_na=False
                        )
                data[key] = df
                log.info("Loaded %s (%d records).", file_name, len(df))

            except pd.errors.EmptyDataError as exc:
                raise ValueError(f"File '{file_name}' in '{gtfs_path}' is empty.") from exc

            except pd.errors.ParserError as exc:
                raise ValueError(f"Parser error in '{file_name}' in '{gtfs_path}': {exc}") from exc

        return data
    finally:
        if archive is not None:
            archive.close()


def prepare_stops(stops: pd.DataFrame) -> pd.DataFrame:
    """Return the stops in stops.txt that buses serve, with numeric coordinates.

    Keeps location_type 0 or blank, the only kind stop_times.txt can name;
    stations, entrances, generic nodes and boarding areas are dropped. Rows
    without usable coordinates are dropped with a warning. stop_code and
    stop_name are optional and become blank when absent.

    Args:
        stops: stops.txt as read by :func:`load_gtfs_data` (string columns).

    Returns:
        Columns stop_id, stop_code, stop_name, stop_lat and stop_lon.

    Raises:
        ValueError: A required column is missing, a stop_id repeats, or no
            stop with coordinates remains.
    """
    missing = [column for column in ("stop_id", "stop_lat", "stop_lon") if column not in stops]
    if missing:
        raise ValueError(f"stops.txt is missing column(s): {', '.join(missing)}.")

    def text(column: str) -> pd.Series:
        values = stops[column] if column in stops else pd.Series("", index=stops.index)
        return values.fillna("").astype(str)

    stop_id = text("stop_id")
    has_id = stop_id.str.strip().ne("")
    if not has_id.all():
        logging.warning("Ignoring %d stops.txt row(s) without a stop_id.", int((~has_id).sum()))
    repeated = stop_id[has_id][stop_id[has_id].duplicated()]
    if not repeated.empty:
        raise ValueError(
            f"stops.txt repeats stop_id {_preview(sorted(set(repeated)))}; stop_ids must be "
            "unique (Step 1 refuses this feed too)."
        )
    is_stop = text("location_type").str.strip().isin(["", "0"])
    lat = pd.to_numeric(stops["stop_lat"], errors="coerce")
    lon = pd.to_numeric(stops["stop_lon"], errors="coerce")
    usable = lat.between(-90, 90) & lon.between(-180, 180)
    unplaced = has_id & is_stop & ~usable
    if unplaced.any():
        logging.warning(
            "Ignoring %d stop(s) without usable coordinates: %s",
            int(unplaced.sum()),
            _preview(stop_id[unplaced].tolist()),
        )
    keep = has_id & is_stop & usable
    if not keep.any():
        raise ValueError("stops.txt has no stops (location_type 0 or blank) with coordinates.")
    prepared = pd.DataFrame(
        {
            "stop_id": stop_id,
            "stop_code": text("stop_code").map(_one_line),
            "stop_name": text("stop_name").map(_one_line),
            "stop_lat": lat,
            "stop_lon": lon,
        }
    )[keep].reset_index(drop=True)
    logging.info("Read %d stops with coordinates.", len(prepared))
    return prepared


# =============================================================================
# ZONES
# =============================================================================


class Zones(NamedTuple):
    """Zone polygons in layer order, as read by :func:`load_zones`."""

    names: List[str]  # each polygon's zone name
    polygons: List[Any]  # each polygon's arcpy.Polygon, in spatial_reference
    spatial_reference: Any  # the zone layer's arcpy.SpatialReference


def load_zones(zones_path: str, name_field: str = "") -> Zones:
    """Read the zone polygons, check them, and name each one.

    Args:
        zones_path: Polygon layer arcpy reads (shapefile, geodatabase feature
            class, ...).
        name_field: Attribute naming each zone, matched regardless of case as
            ArcGIS matches field names, or "" to number zones by position in
            the layer.

    Returns:
        Each polygon's zone name and geometry in layer order (the n-th
        polygon is polygon n in messages), with the layer's spatial reference.

    Raises:
        OSError: *zones_path* does not exist.
        ValueError: The layer cannot be read, is not a polygon layer, is
            empty, has no coordinate system or lacks *name_field*; a polygon
            is missing or breaks Esri's geometry rules; a longitude/latitude
            layer has coordinates out of range; or an unnamed polygon's
            number is another polygon's name.
    """
    if not arcpy.Exists(zones_path):
        raise OSError(f"Zone layer not found: {zones_path}")
    try:
        description = arcpy.Describe(zones_path)
        fields = arcpy.ListFields(zones_path)
    except (OSError, RuntimeError) as exc:
        raise ValueError(f"Could not read zone layer '{zones_path}': {exc}") from exc
    shape_type = getattr(description, "shapeType", "")
    if shape_type != "Polygon":
        kind = shape_type or getattr(description, "dataType", "unknown")
        raise ValueError(f"Zones must be polygons; '{zones_path}' is not a polygon layer ({kind}).")
    sr = description.spatialReference
    if sr is None or sr.name in ("", "Unknown"):
        raise ValueError(
            f"Zone layer '{zones_path}' has no coordinate system (missing or unknown .prj). "
            "Define its projection before running."
        )
    attributes = {field.name.lower(): field.name for field in fields if field.type != "Geometry"}
    field_name = attributes.get(name_field.lower(), "") if name_field else ""
    if name_field and not field_name:
        raise ValueError(
            f"ZONE_NAME_FIELD '{name_field}' not found in '{zones_path}'. "
            f"Available fields: {', '.join(sorted(attributes.values())) or '(none)'}"
        )
    oids: List[int] = []
    polygons: List[Any] = []
    values: List[Any] = []
    cursor_fields = ["OID@", "SHAPE@"] + ([field_name] if field_name else [])
    try:
        # The explicit spatial reference keeps the geometries in the layer's own
        # coordinates even when arcpy.env.outputCoordinateSystem is set.
        with arcpy.da.SearchCursor(zones_path, cursor_fields, spatial_reference=sr) as cursor:
            for row in cursor:
                oids.append(row[0])
                polygons.append(row[1])
                values.append(row[2] if field_name else None)
    except (OSError, RuntimeError, UnicodeDecodeError) as exc:
        raise ValueError(f"Could not read zone layer '{zones_path}': {exc}") from exc
    if not polygons:
        raise ValueError(f"Zone layer '{zones_path}' has no polygons.")
    numbers = list(range(1, len(polygons) + 1))
    _check_polygons(zones_path, polygons, oids)
    if sr.type == "Geographic":
        extents = [polygon.extent for polygon in polygons]
        west = min(extent.XMin for extent in extents)
        south = min(extent.YMin for extent in extents)
        east = max(extent.XMax for extent in extents)
        north = max(extent.YMax for extent in extents)
        if west < -180 or east > 180 or south < -90 or north > 90:
            raise ValueError(
                f"Zone layer '{zones_path}' says it is in longitude/latitude "
                f"({_sr_label(sr)}), but its coordinates are out of range. The layer "
                "probably holds projected coordinates without saying so; set its real "
                "coordinate system with the Define Projection tool and rerun."
            )
    names = _zone_names(values if field_name else None, numbers)
    counts: Dict[str, int] = {}
    for name in names:
        counts[name] = counts.get(name, 0) + 1
    for zone, count in counts.items():
        if count > 1:
            logging.info("Zone '%s' combines %d polygons.", zone, count)
    logging.info("Read %d zone(s) from %d polygon(s).", len(counts), len(names))
    return Zones(names, polygons, sr)


def _check_polygons(zones_path: str, polygons: List[Any], oids: List[int]) -> None:
    """Raise ValueError naming zone polygons that are missing or break Esri's geometry rules.

    Args:
        zones_path: The zone layer, which the Check Geometry tool reads.
        polygons: Each polygon's geometry (None when missing), in layer order.
        oids: Each polygon's ObjectID, aligned with *polygons*.
    """
    missing = [
        n
        for n, polygon in enumerate(polygons, start=1)
        if polygon is None or polygon.pointCount == 0
    ]
    if missing:
        raise ValueError(f"Zone polygon(s) {_preview(missing)} have no geometry.")
    position = {oid: n for n, oid in enumerate(oids, start=1)}
    invalid = [
        f"{position.get(oid, f'OID {oid}')} ({problem})"
        for oid, problem in _geometry_problems(zones_path)
    ]
    if invalid:
        raise ValueError(
            f"Invalid zone polygon(s) {_preview(invalid)}. Repair them (for example with the "
            "Repair Geometry tool) and rerun."
        )


def _geometry_problems(zones_path: str) -> List[Tuple[int, str]]:
    """Return the (ObjectID, problem) pairs the Check Geometry tool reports for *zones_path*.

    Raises:
        ValueError: The tool could not check the layer.
    """
    table = arcpy.CreateUniqueName("zone_geometry_check", "memory")
    try:
        arcpy.management.CheckGeometry(zones_path, table)
        with arcpy.da.SearchCursor(table, ["FEATURE_ID", "PROBLEM"]) as cursor:
            return [(int(oid), _one_line(str(problem))) for oid, problem in cursor]
    except arcpy.ExecuteError as exc:
        raise ValueError(f"Could not check the polygons of '{zones_path}': {exc}") from exc
    finally:
        if arcpy.Exists(table):
            arcpy.management.Delete(table)


def _sr_label(sr: Any) -> str:
    """Name a spatial reference for messages, with its WKID when it has one."""
    code = getattr(sr, "factoryCode", 0)
    return f"{sr.name}, WKID {code}" if code else str(sr.name)


def _zone_names(values: Optional[List[Any]], numbers: List[int]) -> List[str]:
    """Name each polygon from *values*, numbering blank or absent names by position.

    Raises:
        ValueError: An unnamed polygon's number ("Zone 3") is another
            polygon's name, which would merge the two silently.
    """
    given = [_clean_name(value) for value in values] if values is not None else [""] * len(numbers)
    unnamed = [n for name, n in zip(given, numbers) if not name]
    taken = set(given)
    clashes = [f"Zone {n}" for n in unnamed if f"Zone {n}" in taken]
    if clashes:
        raise ValueError(
            f"Unnamed polygons would be numbered {_preview(clashes)}, which other polygons "
            "are named. Name the unnamed polygons or rename the others."
        )
    if values is not None and unnamed:
        logging.warning(
            "%d zone polygon(s) have no name and are numbered by position: %s",
            len(unnamed),
            _preview(unnamed),
        )
    return [name or f"Zone {n}" for name, n in zip(given, numbers)]


def _clean_name(value: Any) -> str:
    """Render an attribute value as a zone name; a missing value is blank.

    Numeric names read back as floats (``12.0``) are written without the
    decimal, and runs of whitespace collapse to single spaces.
    """
    if value is None or (pd.api.types.is_scalar(value) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _one_line(str(value))


def check_zone_overlaps(zones: Zones) -> None:
    """Raise ValueError if two zones share any area; touching edges are allowed.

    Polygons with the same name are one zone, so they may overlap each other.

    Args:
        zones: Output of :func:`load_zones`.

    Raises:
        ValueError: Names each pair of overlapping zones.
    """
    shapes: Dict[str, Any] = {}
    for name, polygon in zip(zones.names, zones.polygons):
        shapes[name] = shapes[name].union(polygon) if name in shapes else polygon
    names = list(shapes)
    geometries = [shapes[name] for name in names]
    # Two polygons share area when they meet (not disjoint) other than along their
    # boundaries (not touching): the DE-9IM "interiors intersect" test.
    overlaps = [
        f"'{names[i]}' and '{names[j]}'"
        for i in range(len(names))
        for j in range(i + 1, len(names))
        if not geometries[i].disjoint(geometries[j]) and not geometries[i].touches(geometries[j])
    ]
    if overlaps:
        raise ValueError(
            f"Zones overlap: {_preview(overlaps, separator='; ')}. A stop may belong to only "
            "one cluster, so zones may touch but not overlap. Fix the polygons and rerun."
        )


# =============================================================================
# STOPS IN ZONES
# =============================================================================


def assign_stops_to_zones(
    stops: pd.DataFrame, zones: Zones, transformation: str = ""
) -> pd.DataFrame:
    """Find the zone each stop falls in, counting a stop on a zone's edge as inside.

    Stops are projected into the zone layer's coordinate system, through the
    geographic (datum) *transformation* when one is named. Zones without
    stops are reported and left out.

    Args:
        stops: Output of :func:`prepare_stops`.
        zones: Output of :func:`load_zones`.
        transformation: Geographic transformation from WGS84 to the zone
            layer's datum, or "" for none.

    Returns:
        One row per stop in a zone (zone, stop_id, stop_code, stop_name),
        zones in layer order and stop_ids in natural order within each zone.

    Raises:
        ValueError: The stops cannot be projected, no stop falls in any
            zone, or a stop lies on the shared edge of two zones.
    """
    points = _project_stops(stops, zones.spatial_reference, transformation)
    # Each zone's extent, grown by the XY tolerance so the quick box test never
    # skips a stop the geometry engine would count as on the zone's edge.
    pad = _xy_tolerance(zones.spatial_reference)
    boxes = [
        (box.XMin - pad, box.YMin - pad, box.XMax + pad, box.YMax + pad)
        for box in (polygon.extent for polygon in zones.polygons)
    ]
    hits = []
    for stop, point in zip(
        stops[["stop_id", "stop_code", "stop_name"]].itertuples(index=False), points
    ):
        x, y = point.firstPoint.X, point.firstPoint.Y
        for zone, polygon, (west, south, east, north) in zip(zones.names, zones.polygons, boxes):
            if west <= x <= east and south <= y <= north and not polygon.disjoint(point):
                hits.append((zone, *stop))
    pairs = pd.DataFrame(hits, columns=ASSIGNMENT_COLUMNS).drop_duplicates(["zone", "stop_id"])
    if pairs.empty:
        raise ValueError(
            "No stop falls inside any zone. Check that the zones cover the stops and that the "
            "zone layer's coordinate system is right (a wrong one puts the polygons somewhere "
            "else)."
        )
    zones_by_stop = pairs.groupby("stop_id", sort=False)["zone"].agg(list)
    shared = [
        f"{stop} ({' | '.join(names)})" for stop, names in zones_by_stop.items() if len(names) > 1
    ]
    if shared:
        raise ValueError(
            f"Stop(s) on the edge of two zones: {_preview(shared, separator='; ')}. A stop may "
            "belong to only one cluster; move the shared edge so each stop is in one zone."
        )
    order = {zone: i for i, zone in enumerate(dict.fromkeys(zones.names))}
    records = sorted(
        pairs.to_dict("records"),
        key=lambda row: (order[row["zone"]], _natural_key(str(row["stop_id"]))),
    )
    assignments = pd.DataFrame(records, columns=ASSIGNMENT_COLUMNS)
    empty = [zone for zone in order if zone not in set(assignments["zone"])]
    if empty:
        logging.warning(
            "No stops in zone(s) %s; they are left out of the config text.", _preview(empty)
        )
    logging.info("Found %d stop(s) in %d zone(s).", len(assignments), assignments["zone"].nunique())
    return assignments


def find_near_misses(
    stops: pd.DataFrame,
    zones: Zones,
    assignments: pd.DataFrame,
    max_feet: float,
    transformation: str = "",
) -> pd.DataFrame:
    """Find the stops outside every zone but within *max_feet* of one, and report them.

    Each stop is listed once, against its nearest zone. Distances are measured
    in the zone layer's coordinate system when it is projected, otherwise in
    the WGS84 UTM zone at the zones' center.

    Args:
        stops: Output of :func:`prepare_stops`.
        zones: Output of :func:`load_zones`.
        assignments: Output of :func:`assign_stops_to_zones`.
        max_feet: How far outside a zone to look, in feet; 0 looks nowhere.
        transformation: Geographic transformation from WGS84 to the zone
            layer's datum, or "" for none (as for :func:`assign_stops_to_zones`).

    Returns:
        One row per near miss (zone, stop_id, stop_code, stop_name, feet),
        zones in layer order and the closest stop first within each zone.

    Raises:
        ValueError: The stops or zones cannot be projected for measuring.
    """
    outside = stops[~stops["stop_id"].isin(assignments["stop_id"])]
    if max_feet <= 0 or outside.empty:
        return pd.DataFrame(columns=NEAR_MISS_COLUMNS)
    points = _project_stops(outside, zones.spatial_reference, transformation)
    polygons = list(zones.polygons)
    sr = zones.spatial_reference
    if sr.type == "Projected":
        feet_per_unit = float(sr.metersPerUnit) / METERS_PER_FOOT
    else:
        # Longitude/latitude: measure in meters in the UTM zone at the zones' center.
        sr = _utm_reference(polygons)
        feet_per_unit = 1 / METERS_PER_FOOT
        try:
            points = [point.projectAs(sr) for point in points]
            polygons = [polygon.projectAs(sr) for polygon in polygons]
        except (RuntimeError, ValueError) as exc:
            raise ValueError(
                f"Could not project the zones and stops into {_sr_label(sr)} to measure "
                f"near misses ({exc})."
            ) from exc
    reach = max_feet / feet_per_unit
    boxes = [
        (box.XMin - reach, box.YMin - reach, box.XMax + reach, box.YMax + reach)
        for box in (polygon.extent for polygon in polygons)
    ]
    rows = []
    for stop, point in zip(
        outside[["stop_id", "stop_code", "stop_name"]].itertuples(index=False), points
    ):
        x, y = point.firstPoint.X, point.firstPoint.Y
        nearest: Optional[Tuple[float, str]] = None
        for zone, polygon, (west, south, east, north) in zip(zones.names, polygons, boxes):
            if not (west <= x <= east and south <= y <= north):
                continue
            feet = polygon.distanceTo(point) * feet_per_unit
            if feet <= max_feet and (nearest is None or feet < nearest[0]):
                nearest = (feet, zone)
        if nearest is not None:
            rows.append((nearest[1], *stop, nearest[0]))
    order = {zone: i for i, zone in enumerate(dict.fromkeys(zones.names))}
    records = sorted(
        pd.DataFrame(rows, columns=NEAR_MISS_COLUMNS).to_dict("records"),
        key=lambda row: (order[row["zone"]], row["feet"], _natural_key(str(row["stop_id"]))),
    )
    near_misses = pd.DataFrame(records, columns=NEAR_MISS_COLUMNS)
    if not near_misses.empty:
        found = [
            f"{row['stop_id']} ({_feet(row['feet'])} ft from {row['zone']})"
            for row in near_misses.to_dict("records")
        ]
        logging.warning(
            "%d stop(s) are outside every zone but within %s ft of one; redraw a zone to add "
            "any that belong (all are listed in the config text header): %s",
            len(near_misses),
            f"{max_feet:,g}",
            _preview(found, separator="; "),
        )
    return near_misses


def _utm_reference(polygons: List[Any]) -> Any:
    """Return the WGS84 UTM spatial reference for the zone at the center of lon/lat *polygons*."""
    extents = [polygon.extent for polygon in polygons]
    lon = (min(box.XMin for box in extents) + max(box.XMax for box in extents)) / 2
    lat = (min(box.YMin for box in extents) + max(box.YMax for box in extents)) / 2
    utm_zone = min(int((lon + 180) // 6) + 1, 60)
    return arcpy.SpatialReference((32600 if lat >= 0 else 32700) + utm_zone)


def _feet(value: float) -> str:
    """Format a stop's distance in whole feet, with a thousands separator."""
    return f"{value:,.0f}"


def _project_stops(stops: pd.DataFrame, sr: Any, transformation: str) -> List[Any]:
    """Project each stop's WGS84 coordinates into *sr*, in *stops* order.

    Raises:
        ValueError: arcpy could not project the stops, for example with the
            named *transformation*.
    """
    wgs84 = arcpy.SpatialReference(WGS84_WKID)
    through = f" through {transformation}" if transformation else ""
    logging.info("Projecting %d stop(s) into %s%s.", len(stops), _sr_label(sr), through)
    points: List[Any] = []
    try:
        for lon, lat in zip(stops["stop_lon"], stops["stop_lat"]):
            point = arcpy.PointGeometry(arcpy.Point(float(lon), float(lat)), wgs84)
            points.append(
                point.projectAs(sr, transformation) if transformation else point.projectAs(sr)
            )
    except (RuntimeError, ValueError) as exc:
        hint = " Check GEO_TRANSFORMATION." if transformation else ""
        raise ValueError(
            f"Could not project the stops into {_sr_label(sr)}{through} ({exc}).{hint}"
        ) from exc
    return points


def _xy_tolerance(sr: Any) -> float:
    """Return the XY tolerance of *sr* in its own units, or 0.0 when it has none."""
    try:
        tolerance = float(sr.XYTolerance)
    except (AttributeError, TypeError, ValueError):
        return 0.0
    return tolerance if math.isfinite(tolerance) and tolerance > 0 else 0.0


def _natural_key(text: str) -> Tuple[Tuple[int, int, str], ...]:
    """Sort key that orders digit runs by value, so "7" < "65" < "2956"."""
    parts = re.split(r"(\d+)", text)
    return tuple((0, int(part), "") if i % 2 else (1, 0, part) for i, part in enumerate(parts)) + (
        (1, 0, text),
    )


# =============================================================================
# CONFIG TEXT
# =============================================================================


def build_config_text(
    assignments: pd.DataFrame,
    zone_order: Sequence[str],
    zones_path: str,
    gtfs_path: str,
    near_misses: Optional[pd.DataFrame] = None,
    near_miss_feet: float = 0.0,
) -> str:
    """Assemble the text to paste: a header, then one section per step.

    Args:
        assignments: Output of :func:`assign_stops_to_zones`.
        zone_order: Every zone name in layer order, including zones without stops.
        zones_path: Zone layer path, quoted in the header.
        gtfs_path: GTFS feed path, quoted in the header.
        near_misses: Output of :func:`find_near_misses`, listed in the header.
        near_miss_feet: The distance *near_misses* was found within; 0 means
            the check was off.

    Returns:
        The config text, ending with a newline.
    """
    counts = assignments.groupby("zone", sort=False).size()
    empty = [zone for zone in zone_order if zone not in counts.index]
    rule = "# " + "=" * 76
    header = [
        rule,
        "# Cluster stop lists from cluster_stops_from_zones_arcpy.py",
        f"# Zones: {zones_path}",
        f"# GTFS:  {gtfs_path}",
        *(f"#   {zone}: {count} stop(s)" for zone, count in counts.items()),
        *([f"#   No stops, left out: {', '.join(empty)}"] if empty else []),
        *_near_miss_lines(near_misses, near_miss_feet),
        "# Every stop is listed as a single bay. Add capacity by hand after pasting.",
        rule,
    ]
    sections = [
        "\n".join(header),
        step1_text(assignments),
        step2_text(assignments),
        step3_text(assignments),
    ]
    return "\n\n".join(sections) + "\n"


def _near_miss_lines(near_misses: Optional[pd.DataFrame], near_miss_feet: float) -> List[str]:
    """Header lines listing each near miss with its zone, distance, code and name."""
    if near_miss_feet <= 0:
        return ["# Near-miss check off (NEAR_MISS_FEET = 0)."]
    within = f"outside every zone but within {near_miss_feet:,g} ft"
    if near_misses is None or near_misses.empty:
        return [f"# Near misses, {within}: none."]
    lines = [f"# Near misses, {within}. Redraw a zone to add any that belong:"]
    for row in near_misses.to_dict("records"):
        stop = str(row["stop_id"])
        note = _describe(stop, str(row["stop_code"]), str(row["stop_name"]))
        line = f"#   {row['zone']}: {stop}, {_feet(row['feet'])} ft"
        lines.append(_fit(f"{line} | {note}" if note else line))
    return lines


def step1_text(assignments: pd.DataFrame) -> str:
    """Step 1's CLUSTER_DEFINITIONS, plus its optional block filters."""
    lines = _banner(
        f"Step 1: {STEP1_SCRIPT}",
        'Replace CLUSTER_DEFINITIONS. Step 1 reads only each zone\'s "stops";',
        "capacity is set in Steps 2 and 3.",
    )
    lines.append("CLUSTER_DEFINITIONS = {")
    for zone, group in assignments.groupby("zone", sort=False):
        lines += [f"    {_literal(str(zone))}: {{", '        "stops": [']
        lines += _stop_lines(group, " " * 12)
        lines += ["        ],", "    },"]
    lines.append("}")
    lines += [
        "",
        "# Optional: process only the blocks that touch these stops. Step 1 matches",
        "# either list, so one is enough.",
        "STOP_ID_FILTER: list[str] = [",
    ]
    for zone, group in assignments.groupby("zone", sort=False):
        lines.append(f"    # {zone}")
        lines += [f"    {_literal(stop)}," for stop in group["stop_id"]]
    lines.append("]")
    codes = assignments[assignments["stop_code"].ne("")].drop_duplicates("stop_code")
    if codes.empty:
        lines.append("STOP_CODE_FILTER: list[str] = []  # no stop in the zones has a stop_code")
        return "\n".join(lines)
    lines.append("STOP_CODE_FILTER: list[str] = [")
    for zone, group in codes.groupby("zone", sort=False):
        lines.append(f"    # {zone}")
        lines += [f"    {_literal(code)}," for code in group["stop_code"]]
    lines.append("]")
    return "\n".join(lines)


def step2_text(assignments: pd.DataFrame) -> str:
    """Step 2's CLUSTER_DEFINITIONS with every stop as a single bay."""
    lines = _banner(
        f"Step 2: {STEP2_SCRIPT}",
        "Replace CLUSTER_DEFINITIONS. Every stop starts as a single bay: move stops",
        "that hold two or three buses to double_bay_stops or triple_bay_stops, and",
        "name any overflow spaces.",
    )
    lines.append("CLUSTER_DEFINITIONS: Dict[str, Dict[str, List[str]]] = {")
    for zone, group in assignments.groupby("zone", sort=False):
        lines += [f"    {_literal(str(zone))}: {{", '        "single_bay_stops": [']
        lines += _stop_lines(group, " " * 12)
        lines += [
            "        ],",
            '        "double_bay_stops": [],',
            '        "triple_bay_stops": [],',
            '        "overflow_bays": [],',
            "    },",
        ]
    lines.append("}")
    return "\n".join(lines)


def step3_text(assignments: pd.DataFrame) -> str:
    """Step 3's CLUSTER_NAME, CLUSTER_STOPS and CLUSTER_CAPACITY, one block per zone."""
    lines = _banner(
        f"Step 3: {STEP3_SCRIPT}",
        "The sweep takes one zone per run: paste that zone's three settings. Bay",
        "labels start as the stop_id; rename them freely (no commas). Bays not",
        "listed in CLUSTER_CAPACITY hold one bus.",
    )
    blocks = []
    for zone, group in assignments.groupby("zone", sort=False):
        block = [
            f"# Zone: {zone}",
            f"CLUSTER_NAME = {_literal(str(zone))}",
            "CLUSTER_STOPS: Dict[str, str] = {  # stop_id -> bay label",
        ]
        block += [
            _with_comment(f"    {_literal(stop)}: {_literal(stop)},", _describe(stop, code, name))
            for stop, code, name in _stop_rows(group)
        ]
        block += ["}", "CLUSTER_CAPACITY: Dict[str, int] = {}  # bay label -> capacity; default 1"]
        blocks.append("\n".join(block))
    return "\n".join(lines) + "\n" + "\n\n".join(blocks)


def _banner(*text: str) -> List[str]:
    """Comment lines introducing one step's section."""
    rule = "# " + "-" * 76
    return [rule, *(f"# {line}" for line in text), rule]


def _stop_rows(group: pd.DataFrame) -> List[Tuple[str, str, str]]:
    """The (stop_id, stop_code, stop_name) of each stop in *group*, as text."""
    return list(
        zip(
            group["stop_id"].astype(str).tolist(),
            group["stop_code"].astype(str).tolist(),
            group["stop_name"].astype(str).tolist(),
        )
    )


def _stop_lines(group: pd.DataFrame, indent: str) -> List[str]:
    """One quoted stop_id per line, with the stop's code and name as a comment."""
    return [
        _with_comment(f"{indent}{_literal(stop)},", _describe(stop, code, name))
        for stop, code, name in _stop_rows(group)
    ]


def _literal(text: str) -> str:
    """Quote *text* as a double-quoted Python string literal."""
    return json.dumps(text, ensure_ascii=False)


def _describe(stop_id: str, stop_code: str, stop_name: str) -> str:
    """Comment text identifying a stop: its stop_code (unless it repeats the stop_id) and name."""
    code = "" if stop_code == stop_id else stop_code
    return " | ".join(part for part in (code, stop_name) if part)


def _with_comment(code: str, note: str) -> str:
    """Append ``  # note`` to *code*, shortening the note to keep within MAX_LINE_LENGTH."""
    room = MAX_LINE_LENGTH - len(code) - len("  # ")
    if not note or room < 4:
        return code
    if len(note) > room:
        note = note[: room - 3].rstrip() + "..."
    return f"{code}  # {note}"


def _fit(line: str) -> str:
    """Shorten *line* to MAX_LINE_LENGTH, marking the cut with "..."."""
    if len(line) <= MAX_LINE_LENGTH:
        return line
    return line[: MAX_LINE_LENGTH - 3].rstrip() + "..."


def _one_line(text: str) -> str:
    """Collapse whitespace, including line breaks, to single spaces."""
    return " ".join(text.split())


def _preview(values: Sequence[Any], limit: int = 10, separator: str = ", ") -> str:
    """List up to *limit* values, noting how many more there are."""
    shown = separator.join(str(value) for value in values[:limit])
    extra = len(values) - limit
    return f"{shown} and {extra} more" if extra > 0 else shown


# =============================================================================
# BAY ASSIGNMENT WORKBOOK
# =============================================================================


def _text_column(frame: pd.DataFrame, name: str) -> pd.Series:
    """Return a column as text, or aligned blanks for an absent optional column."""
    if name not in frame:
        return pd.Series("", index=frame.index, dtype=str)
    return frame[name].fillna("").astype(str)


def _require_columns(frame: pd.DataFrame, columns: Sequence[str], filename: str) -> None:
    """Reject a GTFS table missing columns needed for the workbook."""
    missing = [column for column in columns if column not in frame]
    if missing:
        raise ValueError(f"{filename} is missing column(s): {', '.join(missing)}.")


def _require_unique_ids(frame: pd.DataFrame, column: str, filename: str) -> None:
    """Reject blank or duplicate keys before a many-to-one GTFS join."""
    values = _text_column(frame, column)
    if values.str.strip().eq("").any():
        raise ValueError(f"{filename} contains blank {column} values.")
    repeated = values[values.duplicated()].unique().tolist()
    if repeated:
        raise ValueError(f"{filename} repeats {column}: {_preview(repeated)}.")


def build_bay_assignments(
    assignments: pd.DataFrame,
    stop_times: pd.DataFrame,
    trips: pd.DataFrame,
    routes: pd.DataFrame,
) -> pd.DataFrame:
    """Build unique scheduled stop/route/direction combinations for cluster stops.

    Args:
        assignments: Stops already assigned to zones by the spatial workflow.
        stop_times: GTFS stop_times.txt, read as strings.
        trips: GTFS trips.txt, read as strings.
        routes: GTFS routes.txt, read as strings.

    Returns:
        Columns stop_id, route_id, route_label and direction_id. Unknown
        directions use "?". Route labels prefer route_short_name, then
        route_long_name, then route_id. Duplicate names include route_id.

    Raises:
        ValueError: Required columns or keys are invalid, or a clustered stop
            references an unknown trip or route. No silent dropping of visits.

    Notes:
        All service_ids and all stop visits are included. A route-direction
        may legitimately appear at several stops. Pickup restrictions and
        calendars are not interpreted; this is not a boarding assignment guide.
    """
    _require_columns(stop_times, ("stop_id", "trip_id"), "stop_times.txt")
    _require_columns(trips, ("trip_id", "route_id"), "trips.txt")
    _require_columns(routes, ("route_id",), "routes.txt")
    _require_unique_ids(trips, "trip_id", "trips.txt")
    _require_unique_ids(routes, "route_id", "routes.txt")

    visits = pd.DataFrame(
        {
            "stop_id": _text_column(stop_times, "stop_id"),
            "trip_id": _text_column(stop_times, "trip_id"),
        }
    )
    visits = visits[visits["stop_id"].isin(assignments["stop_id"])].drop_duplicates()
    trip_table = pd.DataFrame(
        {
            "trip_id": _text_column(trips, "trip_id"),
            "route_id": _text_column(trips, "route_id"),
            "direction_id": _text_column(trips, "direction_id").str.strip(),
        }
    )
    joined = visits.merge(
        trip_table, on="trip_id", how="left", validate="many_to_one", indicator=True
    )
    missing_trips = joined.loc[joined["_merge"].eq("left_only"), "trip_id"].unique()
    if len(missing_trips):
        raise ValueError(
            "Cluster stop visits reference missing trips.txt trip_id values: "
            f"{_preview(missing_trips.tolist())}."
        )
    joined = joined.drop(columns="_merge")
    route_table = pd.DataFrame(
        {
            "route_id": _text_column(routes, "route_id"),
            "route_label": _text_column(routes, "route_short_name").map(_one_line),
        }
    )
    long_names = _text_column(routes, "route_long_name").map(_one_line)
    route_table["route_label"] = route_table["route_label"].where(
        route_table["route_label"].ne(""), long_names
    )
    route_table["route_label"] = route_table["route_label"].where(
        route_table["route_label"].ne(""), route_table["route_id"]
    )
    ambiguous = route_table["route_label"].duplicated(keep=False)
    route_table.loc[ambiguous, "route_label"] = (
        route_table.loc[ambiguous, "route_label"]
        + " [route_id="
        + route_table.loc[ambiguous, "route_id"]
        + "]"
    )
    joined = joined.merge(
        route_table, on="route_id", how="left", validate="many_to_one", indicator=True
    )
    missing_routes = joined.loc[joined["_merge"].eq("left_only"), "route_id"].unique()
    if len(missing_routes):
        raise ValueError(
            "Trips serving cluster stops reference missing routes.txt route_id values: "
            f"{_preview(missing_routes.tolist())}."
        )
    invalid = ~joined["direction_id"].isin(["", "0", "1"])
    if invalid.any():
        logging.warning(
            "Invalid direction_id values shown as (?): %s",
            _preview(joined.loc[invalid, "direction_id"].unique().tolist()),
        )
    joined.loc[~joined["direction_id"].isin(["0", "1"]), "direction_id"] = "?"
    columns = ["stop_id", "route_id", "route_label", "direction_id"]
    result = joined[columns].drop_duplicates(["stop_id", "route_id", "direction_id"])
    logging.info("Found %d unique stop/route/direction combination(s).", len(result))
    return result.reset_index(drop=True)


def _excel_text(value: str) -> str:
    """Remove control characters Excel cannot store from display text only."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def _worksheet_name(zone: str, used: set[str]) -> str:
    """Return a unique Excel-safe title; *used* contains casefolded titles."""
    base = re.sub(r"[\\/*?:\[\]]", "_", _excel_text(zone)).strip().strip("'")
    base = base or "Cluster"
    title = base[:31].rstrip("'")
    number = 2
    while title.casefold() in used:
        suffix = f"_{number}"
        title = base[: 31 - len(suffix)].rstrip("'") + suffix
        number += 1
    used.add(title.casefold())
    return title


def write_bay_assignment_workbook(
    output_file: Path,
    assignments: pd.DataFrame,
    bay_assignments: pd.DataFrame,
    zone_order: Sequence[str],
    gtfs_path: str,
) -> None:
    """Write an index and one formatted, printable stop-column sheet per cluster.

    Args:
        output_file: Destination .xlsx file.
        assignments: Cluster membership with stop names and IDs.
        bay_assignments: Unique scheduled uses from build_bay_assignments.
        zone_order: Full cluster order, including clusters with no stops.
        gtfs_path: Source feed path to record in the workbook.

    Raises:
        RuntimeError: openpyxl is unavailable in the Python environment.
        ValueError: A cluster exceeds Excel's column limit.
        OSError: The workbook cannot be saved, e.g. it is open in Excel.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "The Excel export requires openpyxl in this Python environment. "
            "Install openpyxl or set WRITE_BAY_ASSIGNMENT_XLSX = False."
        ) from exc

    workbook = Workbook()
    index = workbook.active
    index.title = "Index"
    used_names = {"index", "history"}
    navy = "17365D"
    pale_blue = "EAF1F8"

    def put(sheet: Any, row: int, column: int, value: Any, header: bool = False) -> Any:
        cell = sheet.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            cell.value = _excel_text(value)
            cell.data_type = "s"  # IDs/names starting with '=' remain literal text.
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(
            name="Calibri", size=11, bold=header, color="FFFFFF" if header else "202B38"
        )
        if header or row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=navy if header else pale_blue)
        return cell

    scope = "All services in feed; no date or day-of-week filter."
    notes = (
        "All scheduled stop visits, including arrivals and no-pickup visits; not a boarding "
        "guide. Entries are route (direction_id); (?) means missing/invalid direction. "
        "Columns are independent lists; rows do not indicate simultaneous use."
    )
    for row, value in enumerate(
        [
            "Scheduled bay use by cluster",
            scope,
            notes,
            f"GTFS: {gtfs_path}",
            f"Generated: {datetime_module.datetime.now().isoformat(timespec='seconds')}",
            "Stop headers use stop_id, not stop_code. No conflict or optimization analysis.",
        ],
        start=1,
    ):
        index.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        put(index, row, 1, value, header=(row == 1))
        index.row_dimensions[row].height = 45 if row in (3, 4) else 30
    for column, label in enumerate(["Cluster", "Worksheet", "Stops", "Route-directions"], 1):
        put(index, 8, column, label, header=True)
    index.column_dimensions["A"].width = 42
    index.column_dimensions["B"].width = 34
    index.column_dimensions["C"].width = 12
    index.column_dimensions["D"].width = 20
    index.freeze_panes = "A9"

    by_stop: dict[str, list[str]] = {}
    records = sorted(
        bay_assignments.to_dict("records"),
        key=lambda row: (_natural_key(row["route_label"]), row["direction_id"], row["route_id"]),
    )
    for record in records:
        by_stop.setdefault(record["stop_id"], []).append(
            f"{record['route_label']} ({record['direction_id']})"
        )

    for row_number, zone in enumerate(dict.fromkeys(zone_order), start=9):
        group = assignments[assignments["zone"].eq(zone)]
        stops = sorted(group.to_dict("records"), key=lambda row: _natural_key(row["stop_id"]))
        if len(stops) > 16384:
            raise ValueError(f"Cluster {zone!r} has more stops than Excel permits columns.")
        title = _worksheet_name(str(zone), used_names)
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        width = max(1, len(stops))
        for row, value in enumerate([str(zone), scope, notes], start=1):
            if width > 1:  # Excel reports a one-cell merge as corrupt content.
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            put(sheet, row, 1, value, header=(row == 1))
            sheet.row_dimensions[row].height = (120 if width == 1 else 72) if row == 3 else 42
        for column, stop in enumerate(stops, start=1):
            name = stop["stop_name"] or "Unnamed stop"
            put(sheet, 5, column, f"{name}\n({stop['stop_id']})", header=True)
            for row, label in enumerate(by_stop.get(stop["stop_id"], ["No scheduled trips"]), 6):
                put(sheet, row, column, label)
                sheet.row_dimensions[row].height = 32
            sheet.column_dimensions[get_column_letter(column)].width = 36
        if not stops:
            sheet.column_dimensions["A"].width = 52
            put(sheet, 5, 1, "No stops in this cluster", header=True)
        sheet.row_dimensions[5].height = 66
        sheet.freeze_panes = "A6"
        sheet.print_title_rows = "1:5"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:{get_column_letter(width)}{sheet.max_row}"
        cluster_uses = bay_assignments[bay_assignments["stop_id"].isin(group["stop_id"])]
        count = len(cluster_uses.drop_duplicates(["route_id", "direction_id"]))
        for column, value in enumerate([str(zone), title, len(stops), count], 1):
            cell = put(index, row_number, column, value)
            if column == 2:
                escaped = title.replace("'", "''")
                cell.hyperlink = f"#'{escaped}'!A1"
        index.row_dimensions[row_number].height = 32

    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_file: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_file.parent, suffix=".xlsx", delete=False
        ) as tmp:
            temporary_file = Path(tmp.name)
        workbook.save(temporary_file)
        os.replace(temporary_file, output_file)
    finally:
        workbook.close()
        if temporary_file is not None and temporary_file.exists():
            temporary_file.unlink()
    logging.info("Bay assignment workbook saved to %s", output_file)


# =============================================================================
# RUN LOG
# =============================================================================


class RunLogError(RuntimeError):
    """Raised when the required ``_runlog.txt`` sidecar could not be written."""


# Canonical version lives in utils/run_log.py -- keep this copy in sync.
def extract_config_block(source_file: Path) -> str:
    r"""Return the text between the CONFIG markers in *source_file*.

    Reads ``source_file`` as UTF-8 text and slices out the lines strictly
    *between* the first occurrence of ``# === BEGIN CONFIG ===`` and the first
    subsequent occurrence of ``# === END CONFIG ===``.  The marker lines
    themselves are excluded; whitespace and inline comments inside the block
    are preserved verbatim.

    Args:
        source_file: Path to the Python source file to scan (typically
            ``Path(__file__)`` from the calling script).

    Returns:
        The verbatim text of the configuration block, joined with ``\n``.

    Raises:
        ValueError: If either marker is missing or they appear out of order.
        OSError: If ``source_file`` cannot be read.
    """
    _BEGIN = "# === BEGIN CONFIG ==="
    _END = "# === END CONFIG ==="

    lines: list[str] = source_file.read_text(encoding="utf-8").splitlines()

    begin_idx: int | None = None
    end_idx: int | None = None
    for i, line in enumerate(lines):
        stripped: str = line.strip()
        if begin_idx is None and stripped == _BEGIN:
            begin_idx = i
        elif begin_idx is not None and stripped == _END:
            end_idx = i
            break

    if begin_idx is None or end_idx is None:
        raise ValueError(
            f"Config markers not found in '{source_file}'. Expected '{_BEGIN}' and '{_END}'."
        )

    return "\n".join(lines[begin_idx + 1 : end_idx])


def resolve_source_file() -> Optional[Path]:
    """Path of this script on disk, or ``None`` in an interactive session without ``__file__``."""
    try:
        return Path(__file__).resolve()
    except NameError:
        return None


def sha256_of_file(path: str) -> Optional[str]:
    """Return the SHA-256 hex digest of *path*, or None if it is not a readable file."""
    try:
        digest = hashlib.sha256()
        with open(path, "rb") as handle:
            for chunk in iter(lambda: handle.read(1 << 20), b""):
                digest.update(chunk)
        return digest.hexdigest()
    except OSError:
        return None


def input_fingerprints(
    gtfs_path: str, zones_path: str, gtfs_files: Sequence[str] = ("stops.txt",)
) -> List[str]:
    """Return SHA-256 fingerprint lines for the inputs, for the run log.

    A GTFS folder fingerprints every file read; a .zip feed the archive. A
    shapefile zone layer also fingerprints its .dbf (names) and .prj
    (coordinate system). A feature class in a geodatabase is not a single
    file and is not fingerprinted.

    Args:
        gtfs_path: The GTFS feed used for this run.
        zones_path: The zone layer used for this run.
        gtfs_files: Files read from the feed, including workbook inputs if enabled.

    Returns:
        "label: path" / "sha256: digest" line pairs.
    """
    targets: List[Tuple[str, str]] = []
    if os.path.isdir(gtfs_path):
        for filename in gtfs_files:
            targets.append((f"GTFS {filename}", os.path.join(gtfs_path, filename)))
    else:
        targets.append(("GTFS feed", gtfs_path))
    targets.append(("Zone layer", zones_path))
    stem, extension = os.path.splitext(zones_path)
    if extension.lower() == ".shp":
        for sidecar in (".dbf", ".prj"):
            for candidate in (stem + sidecar, stem + sidecar.upper()):
                if os.path.isfile(candidate):
                    targets.append((f"Zone layer {sidecar}", candidate))
                    break
    lines: List[str] = []
    for label, path in targets:
        digest = sha256_of_file(path)
        lines += [f"{label}: {path}", f"  sha256: {digest or 'not fingerprinted (not a file)'}"]
    return lines


def write_run_log(
    output_file: Path, effective_settings: Sequence[str], fingerprints: Sequence[str]
) -> bool:
    """Write the ``_runlog.txt`` sidecar for *output_file* (same folder, same stem).

    The log captures this script's CONFIGURATION block verbatim, the settings
    actually used (which CLI flags may have changed), and input fingerprints.

    Args:
        output_file: The config text file or workbook this run wrote.
        effective_settings: Pre-formatted lines of the resolved settings.
        fingerprints: Lines from :func:`input_fingerprints`.

    Returns:
        ``True`` if the log was written, ``False`` otherwise.
    """
    log_path = output_file.with_name(f"{output_file.stem}_runlog.txt")
    source_file = resolve_source_file()
    if source_file is None:
        config_text = "(config block unavailable: interactive session, no __file__ on disk)"
        source_display = "<interactive>"
    else:
        try:
            config_text = extract_config_block(source_file)
        except (OSError, ValueError) as exc:
            logging.error("Could not extract config block for run log: %s", exc)
            return False
        source_display = str(source_file)

    lines: List[str] = [
        "=" * 72,
        "CLUSTER STOPS FROM ZONES RUN LOG",
        "=" * 72,
        f"Run timestamp:    {datetime_module.datetime.now().isoformat(timespec='seconds')}",
        f"Output file:      {output_file}",
        f"Source script:    {source_display}",
        "",
        "-" * 72,
        "EFFECTIVE SETTINGS (constants or CLI flags, as resolved for this run)",
        "-" * 72,
        *effective_settings,
        "",
        "-" * 72,
        "INPUT FINGERPRINTS",
        "-" * 72,
        *fingerprints,
        "",
        "-" * 72,
        "CONFIGURATION (verbatim from source)",
        "-" * 72,
        config_text,
        "=" * 72,
    ]
    try:
        log_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    except OSError as exc:
        logging.error("Error writing run log: %s", exc)
        return False
    logging.info("Run log saved to %s", log_path)
    return True


def require_run_log(written: bool) -> None:
    """Raise :class:`RunLogError` when a run log failed and ``REQUIRE_RUN_LOG`` is set."""
    if not written and REQUIRE_RUN_LOG:
        raise RunLogError(
            "Run log could not be written. Set REQUIRE_RUN_LOG = False to suppress this "
            "error when a sidecar file is genuinely impossible."
        )


# =============================================================================
# PIPELINE
# =============================================================================


def resolve_output_file(output_dir: str, output_filename: str) -> Optional[Path]:
    """Return where to save the config text, or ``None`` to write it to the log only.

    Raises:
        ValueError: *output_filename* is blank or includes a folder.
    """
    if not output_dir:
        return None
    if (
        not output_filename
        or Path(output_filename).name != output_filename
        or "\\" in output_filename
    ):
        raise ValueError(
            f"OUTPUT_FILENAME must be a file name without folders, got {output_filename!r}."
        )
    return Path(output_dir) / output_filename


def check_near_miss_feet(near_miss_feet: float) -> None:
    """Raise ValueError unless *near_miss_feet* is a finite distance of 0 or more."""
    if not math.isfinite(near_miss_feet) or near_miss_feet < 0:
        raise ValueError(
            f"NEAR_MISS_FEET must be 0 (off) or a positive number of feet, got {near_miss_feet}."
        )


def run(args: argparse.Namespace) -> str:
    """Build config text and optionally export scheduled bay use by cluster.

    Args:
        args: Parsed CLI arguments (defaults mirror the CONFIGURATION block).

    Returns:
        The config text.

    Raises:
        OSError: An input is missing or unreadable.
        ValueError: An input or setting is invalid, zones overlap, or a stop
            lies on the edge of two zones (the message says which and how to
            fix it).
        RunLogError: The run log could not be written and REQUIRE_RUN_LOG
            is True.
    """
    check_near_miss_feet(args.near_miss_feet)
    output_file = resolve_output_file(args.output_dir, args.output_filename)
    workbook_file = None
    gtfs_files = ["stops.txt"]
    if args.write_bay_assignment_xlsx:
        if output_file is None:
            raise ValueError(
                "Set OUTPUT_DIR / --output-dir to save the workbook, or disable it with "
                "WRITE_BAY_ASSIGNMENT_XLSX = False / --no-bay-assignment-xlsx."
            )
        workbook_file = resolve_output_file(args.output_dir, args.bay_assignment_filename)
        if workbook_file is None or workbook_file.suffix.lower() != ".xlsx":
            raise ValueError("BAY_ASSIGNMENT_FILENAME must be a filename ending in .xlsx.")
        output_names = [
            output_file.name,
            workbook_file.name,
            f"{output_file.stem}_runlog.txt",
            f"{workbook_file.stem}_runlog.txt",
        ]
        if len({name.casefold() for name in output_names}) != len(output_names):
            raise ValueError("The config text, workbook and their run logs must have unique names.")
        gtfs_files += ["stop_times.txt", "trips.txt", "routes.txt"]
    zones = load_zones(args.zones_path, args.zone_name_field)
    check_zone_overlaps(zones)
    gtfs = load_gtfs_data(args.gtfs_path, files=gtfs_files)
    stops = prepare_stops(gtfs["stops"])
    assignments = assign_stops_to_zones(stops, zones, args.geo_transformation)
    near_misses = find_near_misses(
        stops, zones, assignments, args.near_miss_feet, args.geo_transformation
    )
    zone_order = list(dict.fromkeys(zones.names))
    bay_assignments = None
    if workbook_file is not None:
        bay_assignments = build_bay_assignments(
            assignments, gtfs["stop_times"], gtfs["trips"], gtfs["routes"]
        )
    text = build_config_text(
        assignments,
        zone_order,
        args.zones_path,
        args.gtfs_path,
        near_misses,
        args.near_miss_feet,
    )
    logging.info("Config text to paste:\n\n%s", text)
    if output_file is None:
        return text
    if workbook_file is not None and bay_assignments is not None:
        write_bay_assignment_workbook(
            workbook_file, assignments, bay_assignments, zone_order, args.gtfs_path
        )
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(text, encoding="utf-8")
    logging.info("Config text saved to %s", output_file)
    settings = [
        f"GTFS feed:        {args.gtfs_path}",
        f"Zone layer:       {args.zones_path}",
        f"Zone name field:  {args.zone_name_field or '(none: zones numbered by position)'}",
        f"Near-miss feet:   {args.near_miss_feet:,g}",
        f"Transformation:   {args.geo_transformation or '(none)'}",
        f"Output file:      {output_file}",
        f"Bay workbook:     {workbook_file or '(disabled)'}",
        "Workbook scope:   All services; all scheduled stop visits (including no-pickup).",
    ]
    fingerprints = input_fingerprints(args.gtfs_path, args.zones_path, gtfs_files)
    require_run_log(write_run_log(output_file, settings, fingerprints))
    if workbook_file is not None:
        require_run_log(write_run_log(workbook_file, settings, fingerprints))
    return text


# =============================================================================
# CLI / MAIN
# =============================================================================


# Canonical version lives in utils/cli_helpers.py -- keep this copy in sync.
def notebook_safe_argv(argv: Optional[Sequence[str]]) -> Optional[List[str]]:
    """Return the argv to parse, shielding notebook kernels from stray flags.

    When a script's ``main()`` runs with no explicit ``argv`` inside a
    Jupyter/IPython kernel, ``sys.argv`` holds kernel plumbing (for example
    ``-f /path/kernel.json``) rather than flags meant for the script, and
    strict ``argparse.parse_args`` would reject it and abort.  This helper
    detects the notebook case and substitutes an empty argument list so the
    CONFIGURATION constants stay in charge, while shell runs keep strict
    parsing (a typo in a flag fails loudly instead of being silently ignored).

    Canonical implementation: ``utils/cli_helpers.py``.

    Args:
        argv: Explicit argument list passed to ``main()``, or ``None`` to
            fall back to ``sys.argv``.

    Returns:
        ``list(argv)`` when *argv* was provided; ``[]`` when running inside a
        notebook kernel; otherwise ``None`` so argparse reads ``sys.argv[1:]``.
    """
    if argv is not None:
        return list(argv)
    if "ipykernel" in sys.modules:
        return []
    return None


def build_arg_parser() -> argparse.ArgumentParser:
    """Create the command-line argument parser (defaults mirror CONFIGURATION)."""
    parser = argparse.ArgumentParser(
        description="List the GTFS stops inside zone polygons as config text for the bay scripts.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "--gtfs-path",
        default=GTFS_FOLDER_PATH,
        help="GTFS feed folder or .zip (the feed Step 1 reads).",
    )
    parser.add_argument(
        "--zones-path",
        default=ZONES_PATH,
        help="Zone polygon layer (shapefile, geodatabase feature class, ...).",
    )
    parser.add_argument(
        "--zone-name-field",
        default=ZONE_NAME_FIELD,
        help='Attribute naming each zone; "" numbers zones by position in the layer.',
    )
    parser.add_argument(
        "--near-miss-feet",
        type=float,
        default=NEAR_MISS_FEET,
        help="Report stops outside every zone but within this many feet of one; 0 turns it off.",
    )
    parser.add_argument(
        "--geo-transformation",
        default=GEO_TRANSFORMATION,
        help='Geographic transformation from WGS84 to the zone layer datum; "" applies none.',
    )
    parser.add_argument(
        "--output-dir",
        default=OUTPUT_DIR,
        help='Folder for output files; "" permits log-only output when the workbook is disabled.',
    )
    parser.add_argument(
        "--output-filename",
        default=OUTPUT_FILENAME,
        help="File name for the copy of the config text.",
    )
    parser.add_argument(
        "--bay-assignment-xlsx",
        dest="write_bay_assignment_xlsx",
        action=argparse.BooleanOptionalAction,
        default=WRITE_BAY_ASSIGNMENT_XLSX,
        help="Export an Excel inventory of scheduled stop use, one sheet per cluster.",
    )
    parser.add_argument(
        "--bay-assignment-filename",
        default=BAY_ASSIGNMENT_FILENAME,
        help="File name for the optional Excel workbook (must end in .xlsx).",
    )
    return parser


_PLACEHOLDER_MARKERS: Tuple[str, ...] = (
    "path\\to\\your",
    "your\\file\\path",
    "your\\folder\\path",
    "path/to/your",
    "your/file/path",
    "your/folder/path",
)


def _is_placeholder_path(p: str) -> bool:
    """Return True if *p* still points at a default placeholder location."""
    s = str(p).lower()
    return any(marker in s for marker in _PLACEHOLDER_MARKERS)


def main(argv: Optional[Sequence[str]] = None) -> int:
    """Entry point. Validates placeholder paths before doing any work.

    Args:
        argv: Optional explicit argument list; None reads ``sys.argv``.

    Returns:
        Process exit code: 0 on success, 1 on runtime failure, 2 if required
        CONFIGURATION values are still placeholders.
    """
    logging.basicConfig(
        level=LOG_LEVEL,
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    args = build_arg_parser().parse_args(notebook_safe_argv(argv))
    placeholders = [
        name
        for name, value in (
            ("GTFS_FOLDER_PATH / --gtfs-path", args.gtfs_path),
            ("ZONES_PATH / --zones-path", args.zones_path),
            ("OUTPUT_DIR / --output-dir", args.output_dir),
        )
        if _is_placeholder_path(value)
    ]
    if placeholders:
        logging.warning(
            "Placeholder value(s) still set for: %s. Update the CONFIGURATION section or "
            "pass the matching CLI flags before running.",
            "; ".join(placeholders),
        )
        return 2
    try:
        run(args)
    except (OSError, ValueError, RuntimeError, arcpy.ExecuteError) as exc:
        logging.error("%s", exc)
        return 1
    logging.info("Script completed successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
