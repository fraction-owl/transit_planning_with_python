from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import Any, Optional
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

import scripts.facilities_tools.bay_usage_analyzer as target

# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

TEST_CLUSTERS = {
    "Park & Ride": {
        "single_bay_stops": ["100", "101"],
        "double_bay_stops": [],
        "triple_bay_stops": [],
        "overflow_bays": [],
    },
    "Metro": {
        "single_bay_stops": ["200"],
        "double_bay_stops": ["201"],
        "triple_bay_stops": ["202"],
        "overflow_bays": ["layover_A"],
    },
}


def _make_events_df() -> pd.DataFrame:
    """Three buses present at cluster 'Park & Ride' (capacity 2) at 08:00."""
    return pd.DataFrame(
        {
            "Stop ID": ["100", "100", "101", "200", "999"],
            "Status": ["ARRIVE", "LOADING", "DWELL", "ARRIVE", "ARRIVE"],
            "Timestamp": ["08:00", "08:00", "08:00", "08:00", "08:00"],
        }
    )


# ---------------------------------------------------------------------------
# get_all_official_stops
# ---------------------------------------------------------------------------


def test_get_all_official_stops_combines_bay_lists() -> None:
    cinfo = TEST_CLUSTERS["Metro"]
    assert target.get_all_official_stops(cinfo) == ["200", "201", "202"]


def test_get_all_official_stops_excludes_overflow() -> None:
    cinfo = TEST_CLUSTERS["Metro"]
    assert "layover_A" not in target.get_all_official_stops(cinfo)


def test_get_all_official_stops_missing_keys_returns_empty() -> None:
    assert target.get_all_official_stops({}) == []


# ---------------------------------------------------------------------------
# build_cluster_capacities / build_stop_capacities
# ---------------------------------------------------------------------------


def test_build_cluster_capacities_weights_bay_types() -> None:
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        caps = target.build_cluster_capacities()
    assert caps["Park & Ride"] == 2  # 2 singles
    assert caps["Metro"] == 7  # 1 + 2 + 3 + 1 overflow


def test_build_stop_capacities_per_stop_values() -> None:
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        caps = target.build_stop_capacities()
    assert caps["100"] == 1
    assert caps["201"] == 2
    assert caps["202"] == 3
    assert caps["layover_A"] == 1


# ---------------------------------------------------------------------------
# normalize_stop_id
# ---------------------------------------------------------------------------


def test_normalize_stop_id_strips_float_suffix() -> None:
    assert target.normalize_stop_id("2956.0") == "2956"
    assert target.normalize_stop_id(2956.0) == "2956"


def test_normalize_stop_id_nan_returns_none() -> None:
    assert target.normalize_stop_id(float("nan")) is None


def test_normalize_stop_id_strips_whitespace() -> None:
    assert target.normalize_stop_id(" 42 ") == "42"


def test_normalize_stop_id_plain_string_unchanged() -> None:
    assert target.normalize_stop_id("layover_A") == "layover_A"


# ---------------------------------------------------------------------------
# assign_cluster_name
# ---------------------------------------------------------------------------


def test_assign_cluster_name_maps_official_and_overflow_stops() -> None:
    df = pd.DataFrame({"Stop ID": ["100", "201", "layover_A", "999"]})
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        out = target.assign_cluster_name(df)
    assert out.loc[0, "ClusterName"] == "Park & Ride"
    assert out.loc[1, "ClusterName"] == "Metro"
    assert out.loc[2, "ClusterName"] == "Metro"
    assert out.loc[3, "ClusterName"] is None


def test_assign_cluster_name_does_not_mutate_input() -> None:
    df = pd.DataFrame({"Stop ID": ["100"]})
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        target.assign_cluster_name(df)
    assert "ClusterName" not in df.columns


# ---------------------------------------------------------------------------
# find_cluster_conflicts / find_stop_conflicts
# ---------------------------------------------------------------------------


def test_find_cluster_conflicts_flags_over_capacity_timestamp() -> None:
    df = _make_events_df()
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        df = target.assign_cluster_name(df)
        conflicts = target.find_cluster_conflicts(df)
    assert ("Park & Ride", "08:00") in conflicts
    assert ("Metro", "08:00") not in conflicts  # 1 bus, capacity 7


def test_find_cluster_conflicts_ignores_non_presence_statuses() -> None:
    df = pd.DataFrame(
        {
            "Stop ID": ["100", "100", "101"],
            "Status": ["TRAVELING", "TRAVELING", "TRAVELING"],
            "Timestamp": ["08:00"] * 3,
        }
    )
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        df = target.assign_cluster_name(df)
        assert target.find_cluster_conflicts(df) == set()


def test_find_stop_conflicts_flags_single_bay_double_service() -> None:
    df = _make_events_df()
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        conflicts = target.find_stop_conflicts(df)
    # Stop 100 has two passenger-service buses at 08:00 but capacity 1.
    assert ("100", "08:00") in conflicts
    assert ("101", "08:00") not in conflicts


def test_find_stop_conflicts_overflow_statuses_never_occupy_bay() -> None:
    df = pd.DataFrame(
        {
            "Stop ID": ["100", "100", "100"],
            "Status": ["LAYOVER", "LONG BREAK", "ARRIVE"],
            "Timestamp": ["08:00", "08:00", "08:00"],
        }
    )
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        assert target.find_stop_conflicts(df) == set()


def test_find_stop_conflicts_in_bay_layover_counts_only_when_enabled() -> None:
    df = pd.DataFrame(
        {
            "Stop ID": ["100", "100"],
            "Status": ["DWELL", "ARRIVE"],
            "Timestamp": ["08:00", "08:00"],
        }
    )
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        with patch.object(target, "COUNT_IN_BAY_LAYOVER_AT_STOP", True):
            assert ("100", "08:00") in target.find_stop_conflicts(df)
        with patch.object(target, "COUNT_IN_BAY_LAYOVER_AT_STOP", False):
            assert target.find_stop_conflicts(df) == set()


def test_find_stop_conflicts_ignores_stops_outside_clusters() -> None:
    # Blocks that touch a cluster spend the rest of their day at stops with no
    # defined capacity; those are not assessed (two buses sharing a busy street
    # stop is not a bay conflict).
    df = pd.DataFrame(
        {
            "Stop ID": ["999", "999"],
            "Status": ["ARRIVE", "DEPART"],
            "Timestamp": ["09:00", "09:00"],
        }
    )
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        assert target.find_stop_conflicts(df) == set()


# ---------------------------------------------------------------------------
# annotate_conflicts
# ---------------------------------------------------------------------------


def test_annotate_conflicts_categorises_each_row() -> None:
    df = pd.DataFrame(
        {
            "ClusterName": ["Park & Ride", "Park & Ride", None, "Metro"],
            "Stop ID": ["100", "101", "999", "200"],
            "Timestamp": ["08:00", "08:00", "08:00", "08:00"],
            "Status": ["ARRIVE", "LOADING", "ARRIVE", "DEPART"],
        }
    )
    cluster_conflicts = {("Park & Ride", "08:00")}
    stop_conflicts = {("100", "08:00"), ("999", "08:00")}
    out = target.annotate_conflicts(df, cluster_conflicts, stop_conflicts)
    assert list(out["ConflictType"]) == ["BOTH", "CLUSTER", "STOP", "NONE"]


def test_annotate_conflicts_overflow_rows_are_not_stop_conflicts() -> None:
    # A bus on the overflow space keeps its arrival stop ID but is not in that bay.
    df = pd.DataFrame(
        {
            "ClusterName": ["Park & Ride", "Park & Ride"],
            "Stop ID": ["100", "100"],
            "Timestamp": ["08:00", "08:00"],
            "Status": ["ARRIVE", "LAYOVER"],
        }
    )
    out = target.annotate_conflicts(df, {("Park & Ride", "08:00")}, {("100", "08:00")})
    assert list(out["ConflictType"]) == ["BOTH", "CLUSTER"]


def test_annotate_conflicts_no_conflicts_all_none() -> None:
    df = pd.DataFrame(
        {
            "ClusterName": ["Metro"],
            "Stop ID": ["200"],
            "Timestamp": ["08:00"],
            "Status": ["ARRIVE"],
        }
    )
    out = target.annotate_conflicts(df, set(), set())
    assert list(out["ConflictType"]) == ["NONE"]


# ---------------------------------------------------------------------------
# gather_block_spreadsheets
# ---------------------------------------------------------------------------


def _step1_rows(block: str, ts: str, **columns: Any) -> pd.DataFrame:
    """One Step 1 timeline row with the columns the readers require."""
    row = {"Timestamp": ts, "Block": block, "Trip ID": "T1", "Stop ID": "S1", "Status": "ARRIVE"}
    return pd.DataFrame([{**row, **columns}])


def test_gather_block_spreadsheets_concatenates_block_files(tmp_path: Path) -> None:
    df1 = _step1_rows("B1", "08:00")
    df2 = _step1_rows("B2", "09:00")
    df1.to_excel(tmp_path / "block_101.xlsx", index=False)
    df2.to_excel(tmp_path / "block_102.xlsx", index=False)
    # A non-block file that must be ignored.
    df1.to_excel(tmp_path / "summary.xlsx", index=False)

    combined = target.gather_block_spreadsheets(str(tmp_path))
    assert len(combined) == 2
    assert set(combined["FileName"]) == {"block_101.xlsx", "block_102.xlsx"}


def test_gather_block_spreadsheets_no_files_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        target.gather_block_spreadsheets(str(tmp_path))


def test_gather_block_spreadsheets_keeps_ids_as_text(tmp_path: Path) -> None:
    _step1_rows(101, "08:00", **{"Stop ID": 2956}).to_excel(tmp_path / "block_1.xlsx", index=False)
    out = target.gather_block_spreadsheets(str(tmp_path))
    assert (out.loc[0, "Stop ID"], out.loc[0, "Block"]) == ("2956", "101")


# ---------------------------------------------------------------------------
# load_timeline / read_assumptions / run_input_folder
# ---------------------------------------------------------------------------


def test_load_timeline_prefers_combined_csv(tmp_path: Path) -> None:
    _step1_rows("B1", "08:00", **{"Stop ID": ""}).to_csv(
        tmp_path / "all_blocks_timeline.csv", index=False
    )
    _step1_rows("B9", "09:00").to_excel(tmp_path / "block_9.xlsx", index=False)
    with patch.object(target, "COMBINED_TIMELINE_FILE", "all_blocks_timeline.csv"):
        out = target.load_timeline(str(tmp_path))
    assert list(out["Block"]) == ["B1"]
    assert out.loc[0, "Stop ID"] is None  # blank cells become None, not "" or NaN
    with patch.object(target, "COMBINED_TIMELINE_FILE", ""):
        out = target.load_timeline(str(tmp_path))
    assert list(out["Block"]) == ["B9"]


def test_load_timeline_falls_back_to_workbooks_when_csv_missing(tmp_path: Path) -> None:
    _step1_rows("B9", "09:00").to_excel(tmp_path / "block_9.xlsx", index=False)
    with patch.object(target, "COMBINED_TIMELINE_FILE", "all_blocks_timeline.csv"):
        out = target.load_timeline(str(tmp_path))
    assert list(out["Block"]) == ["B9"]


def test_read_assumptions_returns_non_blank_lines(tmp_path: Path) -> None:
    (tmp_path / "timeline_assumptions.txt").write_text("A=1\n\nB=2\n", encoding="utf-8")
    with patch.object(target, "ASSUMPTIONS_FILE", "timeline_assumptions.txt"):
        assert target.read_assumptions(str(tmp_path)) == ["A=1", "B=2"]
        assert target.read_assumptions(str(tmp_path / "missing")) == []
    with patch.object(target, "ASSUMPTIONS_FILE", ""):
        assert target.read_assumptions(str(tmp_path)) == []


def test_run_input_folder_appends_scenario() -> None:
    with patch.object(target, "BLOCK_OUTPUT_FOLDER", "in"):
        with patch.object(target, "SCENARIO_NAME", "alt"):
            assert target.run_input_folder() == os.path.join("in", "alt")
        with patch.object(target, "SCENARIO_NAME", ""):
            assert target.run_input_folder() == "in"


# ---------------------------------------------------------------------------
# _is_placeholder_path
# ---------------------------------------------------------------------------


def test_is_placeholder_path_detects_defaults() -> None:
    assert target._is_placeholder_path(r"Path\To\Your\Input_Folder") is True
    assert target._is_placeholder_path("path/to/your/output") is True


def test_is_placeholder_path_real_path_is_false() -> None:
    assert target._is_placeholder_path("/home/user/data/blocks") is False


# ---------------------------------------------------------------------------
# bay_statuses / timestamp_to_minutes / bus_label
# ---------------------------------------------------------------------------


def test_in_bay_layover_toggle_excludes_only_between_trip_dwell() -> None:
    df = pd.DataFrame(
        {
            "Status": ["DWELL", "DWELL", "ARRIVE", "LAYOVER"],
            "Prev Trip ID": ["T1", "", "", "T1"],  # row 0: between trips; row 1: a timed hold
            "Next Trip ID": ["T2", "", "", "T2"],
        }
    )
    assert "DWELL" in target.bay_statuses()  # a timed hold can always occupy a bay
    with patch.object(target, "COUNT_IN_BAY_LAYOVER_AT_STOP", True):
        assert list(target.bay_occupancy_mask(df)) == [True, True, True, False]
    with patch.object(target, "COUNT_IN_BAY_LAYOVER_AT_STOP", False):
        assert list(target.bay_occupancy_mask(df)) == [False, True, True, False]
        # Without trip links (older Step 1 output) every DWELL row counts as a layover.
        legacy = df.drop(columns=["Prev Trip ID", "Next Trip ID"])
        assert list(target.bay_occupancy_mask(legacy)) == [False, False, True, False]


@pytest.mark.parametrize(
    ("ts", "expected"),
    [("08:05", 485), ("25:10", 1510), (" 7:30 ", 450), ("abc", None), ("", None), (None, None)],
)
def test_timestamp_to_minutes(ts: Optional[str], expected: Optional[int]) -> None:
    assert target.timestamp_to_minutes(ts) == expected


def test_bus_label_prefers_short_name_and_headsign() -> None:
    row = pd.Series(
        {
            "Route Short Name": "101",
            "Route": "R101",
            "Trip Headsign": "Downtown",
            "Block": "B1",
            "Trip ID": "T1",
        }
    )
    assert target.bus_label(row) == "101 Downtown (blk B1, trip T1)"
    row = pd.Series(
        {
            "Route Short Name": None,
            "Route": "R101",
            "Trip Headsign": None,
            "Block": "B1",
            "Trip ID": "T1",
        }
    )
    assert target.bus_label(row) == "R101 (blk B1, trip T1)"


def test_bus_label_falls_back_to_long_name_then_route_id() -> None:
    row = pd.Series(
        {
            "Route Short Name": None,
            "Route Long Name": "Crosstown",
            "Route": "R7",
            "Trip Headsign": None,
            "Block": "B1",
            "Trip ID": "T1",
        }
    )
    assert target.bus_label(row) == "Crosstown (blk B1, trip T1)"
    assert target.route_display_name("", " ", "R7") == "R7"
    assert target.route_display_name(float("nan"), "Crosstown", "R7") == "Crosstown"


# ---------------------------------------------------------------------------
# Conflict events, tallies, minute rules, overflow events
# ---------------------------------------------------------------------------


def _timeline_row(
    block: str,
    ts: str,
    stop: Optional[str],
    status: str,
    trip: str = "T1",
    route: str = "101",
    headsign: str = "Downtown",
) -> dict[str, Any]:
    return {
        "Timestamp": ts,
        "Block": block,
        "Route": route,
        "Route Short Name": route,
        "Direction": "0",
        "Trip Headsign": headsign,
        "Trip ID": trip,
        "Stop ID": stop,
        "Stop Name": f"Stop {stop}" if stop else "",
        "Stop Sequence": "1",
        "Arrival Time": ts,
        "Departure Time": ts,
        "Status": status,
        "Layover Location": "",
        "Prev Trip ID": "",
        "Next Trip ID": "",
    }


def _conflict_df() -> pd.DataFrame:
    """Stop 100 (capacity 1): two buses at 08:00-08:01, one at 08:02, two again at 08:05."""
    return pd.DataFrame(
        [
            _timeline_row("B1", "08:00", "100", "ARRIVE"),
            _timeline_row("B2", "08:00", "100", "LOADING", trip="T2", route="202"),
            _timeline_row("B1", "08:01", "100", "ARRIVE"),
            _timeline_row("B2", "08:01", "100", "LOADING", trip="T2", route="202"),
            _timeline_row("B1", "08:02", "100", "ARRIVE"),
            _timeline_row("B1", "08:05", "100", "DEPART"),
            _timeline_row("B3", "08:05", "100", "ARRIVE/DEPART", trip="T3", route="303"),
            _timeline_row("B4", "08:05", "101", "ARRIVE", trip="T4", route="404"),
        ]
    )


def test_build_conflict_events_collapses_consecutive_minutes() -> None:
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        events = target.build_conflict_events(_conflict_df())
    assert len(events) == 2
    first = events.iloc[0]
    assert (first["Stop ID"], first["Start"], first["End"], first["Minutes"]) == (
        "100",
        "08:00",
        "08:01",
        2,
    )
    assert (first["Peak Buses"], first["Buses"]) == (2, 2)
    assert first["Bus 1"] == "101 Downtown (blk B1, trip T1)"
    assert first["Bus 2"] == "202 Downtown (blk B2, trip T2)"
    second = events.iloc[1]
    assert (second["Start"], second["End"], second["Minutes"]) == ("08:05", "08:05", 1)
    assert second["Bus 2"] == "303 Downtown (blk B3, trip T3)"


def test_build_conflict_events_flags_events_with_interpolated_times() -> None:
    df = _conflict_df().assign(**{"Estimated Time": "False"})
    df.loc[6, "Estimated Time"] = "True"  # B3's through visit at 08:05 was interpolated
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        events = target.build_conflict_events(df)
        legacy = target.build_conflict_events(_conflict_df())  # no Estimated Time column
    assert list(events["Estimated Times"]) == ["", "YES"]
    assert list(legacy["Estimated Times"]) == ["", ""]


def test_build_conflict_events_empty_when_no_conflicts() -> None:
    df = pd.DataFrame([_timeline_row("B1", "08:00", "100", "ARRIVE")])
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        events = target.build_conflict_events(df)
    assert events.empty
    assert list(events.columns)[:3] == ["Stop ID", "Stop Name", "Start"]


def test_build_bay_tally_counts_minutes_and_events() -> None:
    df = _conflict_df()
    with patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS):
        events = target.build_conflict_events(df)
        tally = target.build_bay_tally(df, events).set_index("Stop ID")
    assert (tally.loc["100", "Capacity"], tally.loc["100", "Conflict Minutes"]) == (1, 3)
    assert tally.loc["100", "Conflict Events"] == 2
    assert tally.loc["101", "Conflict Minutes"] == 0


def _metro_df() -> pd.DataFrame:
    """Metro: stop 200 (cap 1) over by one at 08:00-08:01, stop 201 (cap 2) over by one at 08:01.

    One bus is already on the overflow space at 08:00.
    """
    return pd.DataFrame(
        [
            _timeline_row("B1", "08:00", "200", "ARRIVE"),
            _timeline_row("B2", "08:00", "200", "LOADING", trip="T2"),
            _timeline_row("B1", "08:01", "200", "ARRIVE"),
            _timeline_row("B2", "08:01", "200", "LOADING", trip="T2"),
            _timeline_row("B3", "08:01", "201", "ARRIVE", trip="T3"),
            _timeline_row("B4", "08:01", "201", "ARRIVE", trip="T4"),
            _timeline_row("B5", "08:01", "201", "ARRIVE", trip="T5"),
            _timeline_row("B6", "08:00", "200", "LAYOVER", trip="T6"),
        ]
    )


def test_build_minute_rules_pooled_overflow() -> None:
    with (
        patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS),
        patch.object(target, "OVERFLOW_ROUTING", {}),
        patch.object(target, "BUMPED_TO_SPACE", ""),
    ):
        rules = target.build_minute_rules(_metro_df(), TEST_CLUSTERS["Metro"])
    assert rules["any_stop"] == 2
    assert rules["two_or_more_stops"] == 1
    # 08:00: one bumped bus plus one already on the single overflow bay; 08:01: two bumped.
    assert rules["overflow_adjusted"] == 2
    assert rules["overflow_capacity"] == 1
    assert (rules["peak_on_overflow"], rules["minutes_overflow_over_capacity"]) == (1, 0)
    assert rules["bumped_to"] == "(pool)"
    assert rules["per_space"] == {
        "(pool)": {"capacity": 1, "peak": 1, "minutes_over": 0, "minutes_used": 1}
    }


def test_build_minute_rules_routed_bump_space() -> None:
    with (
        patch.object(target, "CLUSTER_DEFINITIONS", TEST_CLUSTERS),
        patch.object(target, "OVERFLOW_ROUTING", {"layover_A": ["LAYOVER"]}),
        patch.object(target, "BUMPED_TO_SPACE", "layover_A"),
    ):
        rules = target.build_minute_rules(_metro_df(), TEST_CLUSTERS["Metro"])
    assert rules["bumped_to"] == "layover_A"
    assert rules["overflow_adjusted"] == 2
    assert list(rules["per_space"]) == ["layover_A"]


def test_cluster_routing_only_applies_to_own_overflow_bays() -> None:
    routing = {"layover_A": ["LAYOVER"], "elsewhere": ["LONG BREAK"]}
    with patch.object(target, "OVERFLOW_ROUTING", routing):
        assert target._cluster_routing(TEST_CLUSTERS["Metro"]) == {"LAYOVER": "layover_A"}
        assert target._cluster_routing(TEST_CLUSTERS["Park & Ride"]) == {}


def test_overflow_rows_routing_and_pool() -> None:
    df = pd.DataFrame(
        [
            _timeline_row("B6", "08:00", "200", "LAYOVER"),
            _timeline_row("B7", "08:00", "200", "LONG BREAK", trip="T7"),
            _timeline_row("B8", "08:00", "200", "ARRIVE", trip="T8"),
        ]
    )
    with patch.object(target, "OVERFLOW_ROUTING", {"layover_A": ["LAYOVER"]}):
        out = target._overflow_rows(df, TEST_CLUSTERS["Metro"])
    # Only overflow statuses; an unlisted status falls to the first overflow bay.
    assert list(out["Block"]) == ["B6", "B7"]
    assert list(out["Space"]) == ["layover_A", "layover_A"]
    with patch.object(target, "OVERFLOW_ROUTING", {}):
        out = target._overflow_rows(df, TEST_CLUSTERS["Metro"])
    assert set(out["Space"]) == {"(pool)"}


def test_build_overflow_events_flags_space_over_capacity() -> None:
    df = pd.DataFrame(
        [
            _timeline_row("B6", "08:00", "200", "LAYOVER"),
            _timeline_row("B7", "08:00", "200", "LAYOVER", trip="T7"),
            _timeline_row("B6", "08:01", "200", "LAYOVER"),
            _timeline_row("B7", "08:01", "200", "LAYOVER", trip="T7"),
            _timeline_row("B6", "08:02", "200", "LAYOVER"),
        ]
    )
    with patch.object(target, "OVERFLOW_ROUTING", {}):
        events = target.build_overflow_events(df, TEST_CLUSTERS["Metro"])
    assert len(events) == 1
    ev = events.iloc[0]
    assert (ev["Space"], ev["Capacity"], ev["Start"], ev["End"]) == ("(pool)", 1, "08:00", "08:01")
    assert (ev["Minutes"], ev["Peak Buses"], ev["Buses"]) == (2, 2, 2)


# ---------------------------------------------------------------------------
# build_layover_runs
# ---------------------------------------------------------------------------


def _run_row(
    ts: str,
    stop: Optional[str],
    status: str,
    trip: str = "",
    prev: str = "",
    nxt: str = "",
) -> dict[str, Any]:
    row = _timeline_row("B1", ts, stop, status, trip=trip)
    row.update(
        {"Prev Trip ID": prev, "Next Trip ID": nxt, "ClusterName": "Metro" if stop else None}
    )
    return row


def _layover_block_df() -> pd.DataFrame:
    """Arrive at stop 200, sit in the bay, leave 7 minutes later; pull in for the day at 09:00."""
    rows = [
        _run_row("07:58", None, "TRAVELING BETWEEN STOPS", trip="T1"),
        _run_row("07:59", "200", "ARRIVE", trip="T1"),
        *[_run_row(f"08:0{m}", "200", "DWELL", trip="T1", prev="T1", nxt="T2") for m in range(5)],
        _run_row("08:05", "200", "LOADING", trip="T2", prev="T1", nxt="T2"),
        _run_row("08:06", "200", "DEPART", trip="T2"),
        _run_row("09:00", "200", "ARRIVE", trip="T3"),
        _run_row("09:01", "200", "ARRIVE", trip="T3", prev="T3"),
        _run_row("09:02", None, "INACTIVE", prev="T3"),
    ]
    return pd.DataFrame(rows)


def test_build_layover_runs_types_flags_and_trips() -> None:
    with patch.object(target, "COUNT_IN_BAY_LAYOVER_AT_STOP", True):
        runs = target.build_layover_runs(_layover_block_df(), "Metro")
    assert len(runs) == 2
    sit, pull_in = runs.iloc[0], runs.iloc[1]
    assert (sit["Type"], sit["Flag"], sit["Start"], sit["End"]) == (
        "in bay",
        "YES",
        "07:59",
        "08:05",
    )
    assert (sit["Total Minutes"], sit["In-Bay Minutes"], sit["Overflow Minutes"]) == (7, 7, 0)
    assert (sit["Arriving Trip ID"], sit["Departing Trip ID"], sit["Ends With"]) == (
        "T1",
        "T2",
        "DEPART",
    )
    assert (sit["Arriving Trip"], sit["Departure Stop"]) == ("101 Downtown", "200")
    assert (pull_in["Type"], pull_in["Flag"], pull_in["Ends With"]) == ("pull-in", "", "INACTIVE")
    assert (pull_in["Start"], pull_in["End"], pull_in["Total Minutes"]) == ("09:00", "09:01", 2)


def test_build_layover_runs_overflow_type_is_not_flagged() -> None:
    rows = [
        _run_row("07:59", "200", "ARRIVE", trip="T1"),
        *[
            _run_row(f"08:{m:02d}", "200", "LAYOVER", trip="T1", prev="T1", nxt="T2")
            for m in range(15)
        ],
        _run_row("08:15", "200", "LOADING", trip="T2", prev="T1", nxt="T2"),
        _run_row("08:16", "200", "DEPART", trip="T2"),
    ]
    runs = target.build_layover_runs(pd.DataFrame(rows), "Metro")
    assert len(runs) == 1
    run = runs.iloc[0]
    assert (run["Type"], run["Flag"], run["Total Minutes"]) == ("overflow", "", 17)
    assert (run["In-Bay Minutes"], run["Overflow Minutes"]) == (2, 15)


def test_build_layover_runs_pull_out_before_first_trip() -> None:
    rows = [
        _run_row("06:55", "200", "LOADING", trip="T1", nxt="T1"),
        _run_row("06:56", "200", "LOADING", trip="T1", nxt="T1"),
        _run_row("06:57", "200", "DEPART", trip="T1"),
    ]
    runs = target.build_layover_runs(pd.DataFrame(rows), "Metro")
    assert list(runs["Type"]) == ["pull-out"]
    assert runs.loc[0, "Departing Trip ID"] == "T1"


def test_build_layover_runs_marks_trip_only_blocks() -> None:
    rows = [
        {**_run_row("06:55", "200", "LOADING", trip="T1", nxt="T1"), "Block": "trip-only:T1"},
        {**_run_row("06:56", "200", "DEPART", trip="T1"), "Block": "trip-only:T1"},
    ]
    runs = target.build_layover_runs(pd.DataFrame(rows), "Metro")
    assert list(runs["Type"]) == ["trip only"]  # the vehicle's earlier trips are unknown


def test_build_layover_runs_only_reports_the_named_cluster() -> None:
    rows = [
        _run_row("06:55", "200", "LOADING", trip="T1", nxt="T1"),
        _run_row("06:56", "200", "DEPART", trip="T1"),
        {**_run_row("07:30", "100", "LOADING", trip="T2", nxt="T2"), "ClusterName": "Park & Ride"},
        {**_run_row("07:31", "100", "DEPART", trip="T2"), "ClusterName": "Park & Ride"},
    ]
    runs = target.build_layover_runs(pd.DataFrame(rows), "Metro")
    assert list(runs["Departing Trip ID"]) == ["T1"]  # the Park & Ride run is not Metro's


# ---------------------------------------------------------------------------
# End to end: Step 1 output -> Step 2 workbook
# ---------------------------------------------------------------------------


def test_step2_end_to_end_on_step1_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import scripts.gtfs_exports.block_status_timeline_exporter as step1

    # Step 1 stops on gtfs_basic's frequency-based trip T2; treat it as scheduled here.
    fixture = tmp_path / "gtfs_basic_scheduled"
    shutil.copytree(Path(__file__).parent / "fixtures" / "gtfs_basic", fixture)
    (fixture / "frequencies.txt").unlink()
    step1_out = tmp_path / "step1"
    monkeypatch.setattr(step1, "GTFS_FOLDER_PATH", str(fixture))
    monkeypatch.setattr(step1, "BLOCK_OUTPUT_FOLDER", str(step1_out))
    monkeypatch.setattr(step1, "CALENDAR_SERVICE_IDS", ["WKDY"])
    monkeypatch.setattr(step1, "WRITE_PER_BLOCK_FILES", False)
    # Blocks B1-B3 end a trip at S6 and start the next one from S7 five minutes later.
    monkeypatch.setattr(step1, "BUS_STOP_CLUSTERS_STEP1", [{"name": "Hub", "stops": ["S6", "S7"]}])
    step1.run_step1_gtfs_to_blocks()

    clusters = {
        "Hub": {
            "single_bay_stops": ["S6", "S7"],
            "double_bay_stops": [],
            "triple_bay_stops": [],
            "overflow_bays": [],
        }
    }
    monkeypatch.setattr(target, "BLOCK_OUTPUT_FOLDER", str(step1_out))
    monkeypatch.setattr(target, "CLUSTER_CONFLICT_OUTPUT_FOLDER", str(tmp_path / "step2"))
    monkeypatch.setattr(target, "CLUSTER_DEFINITIONS", clusters)
    target.run_step2_conflict_detection()

    workbook = tmp_path / "step2" / "Hub_Conflicts.xlsx"
    assert openpyxl.load_workbook(workbook, read_only=True).sheetnames == [
        "Summary",
        "Conflict Events",
        "Overflow Events",
        "Layover Runs",
        "AllStops",
        "Stop_S6",
        "Stop_S7",
    ]
    runs = pd.read_excel(workbook, sheet_name="Layover Runs", dtype=str)
    assert sorted(runs["Block"]) == ["B1", "B2", "B3"]
    assert set(runs["Type"]) == {"in bay"}
    assert set(runs["Flag"]) == {"YES"}  # a short sit in the bay before departing
    assert set(runs["Arrival Stop"]) == {"S6"}
    assert set(runs["Departure Stop"]) == {"S7"}
    b1 = runs[runs["Block"] == "B1"].iloc[0]
    assert (b1["Start"], b1["End"], b1["Total Minutes"]) == ("07:25", "07:29", "5")
    assert (b1["Arriving Trip ID"], b1["Departing Trip ID"], b1["Ends With"]) == (
        "T1",
        "T4",
        "DEPART",
    )
    # Step 1's assumptions are echoed into the Summary sheet.
    summary = pd.read_excel(workbook, sheet_name="Summary", header=None, dtype=str)
    assert any(str(v).startswith("THROUGH_DWELL_MINUTES=") for v in summary[1])
    # A run-log sidecar sits next to the workbook with the CONFIGURATION block verbatim.
    run_log = (tmp_path / "step2" / "Hub_Conflicts_runlog.txt").read_text(encoding="utf-8")
    assert "COUNT_IN_BAY_LAYOVER_AT_STOP: bool = True" in run_log


def test_require_run_log_honours_switch(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(target, "REQUIRE_RUN_LOG", True)
    target.require_run_log(True)
    with pytest.raises(target.RunLogError):
        target.require_run_log(False)
    monkeypatch.setattr(target, "REQUIRE_RUN_LOG", False)
    target.require_run_log(False)


def test_main_returns_1_when_required_run_log_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "BLOCK_OUTPUT_FOLDER", str(tmp_path))
    monkeypatch.setattr(target, "CLUSTER_CONFLICT_OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(
        target, "run_step2_conflict_detection", lambda: target.require_run_log(False)
    )
    assert target.main() == 1


def test_step2_discloses_estimated_times_and_trip_only_blocks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import scripts.gtfs_exports.block_status_timeline_exporter as step1

    # T1 passes the Hub bay S1 at an untimed stop; T2 ends there. Neither trip has a
    # block_id, the route has only a long name and S1 has no stop_name.
    feed = tmp_path / "feed"
    feed.mkdir()
    tables = {
        "routes": "route_id,route_short_name,route_long_name,route_type\nR1,,Crosstown,3\n",
        "stops": "stop_id,stop_name\nS1,\nS2,Elm St\nS3,Oak St\n",
        "trips": "route_id,service_id,trip_id,direction_id\nR1,WKDY,T1,0\nR1,WKDY,T2,1\n",
        "stop_times": "trip_id,arrival_time,departure_time,stop_id,stop_sequence\n"
        "T1,07:00:00,07:00:00,S3,1\nT1,,,S1,2\nT1,07:10:00,07:10:00,S2,3\n"
        "T2,07:00:00,07:00:00,S2,1\nT2,07:05:00,07:05:00,S1,2\n",
    }
    for name, text in tables.items():
        (feed / f"{name}.txt").write_text(text, encoding="utf-8")
    settings = {
        "GTFS_FOLDER_PATH": str(feed),
        "BLOCK_OUTPUT_FOLDER": str(tmp_path / "step1"),
        "SCENARIO_NAME": "",
        "CALENDAR_SERVICE_IDS": ["WKDY"],
        "SERVICE_DATE": "",
        "WRITE_PER_BLOCK_FILES": False,
        "BUS_STOP_CLUSTERS_STEP1": [{"name": "Hub", "stops": ["S1"]}],
        "ROUTE_SHORTNAME_FILTER": [],
        "STOP_ID_FILTER": [],
        "STOP_CODE_FILTER": [],
        "BAY_OVERRIDES": [],
        "THROUGH_DWELL_MINUTES": 2,
        "POST_ARRIVAL_MINUTES": 2,
        "INTERPOLATE_UNTIMED_STOPS": True,
        "TRIP_ONLY_WITHOUT_BLOCK_ID": True,
    }
    for name, value in settings.items():
        monkeypatch.setattr(step1, name, value)
    step1.run_step1_gtfs_to_blocks()

    clusters = {
        "Hub": {
            "single_bay_stops": ["S1"],
            "double_bay_stops": [],
            "triple_bay_stops": [],
            "overflow_bays": [],
        }
    }
    monkeypatch.setattr(target, "BLOCK_OUTPUT_FOLDER", str(tmp_path / "step1"))
    monkeypatch.setattr(target, "SCENARIO_NAME", "")
    monkeypatch.setattr(target, "CLUSTER_CONFLICT_OUTPUT_FOLDER", str(tmp_path / "step2"))
    monkeypatch.setattr(target, "CLUSTER_DEFINITIONS", clusters)
    target.run_step2_conflict_detection()

    workbook = tmp_path / "step2" / "Hub_Conflicts.xlsx"
    events = pd.read_excel(workbook, sheet_name="Conflict Events", dtype=str)
    assert len(events) == 1
    event = events.iloc[0]
    # T1's interpolated 07:05 pass and T2's 07:05 arrival share the one-bus bay.
    assert (event["Stop Name"], event["Start"], event["End"]) == ("S1", "07:05", "07:06")
    assert event["Estimated Times"] == "YES"
    assert event["Bus 1"] == "Crosstown (blk trip-only:T1, trip T1)"
    rows = list(openpyxl.load_workbook(workbook)["Summary"].iter_rows(values_only=True))
    values = {row[0]: row[1] for row in rows}
    assert values["Conflict events involving interpolated stop times"] == 1
    assert (
        values["Trips analyzed without a block_id (layovers and vehicle continuity unknown)"] == 2
    )
    notes = " ".join(str(row[1]) for row in rows)
    assert "ESTIMATED_STOP_VISITS=1" in notes
    assert "TRIP_ONLY_TRIPS=2" in notes
    runs = pd.read_excel(workbook, sheet_name="Layover Runs", dtype=str)
    assert set(runs["Type"]) == {"trip only"}
