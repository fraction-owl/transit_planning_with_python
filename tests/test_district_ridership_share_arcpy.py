"""Tests for scripts/ridership_tools/district_ridership_share_arcpy.py.

ArcPy is unavailable outside ArcGIS Pro, so a stand-in module records the
geoprocessing calls (and the environment active during each) instead of
running them. The ridership loading, allocation, reconciliation, and workbook
steps are plain pandas/openpyxl and run for real.
"""

from __future__ import annotations

import argparse
import importlib
import sys
import types
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook

MODULE = "scripts.ridership_tools.district_ridership_share_arcpy"


class _FakeArcpy(types.ModuleType):
    """Stand-in arcpy that records Project/SpatialJoin calls with the active env."""

    def __init__(self) -> None:
        super().__init__("arcpy")
        self.env = types.SimpleNamespace()
        self.calls: list[tuple[str, tuple[Any, ...], dict[str, Any]]] = []
        self.transformations: list[str] = []
        self.cursor_rows: list[tuple[Any, ...]] = []
        self._env_stack: list[dict[str, Any]] = []
        self.management = types.SimpleNamespace(Project=self._project)
        self.analysis = types.SimpleNamespace(SpatialJoin=self._spatial_join)
        self.da = types.SimpleNamespace(SearchCursor=self._search_cursor)
        self.ExecuteError = type("ExecuteError", (Exception,), {})

    def _active_env(self) -> dict[str, Any]:
        merged: dict[str, Any] = {}
        for layer in self._env_stack:
            merged.update(layer)
        return merged

    def EnvManager(self, **kwargs: Any) -> Any:  # noqa: N802 - arcpy API name
        stack = self._env_stack

        class _Ctx:
            def __enter__(self) -> None:
                stack.append(kwargs)

            def __exit__(self, *exc: object) -> None:
                stack.pop()

        return _Ctx()

    def _project(self, *args: Any) -> None:
        self.calls.append(("Project", args, self._active_env()))

    def _spatial_join(self, **kwargs: Any) -> None:
        self.calls.append(("SpatialJoin", tuple(kwargs.items()), self._active_env()))

    def _search_cursor(self, fc: str, fields: list[str]) -> Any:
        rows = self.cursor_rows

        class _Cursor:
            def __enter__(self) -> list[tuple[Any, ...]]:
                return rows

            def __exit__(self, *exc: object) -> None:
                return None

        return _Cursor()

    def Describe(self, path: str) -> Any:  # noqa: N802 - arcpy API name
        return types.SimpleNamespace(
            spatialReference=types.SimpleNamespace(name="NAD_1983_StatePlane"),
            extent="stops-extent",
        )

    def SpatialReference(self, code: int) -> Any:  # noqa: N802 - arcpy API name
        return types.SimpleNamespace(name=f"EPSG:{code}")

    def ListTransformations(self, from_sr: Any, to_sr: Any, extent: Any) -> list[str]:  # noqa: N802
        return list(self.transformations)

    def ValidateTableName(self, name: str, workspace: str) -> str:  # noqa: N802
        return name

    def GetMessages(self) -> str:  # noqa: N802 - arcpy API name
        return ""


@pytest.fixture()
def arcpy_stub() -> _FakeArcpy:
    return _FakeArcpy()


@pytest.fixture()
def mod(monkeypatch: pytest.MonkeyPatch, arcpy_stub: _FakeArcpy) -> Iterator[types.ModuleType]:
    """Import the script against the stand-in ``arcpy`` without leaking either module."""
    monkeypatch.setitem(sys.modules, "arcpy", arcpy_stub)
    sys.modules.pop(MODULE, None)
    yield importlib.import_module(MODULE)
    sys.modules.pop(MODULE, None)


def _excel(tmp_path: Path, df: pd.DataFrame) -> str:
    path = tmp_path / "ridership.xlsx"
    df.to_excel(path, index=False)
    return str(path)


def _ridership(stop_ids: list[Any], boardings: list[float], alightings: list[float]) -> dict:
    return {"STOP_ID": stop_ids, "BOARD_ALL": boardings, "ALIGHT_ALL": alightings}


# --- ridership loading -------------------------------------------------------


def test_time_period_filter_without_column_fails_loud(
    mod: types.ModuleType, tmp_path: Path
) -> None:
    path = _excel(tmp_path, pd.DataFrame(_ridership(["1", "2"], [10, 90], [0, 0])))
    with pytest.raises(ValueError, match="TIME_PERIODS given"):
        mod.load_stop_ridership(path, time_periods=["AM PEAK"])
    # No filter requested: every row is kept.
    assert sum(v["boardings"] for v in mod.load_stop_ridership(path).values()) == 100


def test_text_stop_ids_keep_leading_zeros_and_literal_na(
    mod: types.ModuleType, tmp_path: Path
) -> None:
    df = pd.DataFrame(_ridership(["00123", "00456", "NA"], [1, 2, 3], [0, 0, 0]))
    result = mod.load_stop_ridership(_excel(tmp_path, df))
    assert sorted(result) == ["00123", "00456", "NA"]


def test_numeric_stop_ids_render_without_decimal(mod: types.ModuleType, tmp_path: Path) -> None:
    df = pd.DataFrame(_ridership([1001, 1002], [1, 2], [0, 0]))
    assert sorted(mod.load_stop_ridership(_excel(tmp_path, df))) == ["1001", "1002"]


def test_blank_stop_id_ridership_reaches_reconciliation(
    mod: types.ModuleType, tmp_path: Path
) -> None:
    df = pd.DataFrame(_ridership(["1", None], [10.0, 90.0], [0.0, 0.0]))
    stop_ridership = mod.load_stop_ridership(_excel(tmp_path, df))
    _, _, diag = mod.allocate_ridership_to_districts({"1": {"A"}}, stop_ridership, {"1"})
    recon = mod.build_reconciliation_frame(diag).set_index("category")
    assert recon.loc["missing_stop_id", "boardings"] == 90.0
    assert recon.loc["grand_total", "boardings"] == 100.0


def test_all_blank_stop_ids_fail_loud(mod: types.ModuleType, tmp_path: Path) -> None:
    df = pd.DataFrame(_ridership([None, None], [1.0, 2.0], [0.0, 0.0]))
    with pytest.raises(ValueError, match="No usable ridership rows"):
        mod.load_stop_ridership(_excel(tmp_path, df))


# --- rounding ----------------------------------------------------------------


def test_per_stop_values_are_not_rounded(mod: types.ModuleType, tmp_path: Path) -> None:
    df = pd.DataFrame(_ridership(["1"], [1.24], [1.24]))
    vals = mod.load_stop_ridership(_excel(tmp_path, df))["1"]
    assert vals["total"] == pytest.approx(vals["boardings"] + vals["alightings"])


def test_split_district_totals_reconcile(mod: types.ModuleType) -> None:
    ridership = {"1": {"boardings": 1.1, "alightings": 0.0, "total": 1.1}}
    district_df, alloc_df, diag = mod.allocate_ridership_to_districts(
        {"1": {"A", "B"}}, ridership, {"1"}
    )
    recon = mod.build_reconciliation_frame(diag).set_index("category")
    assert district_df["boardings"].sum() == pytest.approx(1.1)
    assert alloc_df["boardings_alloc"].sum() == pytest.approx(1.1)
    assert recon.loc["allocated_to_districts", "boardings"] == pytest.approx(1.1)


def test_workbook_formats_display_but_keeps_values(mod: types.ModuleType, tmp_path: Path) -> None:
    ridership = {"1": {"boardings": 1.1, "alightings": 0.0, "total": 1.1}}
    district_df, alloc_df, diag = mod.allocate_ridership_to_districts(
        {"1": {"A", "B"}}, ridership, {"1"}
    )
    out = tmp_path / "out.xlsx"
    mod.write_output_workbook(str(out), district_df, alloc_df, mod.build_reconciliation_frame(diag))

    ws = load_workbook(out)["district_ridership"]
    headers = [c.value for c in ws[1]]
    boardings = ws.cell(row=2, column=headers.index("boardings") + 1)
    pct = ws.cell(row=2, column=headers.index("pct_boardings") + 1)
    district = ws.cell(row=2, column=headers.index("district") + 1)
    assert boardings.value == pytest.approx(0.55)
    assert boardings.number_format == "0.0"
    assert pct.number_format == "0.0"
    assert district.number_format == "General"


# --- GTFS stops / districts ---------------------------------------------------


def test_filter_stops_drops_impossible_coordinates(mod: types.ModuleType) -> None:
    stops = pd.DataFrame(
        {
            "stop_code": ["lat91", "inf", "text", "blank", "ok"],
            "stop_lat": ["91", "inf", "abc", None, "38.9"],
            "stop_lon": ["-77", "-77", "-77", "-77", "-77.0"],
        }
    )
    out = mod.filter_stops(stops, "stop_code")
    assert out["stop_key"].tolist() == ["ok"]
    assert out["stop_lat"].tolist() == [38.9]


def test_blank_district_label_fails_loud(mod: types.ModuleType, arcpy_stub: _FakeArcpy) -> None:
    arcpy_stub.cursor_rows = [("1", "A"), ("2", "  "), ("3", None)]
    with pytest.raises(ValueError, match="blank DISTRICT_FIELD"):
        mod.extract_stop_districts("joined", "DISTRICT")


# --- spatial environment / transformation ---------------------------------------


def test_unset_transformation_is_resolved_and_recorded(
    mod: types.ModuleType, arcpy_stub: _FakeArcpy
) -> None:
    arcpy_stub.transformations = ["T_FIRST", "T_SECOND"]
    _, applied = mod.project_stops_to_district_sr("stops", "districts", "gdb", None)
    name, args, env = arcpy_stub.calls[-1]
    assert name == "Project"
    assert args[-1] == "T_FIRST"
    assert env == {"extent": "MAXOF"}
    assert applied.startswith("T_FIRST")


def test_no_listed_transformation_is_recorded_as_none(
    mod: types.ModuleType, arcpy_stub: _FakeArcpy
) -> None:
    _, applied = mod.project_stops_to_district_sr("stops", "districts", "gdb", None)
    assert len(arcpy_stub.calls[-1][1]) == 3  # no transformation argument
    assert applied.startswith("none")


def test_named_transformation_is_used_verbatim(
    mod: types.ModuleType, arcpy_stub: _FakeArcpy
) -> None:
    arcpy_stub.transformations = ["T_FIRST"]
    _, applied = mod.project_stops_to_district_sr("stops", "districts", "gdb", "T_CHOSEN")
    assert arcpy_stub.calls[-1][1][-1] == "T_CHOSEN"
    assert applied == "T_CHOSEN"


def test_spatial_join_runs_in_controlled_environment(
    mod: types.ModuleType, arcpy_stub: _FakeArcpy
) -> None:
    mod.containment_join_stops_to_districts("stops", "districts", "gdb")
    name, _, env = arcpy_stub.calls[-1]
    assert name == "SpatialJoin"
    assert env["extent"] == "MAXOF"
    assert env["outputCoordinateSystem"].name == "NAD_1983_StatePlane"


# --- run log requirement ------------------------------------------------------------


def test_failed_run_log_removes_workbook(
    mod: types.ModuleType, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    ridership = {"1": {"boardings": 5.0, "alightings": 1.0, "total": 6.0}}
    stops = pd.DataFrame({"stop_key": ["1"], "stop_lon": [-77.0], "stop_lat": [38.9]})
    monkeypatch.setattr(mod, "ensure_work_gdb", lambda work_dir, name: "gdb")
    monkeypatch.setattr(mod, "load_gtfs_stops", lambda path: stops)
    monkeypatch.setattr(mod, "filter_stops", lambda df, key: df)
    monkeypatch.setattr(mod, "load_stop_ridership", lambda *a: ridership)
    monkeypatch.setattr(mod, "validate_district_field", lambda fc, field: None)
    monkeypatch.setattr(mod, "stops_to_points", lambda df, gdb: "pts")
    monkeypatch.setattr(mod, "project_stops_to_district_sr", lambda *a: ("proj", "none"))
    monkeypatch.setattr(mod, "containment_join_stops_to_districts", lambda *a: "sj")
    monkeypatch.setattr(mod, "extract_stop_districts", lambda fc, field: {"1": {"A"}})
    monkeypatch.setattr(mod, "write_run_log", lambda *a, **k: False)
    args = argparse.Namespace(
        work_dir=str(tmp_path),
        gtfs_input="gtfs",
        excel_file="ridership.xlsx",
        districts_fc="districts.shp",
        district_field="DISTRICT",
        gtfs_join_key="stop_code",
        routes=[],
        routes_exclude=[],
        time_periods=[],
        geo_transformation=None,
        boundary_allocation="split",
        pct_base="allocated",
        output_dir=str(tmp_path),
    )
    with pytest.raises(RuntimeError, match="workbook was removed"):
        mod.run_analysis(args)
    assert not (tmp_path / mod.OUTPUT_EXCEL_NAME).exists()
