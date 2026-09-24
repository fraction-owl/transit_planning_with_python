from __future__ import annotations

from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

import scripts.facilities_tools.bay_change_sweep as target
import scripts.gtfs_exports.block_status_timeline_exporter as step1

# ---------------------------------------------------------------------------
# Synthetic Step 1 timeline: two blocks meeting at a two-bay cluster
# ---------------------------------------------------------------------------

CLUSTER = {"S1": "A", "S2": "B"}


def _row(
    block: str,
    ts: str,
    status: str,
    trip: str,
    route: str,
    stop: str = "",
    arr: str = "",
    dep: str = "",
) -> dict[str, Any]:
    return {
        "Timestamp": ts,
        "Block": block,
        "Route": f"R{route}",
        "Route Short Name": route,
        "Direction": "0",
        "Trip Headsign": f"Route {route} headsign",
        "Trip ID": trip,
        "Stop ID": stop,
        "Stop Name": f"Stop {stop}" if stop else "",
        "Stop Sequence": "1" if stop else "",
        "Arrival Time": arr,
        "Departure Time": dep,
        "Status": status,
        "Layover Location": "",
        "Prev Trip ID": "",
        "Next Trip ID": "",
        "Timepoint": "0",
    }


def _timeline() -> pd.DataFrame:
    """Block B1 (route 10) and block B2 (routes 20 then 21) both arrive at bay A at 08:00.

    B1: T1 arrives S1 08:00, sits, departs as T2 from S1 at 08:10 (gap 10).
    B2: T3 arrives S1 08:00, sits, departs as T4 (route 21) from S1 at 08:20 (interline, gap 20).
    Bay A therefore has two buses at 08:00 and 08:01: two conflict minutes.
    """
    rows = [
        # B1 / T1: route 10, starts at S9 07:50, ends at S1 08:00
        _row("B1", "07:50", "DEPART", "T1", "10", "S9", "07:50", "07:50"),
        *[_row("B1", f"07:5{m}", "TRAVELING BETWEEN STOPS", "T1", "10") for m in range(1, 10)],
        _row("B1", "08:00", "ARRIVE", "T1", "10", "S1", "08:00", "08:00"),
        _row("B1", "08:01", "ARRIVE", "T1", "10", "S1", "08:00", "08:10"),
        *[_row("B1", f"08:0{m}", "DWELL", "T1", "10", "S1", "08:00", "08:10") for m in range(2, 5)],
        *[
            _row("B1", f"08:0{m}", "LOADING", "T2", "10", "S1", "08:00", "08:10")
            for m in range(5, 10)
        ],
        # B1 / T2: route 10, starts at S1 08:10, ends at S9 08:20
        _row("B1", "08:10", "DEPART", "T2", "10", "S1", "08:10", "08:10"),
        *[_row("B1", f"08:1{m}", "TRAVELING BETWEEN STOPS", "T2", "10") for m in range(1, 10)],
        _row("B1", "08:20", "ARRIVE", "T2", "10", "S9", "08:20", "08:20"),
        # B2 / T3: route 20, starts at S8 07:45, ends at S1 08:00
        _row("B2", "07:45", "DEPART", "T3", "20", "S8", "07:45", "07:45"),
        *[_row("B2", f"07:{m}", "TRAVELING BETWEEN STOPS", "T3", "20") for m in range(46, 60)],
        _row("B2", "08:00", "ARRIVE", "T3", "20", "S1", "08:00", "08:00"),
        _row("B2", "08:01", "ARRIVE", "T3", "20", "S1", "08:00", "08:20"),
        *[
            _row("B2", f"08:{m:02d}", "LAYOVER", "T3", "20", "S1", "08:00", "08:20")
            for m in range(2, 15)
        ],
        *[
            _row("B2", f"08:{m}", "LOADING", "T4", "21", "S1", "08:00", "08:20")
            for m in range(15, 20)
        ],
        # B2 / T4: route 21, starts at S1 08:20, ends at S8 08:35
        _row("B2", "08:20", "DEPART", "T4", "21", "S1", "08:20", "08:20"),
        *[_row("B2", f"08:{m}", "TRAVELING BETWEEN STOPS", "T4", "21") for m in range(21, 35)],
        _row("B2", "08:35", "ARRIVE", "T4", "21", "S8", "08:35", "08:35"),
    ]
    return pd.DataFrame(rows)


@pytest.fixture
def step1_folder(tmp_path: Path) -> Path:
    folder = tmp_path / "step1"
    folder.mkdir()
    _timeline().to_csv(folder / "all_blocks_timeline.csv", index=False)
    return folder


# The same schedule as GTFS, exported by Step 1 with the manifest and schedule snapshot the
# sweep needs. POST_ARRIVAL_MINUTES = 1 reproduces the hand-built timeline above.
# A visit is (stop, time) or (stop, arrival, departure).
_GTFS_TRIPS: list[tuple[str, str, str, list[tuple[str, ...]]]] = [
    ("T1", "10", "B1", [("S9", "07:50:00"), ("S1", "08:00:00")]),
    ("T2", "10", "B1", [("S1", "08:10:00"), ("S9", "08:20:00")]),
    ("T3", "20", "B2", [("S8", "07:45:00"), ("S1", "08:00:00")]),
    ("T4", "21", "B2", [("S1", "08:20:00"), ("S8", "08:35:00")]),
]


def _export(
    gtfs: Path,
    out: Path,
    monkeypatch: pytest.MonkeyPatch,
    trips: list[tuple[str, str, str, list[tuple[str, ...]]]] = _GTFS_TRIPS,
    route_names: dict[str, tuple[str, str]] | None = None,
    **step1_settings: Any,
) -> Path:
    """Write *trips* as a GTFS feed and run Step 1 on it; return the run folder.

    *route_names* maps a route to its (short, long) names; by default route "10" is
    route_id R10 with short name "10".
    """
    gtfs.mkdir(parents=True)
    stops = ["S1", "S2", "S8", "S9"]
    pd.DataFrame({"stop_id": stops, "stop_name": [f"Stop {s}" for s in stops]}).to_csv(
        gtfs / "stops.txt", index=False
    )
    names = route_names or {r: (r, "") for r in ("10", "20", "21")}
    pd.DataFrame(
        {
            "route_id": [f"R{r}" for r in names],
            "route_short_name": [short for short, _ in names.values()],
            "route_long_name": [long_name for _, long_name in names.values()],
            "route_type": 3,
        }
    ).to_csv(gtfs / "routes.txt", index=False)
    pd.DataFrame(
        [
            {
                "route_id": f"R{route}",
                "service_id": "WKDY",
                "trip_id": trip_id,
                "direction_id": "0",
                "block_id": block,
                "trip_headsign": f"Route {route} headsign",
            }
            for trip_id, route, block, _ in trips
        ]
    ).to_csv(gtfs / "trips.txt", index=False)
    pd.DataFrame(
        [
            {
                "trip_id": trip_id,
                "arrival_time": times[0],
                "departure_time": times[-1],
                "stop_id": stop,
                "stop_sequence": seq,
            }
            for trip_id, _, _, visits in trips
            for seq, (stop, *times) in enumerate(visits, start=1)
        ]
    ).to_csv(gtfs / "stop_times.txt", index=False)
    settings: dict[str, Any] = {
        "GTFS_FOLDER_PATH": str(gtfs),
        "BLOCK_OUTPUT_FOLDER": str(out),
        "SCENARIO_NAME": "",
        "CALENDAR_SERVICE_IDS": ["WKDY"],
        "SERVICE_DATE": "",
        "ROUTE_SHORTNAME_FILTER": [],
        "STOP_ID_FILTER": [],
        "STOP_CODE_FILTER": [],
        "BAY_OVERRIDES": [],
        "WRITE_PER_BLOCK_FILES": False,
        "BUS_STOP_CLUSTERS_STEP1": [{"name": "TC", "stops": ["S1", "S2"]}],
        "THROUGH_DWELL_MINUTES": 2,
        "PRE_DEPARTURE_MINUTES": 5,
        "POST_ARRIVAL_MINUTES": 1,
        "IN_BAY_LAYOVER_MAX_MINUTES": 10,
        "LAYOVER_THRESHOLD": 20,
        "INTERPOLATE_UNTIMED_STOPS": False,
        "TRIP_ONLY_WITHOUT_BLOCK_ID": False,
        **step1_settings,
    }
    for name, value in settings.items():
        monkeypatch.setattr(step1, name, value)
    step1.run_step1_gtfs_to_blocks()
    return out


@pytest.fixture
def exported_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    return _export(tmp_path / "gtfs", tmp_path / "step1_run", monkeypatch)


@pytest.fixture
def cluster(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(target, "CLUSTER_STOPS", CLUSTER)
    monkeypatch.setattr(target, "CLUSTER_CAPACITY", {})
    monkeypatch.setattr(target, "SAME_BAY_GROUPS", [])
    monkeypatch.setattr(target, "SHIFT_TOGETHER", [])


_NAMES = target.RouteNames({"R10": "10", "R20": "20", "R21": "21"})


def _end(text: str) -> Any:
    """The structured route-end CONFIGURATION text such as "10 arrive" stands for."""
    route_end = _NAMES.parse(text)
    assert route_end is not None, text
    return route_end


def _change(moves: dict[str, str] | None = None, shifts: dict[str, int] | None = None) -> Any:
    """A Change written with CONFIGURATION route-end names."""
    return target.Change(
        {_end(key): bay for key, bay in (moves or {}).items()},
        {_end(key): value for key, value in (shifts or {}).items()},
    )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("ts", "expected"), [("08:05", 485), ("25:10", 1510), ("", None), ("x", None), (None, None)]
)
def test_timestamp_to_minutes(ts: str | None, expected: int | None) -> None:
    assert target.timestamp_to_minutes(ts) == expected


def test_minutes_to_hhmm_past_midnight() -> None:
    assert target.minutes_to_hhmm(1590) == "26:30"
    assert target.minutes_to_hhmm(None) == ""


# ---------------------------------------------------------------------------
# Trips, chains, route-ends
# ---------------------------------------------------------------------------


def test_load_timeline_reads_csv_and_labels_bays(step1_folder: Path, cluster: None) -> None:
    df = target.load_timeline(str(step1_folder))
    assert set(df["Bay"]) == {"A", ""}
    assert df["Minute"].dtype.kind == "i"


def test_build_trips_scheduled_ends(step1_folder: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(step1_folder))).set_index("trip_id")
    assert list(trips.index) == ["T1", "T2", "T3", "T4"]
    assert (trips.loc["T1", "start"], trips.loc["T1", "end"]) == (470, 480)
    assert (trips.loc["T1", "first_bay"], trips.loc["T1", "last_bay"]) == ("", "A")
    assert (trips.loc["T2", "start"], trips.loc["T2", "first_bay"]) == (490, "A")
    # Routes are identified by route_id and labeled by their short name.
    assert (trips.loc["T4", "route"], trips.loc["T4", "route_name"]) == ("R21", "21")


def test_build_trips_empty_input_has_columns() -> None:
    empty = pd.DataFrame(columns=list(_row("B", "08:00", "DEPART", "T", "1")) + ["Minute", "Bay"])
    trips = target.build_trips(empty)
    assert trips.empty
    assert list(trips.columns) == target.TRIP_COLUMNS


def test_block_chains_gaps_and_interlines(step1_folder: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(step1_folder)))
    chains = target.build_block_chains(trips).set_index("block")
    assert chains.loc["B1", ["schedule_gap_min", "occupancy_gap_min"]].tolist() == [10, 10]
    assert chains.loc["B1", "where"] == "same stop (Bay A)"
    assert not chains.loc["B1", "interline"]
    assert chains.loc["B2", "schedule_gap_min"] == 20
    assert chains.loc["B2", "interline"]
    assert (chains.loc["B2", "from_route_end"], chains.loc["B2", "to_route_end"]) == (
        "20 arrive",
        "21 depart",
    )


def test_route_ends_and_interline_summary(step1_folder: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(step1_folder)))
    chains = target.build_block_chains(trips)
    ends = target.build_route_ends(trips, chains).set_index("route_end")
    assert sorted(ends.index) == ["10 arrive", "10 depart", "20 arrive", "21 depart"]
    assert ends.loc["10 arrive", "bays"] == "A"
    assert ends.loc["10 arrive", "min_schedule_gap_after"] == 10
    assert ends.loc["21 depart", "interlines_with"] == "20"
    inter = target.interline_summary(chains)
    assert len(inter) == 1
    assert (inter.iloc[0]["from_route"], inter.iloc[0]["to_route"]) == ("20", "21")
    assert not inter.iloc[0]["shift_together_needed"]


def test_route_ends_with_no_chains() -> None:
    trips = pd.DataFrame(
        [
            {
                "trip_id": "T1",
                "block": "B1",
                "route": "R10",
                "route_name": "10",
                "headsign": "h",
                "direction": "0",
                "start": 470,
                "end": 480,
                "first_stop": "S9",
                "first_bay": "",
                "last_stop": "S1",
                "last_bay": "A",
                "through_bays": "",
                "departure": 470,
                "arrival": 480,
                "through_visits": [],
            }
        ]
    )
    chains = target.build_block_chains(trips)
    assert chains.empty
    ends = target.build_route_ends(trips, chains)
    assert ends.loc[0, "min_schedule_gap_before"] == ""
    assert target.interline_summary(chains).empty


# ---------------------------------------------------------------------------
# Change / expand_groups / Standard / Feasibility
# ---------------------------------------------------------------------------


def test_change_combine_and_compatible() -> None:
    a = _change(moves={"10 arrive": "B"})
    b = _change(shifts={"20": 2})
    c = _change(shifts={"20": -2})
    assert a.compatible(b)
    assert not b.compatible(c)  # one route cannot shift both ways
    assert not b.compatible(b)  # nothing new to add
    combined = a.combine(b)
    assert combined.label() == "10 arrive to Bay B; 20 +2 min"
    assert b.combine(_change(shifts={"20": 2})).shifts == {_end("20"): 2}  # applied once
    with pytest.raises(ValueError, match="conflicting"):
        b.combine(c)


def _groups(same_bay: list[list[str]], together: list[list[str]]) -> Any:
    return target.Selections(
        same_bay_groups=[[_end(member) for member in group] for group in same_bay],
        shift_together=[[_end(member) for member in group] for group in together],
    )


def test_expand_groups_refuses_contradictions_and_follows_overlaps() -> None:
    groups = _groups([], [["20 arrive", "21 depart"]])
    with pytest.raises(ValueError, match="Conflicting"):
        target.expand_groups(_change(shifts={"20 arrive": 2, "21 depart": 1}), groups)
    groups = _groups([["10 depart", "20 arrive"], ["10 arrive", "10 depart"]], [])
    out = target.expand_groups(_change(moves={"10 arrive": "B"}), groups)
    assert out.moves == {_end("10 arrive"): "B", _end("10 depart"): "B", _end("20 arrive"): "B"}


def test_expand_groups_moves_members_together() -> None:
    groups = _groups([["10 arrive", "10 depart"]], [["20 arrive", "21 depart"]])
    out = target.expand_groups(_change(moves={"10 arrive": "B"}, shifts={"20 arrive": 3}), groups)
    assert out.moves == {_end("10 arrive"): "B", _end("10 depart"): "B"}
    assert out.shifts == {_end("20 arrive"): 3, _end("21 depart"): 3}


def test_standard_baseline_conflicts(exported_run: Path, cluster: None) -> None:
    df = target.load_timeline(str(exported_run))
    std = target.Standard("x", df, target.build_trips(df))
    total, per_bay, conflict = std.score(std.occ)
    assert (total, per_bay) == (2, {"A": 2})
    who, kinds = std.detail(std.occ, conflict)
    assert kinds == {"boarding/boarding": 2}
    assert {tuple(map(str, ends)) for ends in who.values()} == {("10 arrive", "20 arrive")}


def test_standard_bay_move_clears_conflict(exported_run: Path, cluster: None) -> None:
    df = target.load_timeline(str(exported_run))
    std = target.Standard("x", df, target.build_trips(df))
    assert std.score(std.apply(_change(moves={"20 arrive": "B"})))[0] == 0


def test_standard_shift_rebuilds_the_block(exported_run: Path, cluster: None) -> None:
    df = target.load_timeline(str(exported_run))
    std = target.Standard("x", df, target.build_trips(df))
    # Route 10 departing 5 minutes earlier turns T1's arrival buffer into loading for T2;
    # bay A still holds one route-10 bus and one route-20 bus at 08:00-08:01.
    rebuilt = std.apply(_change(shifts={"10 depart": -5}))
    assert std.score(rebuilt)[0] == 2
    assert not rebuilt.duplicated(["Block", "Minute"]).any()
    b1 = rebuilt[rebuilt["Block"] == "B1"].set_index("Minute")["Status"]
    assert list(b1.loc[480:485]) == ["ARRIVE", "LOADING", "LOADING", "LOADING", "LOADING", "DEPART"]


# Each change, and the GTFS edit that makes the same change for a fresh Step 1 run.
_EDITS: dict[str, tuple[dict[str, Any], dict[str, list[tuple[str, str]]]]] = {
    "bay move": ({"moves": {"20 arrive": "B"}}, {"T3": [("S8", "07:45:00"), ("S2", "08:00:00")]}),
    "route-end shift": (
        {"shifts": {"10 depart": -5}},
        {"T2": [("S1", "08:05:00"), ("S9", "08:15:00")]},
    ),
    "whole-route shift": ({"shifts": {"21": 3}}, {"T4": [("S1", "08:23:00"), ("S8", "08:38:00")]}),
    # Changes that move a turn between the bay and the overflow space.
    "earlier arrival leaves the bay": (
        {"shifts": {"10 arrive": -3}},
        {"T1": [("S9", "07:47:00"), ("S1", "07:57:00")]},  # 13-minute turn: overflow
    ),
    "interline shift into the bay": (
        {"shifts": {"21": -10}},
        {"T4": [("S1", "08:10:00"), ("S8", "08:25:00")]},  # 10-minute turn: in the bay
    ),
}


@pytest.mark.parametrize("case", sorted(_EDITS))
def test_standard_rebuild_matches_a_fresh_export(
    case: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, cluster: None
) -> None:
    change, edits = _EDITS[case]
    baseline = _export(tmp_path / "gtfs", tmp_path / "base", monkeypatch)
    edited = [(tid, route, blk, edits.get(tid, visits)) for tid, route, blk, visits in _GTFS_TRIPS]
    fresh = _export(tmp_path / "gtfs_edited", tmp_path / "edited", monkeypatch, edited)
    df = target.load_timeline(str(baseline))
    std = target.Standard("x", df, target.build_trips(df))
    columns = ["Block", "Minute", "Bay", "Status", "route_end"]
    rebuilt = std.apply(_change(**change))[columns].astype(str)
    expected = std._occupancy(target.load_timeline(str(fresh)))[columns].astype(str)
    assert sorted(map(tuple, rebuilt.to_numpy())) == sorted(map(tuple, expected.to_numpy()))


def test_feasibility_rejects_gap_below_minimum(step1_folder: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(step1_folder)))
    feas = target.Feasibility(trips, target.build_block_chains(trips))
    ok, reason, tightest = feas.check(_change(shifts={"10 depart": -8}))
    assert not ok
    assert "schedule gap 10 08:00 -> 10 08:10 (same stop (Bay A)) goes 10 -> 2 min" in reason
    assert reason.endswith("below 3")
    ok, reason, tightest = feas.check(_change(shifts={"10 depart": -5}))
    assert (ok, reason, tightest) == (True, "", 5)
    # Interline hand-offs use the lower minimum.
    ok, _, tightest = feas.check(_change(shifts={"21 depart": -18}))
    assert (ok, tightest) == (True, 2)
    assert not feas.check(_change(shifts={"21 depart": -19}))[0]


def test_trip_shift_applies_overlapping_selectors_once(exported_run: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(exported_run)))
    feas = target.Feasibility(trips, target.build_block_chains(trips))
    shifts = feas.trip_shift({_end("10"): 2, _end("10 depart"): 2})
    shift = dict(zip(feas.trips["trip_id"], shifts))
    assert (shift["T1"], shift["T2"], shift["T3"]) == (2, 2, 0)  # T2 moves 2, not 4
    with pytest.raises(ValueError, match="different shifts"):
        feas.trip_shift({_end("10"): 2, _end("10 depart"): 3})


def test_feasibility_off_clockface_note(step1_folder: Path, cluster: None) -> None:
    trips = target.build_trips(target.load_timeline(str(step1_folder)))
    feas = target.Feasibility(trips, target.build_block_chains(trips))
    # T4 departs the cluster at 08:20 (not on :00/:30), so a shift is not flagged...
    assert feas.off_clockface(_change(shifts={"21 depart": 1})) == ""
    # ...while a whole-route shift of route 20 moves nothing that departs the cluster.
    assert feas.off_clockface(_change(shifts={"20": 1})) == ""


def test_is_win_and_rank_key(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(target, "REQUIRE_NO_WORSE_ON", "all")
    assert target.is_win({"a": 2, "b": 0})
    assert not target.is_win({"a": 2, "b": -1})
    assert not target.is_win({"a": 0, "b": 0})
    assert not target.is_win(None)
    monkeypatch.setattr(target, "REQUIRE_NO_WORSE_ON", "any")
    assert target.is_win({"a": 2, "b": -1})
    assert target.rank_key({"a": 1, "b": 5}) == (-5, -1)


# ---------------------------------------------------------------------------
# Discover and sweep end to end
# ---------------------------------------------------------------------------


def test_run_discover_writes_workbook_and_stub(
    step1_folder: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(step1_folder)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "SCENARIO_LABEL", "test")
    out = target.run_discover()
    wb = openpyxl.load_workbook(out, read_only=True)
    assert wb.sheetnames == [
        "Route-ends Direct",
        "Block chains Direct",
        "Interlines Direct",
        "Config stub",
    ]
    stub = "\n".join(str(v) for v in pd.read_excel(out, sheet_name="Config stub")["config_stub"])
    assert '"route_end": "20 arrive"' in stub
    assert '"21": [-3, -2, -1, 1, 2, 3],  # whole route' in stub
    run_log = (tmp_path / "out" / "test_discover_runlog.txt").read_text(encoding="utf-8")
    assert "MIN_LAYOVER_MINUTES = 3" in run_log


def test_run_sweep_finds_bay_move(
    exported_run: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        target, "TIMELINES", {"Direct": str(exported_run), "Likely": str(exported_run)}
    )
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "SCENARIO_LABEL", "test")
    monkeypatch.setattr(target, "BAY_CANDIDATES", [{"route_end": "20 arrive", "bays": ["A", "B"]}])
    monkeypatch.setattr(target, "SHIFT_CANDIDATES", {"10 depart": [-8, -5]})
    out = target.run_sweep()

    sheets = openpyxl.load_workbook(out, read_only=True).sheetnames
    assert sheets == [
        "Baseline",
        "Easy wins",
        "Pairs",
        "Packages",
        "Package finalists",
        "Rejected",
        "Config used",
    ]
    baseline = pd.read_excel(out, sheet_name="Baseline")
    assert list(baseline["standard"]) == ["Direct", "Likely"]
    assert list(baseline["total"]) == [2, 2]
    assert list(baseline["Bay A"]) == [2, 2]

    wins = pd.read_excel(out, sheet_name="Easy wins")
    assert list(wins["change"]) == ["20 arrive to Bay B"]
    assert (wins.loc[0, "Direct before"], wins.loc[0, "Direct after"]) == (2, 0)
    assert wins.loc[0, "Direct removed"] == "A 08:00-08:01 10/20"

    rejected = pd.read_excel(out, sheet_name="Rejected").set_index("change")
    assert "below 3" in rejected.loc["10 depart -8 min", "reason"]
    assert rejected.loc["10 depart -5 min", "reason"] == "no improvement"

    pairs = pd.read_excel(out, sheet_name="Pairs")
    assert list(pairs["change"]) == ["20 arrive to Bay B; 10 depart -5 min"]

    packages = pd.read_excel(out, sheet_name="Packages")
    assert list(packages["package"]) == ["20 arrive to Bay B"]
    assert packages.loc[0, "Direct conflict minutes"] == 0
    finalists = pd.read_excel(out, sheet_name="Package finalists")
    assert finalists.loc[0, "change"] == "20 arrive to Bay B"
    assert finalists.loc[0, "Direct before"] == 2  # scored against the untouched baseline
    assert (tmp_path / "out" / "test_sweep_runlog.txt").is_file()


def test_run_sweep_refuses_legacy_timelines(
    step1_folder: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(step1_folder)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    with pytest.raises(ValueError, match="revised exporter"):
        target.run_sweep()


def test_run_discover_on_exported_run_with_pull_out_loading(
    exported_run: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Pull-out LOADING rows carry no Arrival Time; trip ends must still be read correctly.
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(exported_run)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "SCENARIO_LABEL", "test")
    chains = pd.read_excel(target.run_discover(), sheet_name="Block chains Direct")
    assert sorted(chains["schedule_gap_min"]) == [10, 20]


def test_legacy_trips_ignore_blank_arrival_times(tmp_path: Path, cluster: None) -> None:
    rows = [
        *[_row("B1", f"07:4{m}", "LOADING", "T1", "10", "S9", "", "07:50") for m in range(5, 10)],
        *_timeline().to_dict("records"),
    ]
    folder = tmp_path / "legacy"
    folder.mkdir()
    pd.DataFrame(rows).to_csv(folder / "all_blocks_timeline.csv", index=False)
    trips = target.build_trips(target.load_timeline(str(folder))).set_index("trip_id")
    assert (trips.loc["T1", "start"], trips.loc["T1", "end"]) == (470, 480)
    chains = target.build_block_chains(trips.reset_index())
    assert chains["schedule_gap_min"].tolist() == [10, 20]


def _legacy_trips(folder: Path, edit: dict[tuple[str, str], dict[str, str]]) -> Any:
    """build_trips on the hand-built legacy timeline with rows of (trip, status) edited."""
    rows = [
        {**row, **edit.get((row["Trip ID"], row["Status"]), edit.get((row["Trip ID"], "*"), {}))}
        for row in _timeline().to_dict("records")
    ]
    folder.mkdir()
    pd.DataFrame(rows).to_csv(folder / "all_blocks_timeline.csv", index=False)
    return target.build_trips(target.load_timeline(str(folder)))


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        ({("T4", "*"): {"Arrival Time": ""}}, "Trip 'T4' in block 'B2' has no readable arrival"),
        ({("T2", "DEPART"): {"Arrival Time": ""}}, "Trip 'T2' in block 'B1' has no readable first"),
        (
            {("T2", "DEPART"): {"Departure Time": ""}},
            "Trip 'T2' in block 'B1' has no readable first",
        ),
        ({("T1", "ARRIVE"): {"Arrival Time": ""}}, "Trip 'T1' in block 'B1' has no readable final"),
        ({("T2", "DEPART"): {"Status": "LOADING"}}, "Trip 'T2' in block 'B1' has no DEPART row"),
    ],
)
def test_legacy_trips_name_the_trip_and_block_with_unreadable_times(
    tmp_path: Path, cluster: None, edit: dict[tuple[str, str], dict[str, str]], message: str
) -> None:
    with pytest.raises(ValueError, match=message):
        _legacy_trips(tmp_path / "legacy", edit)


def test_legacy_trips_read_the_scheduled_departure_after_a_hold(
    tmp_path: Path, cluster: None
) -> None:
    # T2's DEPART row: it reached bay A at 08:02 and departs at 08:10.
    trips = _legacy_trips(tmp_path / "legacy", {("T2", "DEPART"): {"Arrival Time": "08:02"}})
    t2 = trips.set_index("trip_id").loc["T2"]
    assert (t2["start"], t2["departure"]) == (482, 490)
    b1 = target.build_block_chains(trips).set_index("block").loc["B1"]
    assert (b1["departs"], b1["schedule_gap_min"], b1["occupancy_gap_min"]) == ("08:10", 10, 2)


def test_run_discover_requires_timelines(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    with pytest.raises(ValueError, match="TIMELINES is empty; list at least one Step 1"):
        target.run_discover()
    assert not (tmp_path / "out").exists()


def test_main_returns_1_when_required_run_log_fails(
    step1_folder: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(step1_folder)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "DISCOVER_ONLY", True)
    monkeypatch.setattr(target, "write_run_log", lambda output_file: False)
    assert target.main() == 1
    monkeypatch.setattr(target, "REQUIRE_RUN_LOG", False)
    assert target.main() == 0


def test_main_returns_2_on_placeholders(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": r"Path\To\block_status_output\x"})
    assert target.main() == 2


def test_main_returns_2_when_folder_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(tmp_path / "missing")})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    assert target.main() == 2


def test_main_runs_discover(
    step1_folder: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(step1_folder)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "DISCOVER_ONLY", True)
    with patch.object(target, "run_sweep") as sweep:
        assert target.main() == 0
    sweep.assert_not_called()
    assert (tmp_path / "out" / f"{target.SCENARIO_LABEL}_discover.xlsx").is_file()


# ---------------------------------------------------------------------------
# Route identity and names, schedule vs occupancy gaps, trip-only runs, errors
# ---------------------------------------------------------------------------


def test_route_names_label_shared_names_and_parse_structured_route_ends() -> None:
    names = target.RouteNames(
        {"R1": "Express", "R2": "Express", "R3": "R1", "R4": "Green Line", "R5": "R5"}
    )
    assert names.label == {
        "R1": "Express [R1]",
        "R2": "Express [R2]",
        "R3": "R1 [R3]",  # its name is another route's route_id
        "R4": "Green Line",
        "R5": "R5",  # no short or long name: named by its route_id
    }
    assert names.parse("Green Line through 0") == target.RouteEnd("R4", "through", "0")
    assert str(names.parse("Green Line through 0")) == "Green Line through 0"
    assert names.parse("Green Line") == target.RouteEnd("R4")
    assert names.parse("Express [R2] arrive") == target.RouteEnd("R2", "arrive")
    assert names.parse("R2 depart") == target.RouteEnd("R2", "depart")  # a route_id
    assert names.parse("Blue Line arrive") is None
    with pytest.raises(ValueError, match=r"could mean Express \[R1\] arrive, Express \[R2\]"):
        names.parse("Express arrive")
    with pytest.raises(ValueError, match="could mean"):
        names.parse("R1")  # route R1's ID, and route R3's name


def test_sweep_names_routes_with_spaces_and_without_short_names(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    run = _export(
        tmp_path / "gtfs",
        tmp_path / "run",
        monkeypatch,
        route_names={"10": ("", "Green Line"), "20": ("20", "Crosstown"), "21": ("", "")},
    )
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(run)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "SCENARIO_LABEL", "names")
    monkeypatch.setattr(
        target, "BAY_CANDIDATES", [{"route_end": "Green Line arrive", "bays": ["B"]}]
    )
    monkeypatch.setattr(target, "SHIFT_CANDIDATES", {"Green Line": [2], "R21 depart": [1]})
    out = target.run_sweep()
    wins = pd.read_excel(out, sheet_name="Easy wins").set_index("change")
    # Moving route 10's arrivals to bay B, or running route 10 two minutes later, each
    # clears the 08:00-08:01 conflict with route 20 in bay A.
    assert set(wins.index) == {"Green Line arrive to Bay B", "Green Line +2 min"}
    assert wins.loc["Green Line arrive to Bay B", "Direct removed"] == (
        "A 08:00-08:01 Green Line/20"
    )
    rejected = pd.read_excel(out, sheet_name="Rejected")
    assert list(rejected["change"]) == ["R21 depart +1 min"]  # no short or long name

    monkeypatch.setattr(target, "DISCOVER_ONLY", True)
    discover = target.run_discover()
    ends = pd.read_excel(discover, sheet_name="Route-ends Direct")
    assert set(ends["route_end"]) == {
        "Green Line arrive",
        "Green Line depart",
        "20 arrive",
        "R21 depart",
    }
    stub = "\n".join(
        str(v) for v in pd.read_excel(discover, sheet_name="Config stub")["config_stub"]
    )
    assert '"Green Line": [-3, -2, -1, 1, 2, 3],  # whole route' in stub


# T2 reaches bay A at 08:02 and holds there until its scheduled 08:10 departure.
_HOLD_TRIPS: list[tuple[str, str, str, list[tuple[str, ...]]]] = [
    ("T1", "10", "B1", [("S9", "07:50:00"), ("S1", "08:00:00")]),
    ("T2", "10", "B1", [("S1", "08:02:00", "08:10:00"), ("S9", "08:20:00")]),
    ("T3", "20", "B2", [("S8", "07:45:00"), ("S1", "08:00:00")]),
    ("T4", "21", "B2", [("S1", "08:20:00"), ("S8", "08:35:00")]),
]


def test_block_chains_report_scheduled_times_and_name_both_gaps(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    run = _export(tmp_path / "gtfs", tmp_path / "run", monkeypatch, _HOLD_TRIPS)
    trips = target.build_trips(target.load_timeline(str(run)))
    b1 = target.build_block_chains(trips).set_index("block").loc["B1"]
    assert (b1["arrives"], b1["departs"]) == ("08:00", "08:10")  # T2's scheduled departure
    assert (b1["schedule_gap_min"], b1["occupancy_gap_min"]) == (10, 2)


def test_feasibility_keeps_the_hold_when_a_shift_leaves_the_schedule_gap_legal(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    run = _export(tmp_path / "gtfs", tmp_path / "run", monkeypatch, _HOLD_TRIPS)
    trips = target.build_trips(target.load_timeline(str(run)))
    feas = target.Feasibility(trips, target.build_block_chains(trips))
    # Three minutes earlier leaves a 7-minute schedule gap, but T2 would reach the bay
    # at 07:59, before T1 has arrived: the vehicle cannot run both.
    ok, reason, _ = feas.check(_change(shifts={"10 depart": -3}))
    assert not ok
    assert reason == "trip T2 (10) would begin 1 min before trip T1 (10) ends on block B1"
    assert feas.check(_change(shifts={"10 depart": -2})) == (True, "", 8)


def test_standard_rebuild_of_a_held_trip_matches_a_fresh_export(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    baseline = _export(tmp_path / "gtfs", tmp_path / "base", monkeypatch, _HOLD_TRIPS)
    edited = [
        (
            tid,
            route,
            blk,
            [("S1", "08:00:00", "08:08:00"), ("S9", "08:18:00")] if tid == "T2" else v,
        )
        for tid, route, blk, v in _HOLD_TRIPS
    ]
    fresh = _export(tmp_path / "gtfs_edited", tmp_path / "edited", monkeypatch, edited)
    df = target.load_timeline(str(baseline))
    std = target.Standard("x", df, target.build_trips(df))
    columns = ["Block", "Minute", "Bay", "Status", "route_end"]
    rebuilt = std.apply(_change(shifts={"10 depart": -2}))[columns].astype(str)
    expected = std._occupancy(target.load_timeline(str(fresh)))[columns].astype(str)
    assert sorted(map(tuple, rebuilt.to_numpy())) == sorted(map(tuple, expected.to_numpy()))


def test_discover_and_sweep_refuse_trip_only_runs(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    blockless = [(tid, route, "", visits) for tid, route, _, visits in _GTFS_TRIPS]
    run = _export(
        tmp_path / "gtfs", tmp_path / "run", monkeypatch, blockless, TRIP_ONLY_WITHOUT_BLOCK_ID=True
    )
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(run)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    with pytest.raises(ValueError, match=r"trip-only mode: 4 trip\(s\) have no block_id"):
        target.run_discover()
    with pytest.raises(ValueError, match="trip-only mode"):
        target.run_sweep()


def test_run_discover_reports_the_real_error_and_writes_no_workbook(
    step1_folder: Path, tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    # An interrupted Step 1 run: the manifest says it never completed.
    (step1_folder / "timeline_manifest.json").write_text(
        '{"schema_version": 1, "status": "in_progress", "interval_minutes": 1}', encoding="utf-8"
    )
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(step1_folder)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    monkeypatch.setattr(target, "SCENARIO_LABEL", "test")
    monkeypatch.setattr(target, "DISCOVER_ONLY", True)
    with pytest.raises(ValueError, match="not complete"):  # not "no visible sheet"
        target.run_discover()
    assert not (tmp_path / "out" / "test_discover.xlsx").exists()
    assert target.main() == 1


def test_sweep_discloses_interpolated_stop_times(
    tmp_path: Path, cluster: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    trips = [
        (tid, route, blk, [visits[0], ("S2", ""), visits[1]] if tid == "T1" else visits)
        for tid, route, blk, visits in _GTFS_TRIPS
    ]
    run = _export(
        tmp_path / "gtfs", tmp_path / "run", monkeypatch, trips, INTERPOLATE_UNTIMED_STOPS=True
    )
    monkeypatch.setattr(target, "TIMELINES", {"Direct": str(run)})
    monkeypatch.setattr(target, "OUTPUT_FOLDER", str(tmp_path / "out"))
    config = pd.read_excel(target.run_sweep(), sheet_name="Config used").set_index("setting")
    assert str(config.loc["Stop visits with interpolated times (Step 1)", "value"]) == "1"
