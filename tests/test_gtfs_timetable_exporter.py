from __future__ import annotations

import logging
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from scripts.gtfs_exports.gtfs_timetable_exporter import (
    MAX_COLUMN_WIDTH,
    MISSING_TIME,
    adjust_time,
    apply_in_out_filters,
    build_column_headers,
    build_column_layouts,
    build_service_calendar,
    build_service_id_schedule_map,
    check_schedule_order,
    embed_pattern,
    export_to_excel_multiple_sheets,
    format_output_folder_name,
    format_workbook_name,
    get_all_route_short_names,
    map_service_id_to_schedule,
    merge_stop_patterns,
    prepare_timepoints,
    remove_empty_schedule_columns,
    time_to_minutes,
)

GTFS_BASIC = Path(__file__).parent / "fixtures" / "gtfs_basic"

# ---------------------------------------------------------------------------
# time_to_minutes
# ---------------------------------------------------------------------------


def test_time_to_minutes_24h() -> None:
    assert time_to_minutes("08:30") == 510


def test_time_to_minutes_midnight() -> None:
    assert time_to_minutes("00:00") == 0


def test_time_to_minutes_past_midnight() -> None:
    assert time_to_minutes("25:00") == 1500


def test_time_to_minutes_sentinel_returns_none() -> None:
    assert time_to_minutes(MISSING_TIME) is None


def test_time_to_minutes_non_string_returns_none() -> None:
    assert time_to_minutes(None) is None  # type: ignore[arg-type]


def test_time_to_minutes_am_pm() -> None:
    assert time_to_minutes("8:30 AM") == 510


def test_time_to_minutes_pm() -> None:
    assert time_to_minutes("1:00 PM") == 780


# ---------------------------------------------------------------------------
# adjust_time
# ---------------------------------------------------------------------------


def test_adjust_time_24h_format() -> None:
    assert adjust_time("8:05:00", "24") == "08:05"


def test_adjust_time_12h_am() -> None:
    assert adjust_time("08:05", "12") == "8:05 AM"


def test_adjust_time_12h_pm() -> None:
    assert adjust_time("13:00", "12") == "1:00 PM"


def test_adjust_time_none_returns_none() -> None:
    assert adjust_time(None, "24") is None  # type: ignore[arg-type]


def test_adjust_time_sentinel_returns_sentinel() -> None:
    assert adjust_time(MISSING_TIME, "24") == MISSING_TIME


# ---------------------------------------------------------------------------
# prepare_timepoints
# ---------------------------------------------------------------------------


def _stop_times_with_timepoint() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "trip_id": ["T1", "T1", "T1"],
            "stop_sequence": ["1", "2", "3"],
            "stop_id": ["S1", "S2", "S3"],
            "timepoint": ["1", "0", "1"],
        }
    )


def test_prepare_timepoints_filters_to_one() -> None:
    df = prepare_timepoints(_stop_times_with_timepoint())
    assert len(df) == 2


def test_prepare_timepoints_all_rows_returned_when_no_column() -> None:
    df = pd.DataFrame(
        {
            "trip_id": ["T1", "T1"],
            "stop_sequence": ["1", "2"],
            "stop_id": ["S1", "S2"],
        }
    )
    result = prepare_timepoints(df)
    assert len(result) == 2


def test_prepare_timepoints_stop_sequence_is_numeric() -> None:
    df = prepare_timepoints(_stop_times_with_timepoint())
    assert pd.api.types.is_numeric_dtype(df["stop_sequence"])


# ---------------------------------------------------------------------------
# remove_empty_schedule_columns
# ---------------------------------------------------------------------------


def test_remove_empty_schedule_columns_drops_blank_cols() -> None:
    df = pd.DataFrame(
        {
            "Stop A Schedule": [MISSING_TIME, MISSING_TIME],
            "Stop B Schedule": ["07:05", MISSING_TIME],
            "Route Name": ["101", "101"],
        }
    )
    result = remove_empty_schedule_columns(df)
    assert "Stop A Schedule" not in result.columns
    assert "Stop B Schedule" in result.columns


def test_remove_empty_schedule_columns_preserves_non_schedule_cols() -> None:
    df = pd.DataFrame(
        {
            "Route Name": ["101"],
            "Stop A Schedule": [MISSING_TIME],
        }
    )
    result = remove_empty_schedule_columns(df)
    assert "Route Name" in result.columns


# ---------------------------------------------------------------------------
# map_service_id_to_schedule
# ---------------------------------------------------------------------------


def _cal_row(days: dict) -> pd.Series:
    base = {
        d: "0"
        for d in ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
    }
    base.update(days)
    return pd.Series(base)


def test_map_weekday() -> None:
    row = _cal_row(
        {"monday": "1", "tuesday": "1", "wednesday": "1", "thursday": "1", "friday": "1"}
    )
    assert map_service_id_to_schedule(row) == "Weekday"


def test_map_saturday() -> None:
    row = _cal_row({"saturday": "1"})
    assert map_service_id_to_schedule(row) == "Saturday"


def test_map_sunday() -> None:
    row = _cal_row({"sunday": "1"})
    assert map_service_id_to_schedule(row) == "Sunday"


def test_map_weekend() -> None:
    row = _cal_row({"saturday": "1", "sunday": "1"})
    assert map_service_id_to_schedule(row) == "Weekend"


def test_map_daily() -> None:
    row = _cal_row(
        {
            d: "1"
            for d in ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
        }
    )
    assert map_service_id_to_schedule(row) == "Daily"


def test_map_no_days_returns_holiday() -> None:
    row = _cal_row({})
    assert map_service_id_to_schedule(row) == "Holiday"


# ---------------------------------------------------------------------------
# format_output_folder_name
# ---------------------------------------------------------------------------


def test_format_output_folder_name_weekday() -> None:
    name = format_output_folder_name("1", "Weekday")
    assert name == "weekday_sid_1"


def test_format_output_folder_name_saturday() -> None:
    name = format_output_folder_name("2", "Saturday")
    assert name == "saturday_sid_2"


# ---------------------------------------------------------------------------
# get_all_route_short_names
# ---------------------------------------------------------------------------


def test_get_all_route_short_names_sorted() -> None:
    routes = pd.DataFrame({"route_short_name": ["202", "101", "303"]})
    assert get_all_route_short_names(routes) == ["101", "202", "303"]


def test_get_all_route_short_names_drops_na() -> None:
    routes = pd.DataFrame({"route_short_name": ["101", None, "202"]})
    result = get_all_route_short_names(routes)
    assert None not in result
    assert len(result) == 2


# ---------------------------------------------------------------------------
# apply_in_out_filters
# ---------------------------------------------------------------------------


def test_apply_in_out_filters_filter_in() -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    orig_in = mod.FILTER_IN_ROUTES
    orig_out = mod.FILTER_OUT_ROUTES
    mod.FILTER_IN_ROUTES = ["101"]
    mod.FILTER_OUT_ROUTES = []
    try:
        result = apply_in_out_filters(["101", "202", "303"])
        assert result == ["101"]
    finally:
        mod.FILTER_IN_ROUTES = orig_in
        mod.FILTER_OUT_ROUTES = orig_out


def test_apply_in_out_filters_filter_out() -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    orig_in = mod.FILTER_IN_ROUTES
    orig_out = mod.FILTER_OUT_ROUTES
    mod.FILTER_IN_ROUTES = []
    mod.FILTER_OUT_ROUTES = ["202"]
    try:
        result = apply_in_out_filters(["101", "202", "303"])
        assert "202" not in result
        assert len(result) == 2
    finally:
        mod.FILTER_IN_ROUTES = orig_in
        mod.FILTER_OUT_ROUTES = orig_out


# ---------------------------------------------------------------------------
# build_service_id_schedule_map
# ---------------------------------------------------------------------------


def test_build_service_id_schedule_map_basic() -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    orig = mod.SERVICE_LABEL_OVERRIDES
    mod.SERVICE_LABEL_OVERRIDES = {}
    cal = pd.DataFrame(
        {
            "service_id": ["1"],
            "monday": ["1"],
            "tuesday": ["1"],
            "wednesday": ["1"],
            "thursday": ["1"],
            "friday": ["1"],
            "saturday": ["0"],
            "sunday": ["0"],
        }
    )
    try:
        mapping = build_service_id_schedule_map(cal)
        assert mapping["1"] == "Weekday"
    finally:
        mod.SERVICE_LABEL_OVERRIDES = orig


def test_build_service_id_schedule_map_override_takes_precedence() -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    orig = mod.SERVICE_LABEL_OVERRIDES
    mod.SERVICE_LABEL_OVERRIDES = {"1": "CustomLabel"}
    cal = pd.DataFrame(
        {
            "service_id": ["1"],
            "monday": ["1"],
            "tuesday": ["1"],
            "wednesday": ["1"],
            "thursday": ["1"],
            "friday": ["1"],
            "saturday": ["0"],
            "sunday": ["0"],
        }
    )
    try:
        mapping = build_service_id_schedule_map(cal)
        assert mapping["1"] == "CustomLabel"
    finally:
        mod.SERVICE_LABEL_OVERRIDES = orig


# ---------------------------------------------------------------------------
# export_to_excel_multiple_sheets — real openpyxl output
# ---------------------------------------------------------------------------


def test_export_to_excel_multiple_sheets_writes_real_workbook(tmp_path: Path) -> None:
    out = tmp_path / "schedule.xlsx"
    df = pd.DataFrame({"Route": ["R1", "R1"], "Main St & Doe St": ["07:10", "12:10"]})

    export_to_excel_multiple_sheets({"Weekday_0": df, "Saturday_0": pd.DataFrame()}, str(out))

    wb = openpyxl.load_workbook(out)
    # The empty DataFrame's sheet is skipped, the populated one written.
    assert wb.sheetnames == ["Weekday_0"]
    ws = wb["Weekday_0"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Route", "Main St & Doe St")
    assert rows[1] == ("R1", "07:10")
    assert rows[2] == ("R1", "12:10")
    # Styling applied by the real export path.
    assert ws["A1"].alignment.wrap_text is True
    assert ws["A2"].alignment.horizontal == "left"
    assert ws.column_dimensions["A"].width <= MAX_COLUMN_WIDTH


def test_export_to_excel_multiple_sheets_empty_dict_writes_nothing(tmp_path: Path) -> None:
    out = tmp_path / "schedule.xlsx"
    export_to_excel_multiple_sheets({}, str(out))
    assert not out.exists()


# ---------------------------------------------------------------------------
# main — full pipeline against gtfs_basic, real xlsx output
# ---------------------------------------------------------------------------


def test_main_writes_real_schedule_workbooks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    monkeypatch.setattr(mod, "GTFS_FOLDER_PATH", str(GTFS_BASIC))
    monkeypatch.setattr(mod, "BASE_OUTPUT_PATH", str(tmp_path))
    monkeypatch.setattr(mod, "FILTER_SERVICE_IDS", ["WKDY"])
    monkeypatch.setattr(mod, "FILTER_IN_ROUTES", [])

    assert mod.main() == 0

    # WKDY has no label override; its Monday-Friday calendar row makes it "Weekday".
    folder = tmp_path / "weekday_sid_WKDY"
    assert sorted(p.name for p in folder.glob("*.xlsx")) == [
        f"route_R{n}_timetable_weekday_sid_WKDY.xlsx" for n in (1, 2, 3)
    ]

    wb = openpyxl.load_workbook(folder / "route_R3_timetable_weekday_sid_WKDY.xlsx")
    assert wb.sheetnames == ["Direction_0"]
    header, *rows = wb["Direction_0"].iter_rows(values_only=True)
    # gtfs_basic has no timepoint column, so every stop gets a schedule column.
    assert header[3:] == (
        "Robin Hood Rd & Old Mill Rd Schedule",
        "Main St & Robin Hood Rd Schedule",
        "Main St & Doe St Schedule",
        "Doe St & Fordson Rd Schedule",
    )
    # One row per trip, in time order.
    assert [row[:2] for row in rows] == [("R3", "0")] * 3
    assert [row[3:] for row in rows] == [
        ("08:00", "08:05", "08:10", "08:15"),
        ("14:00", "14:05", "14:10", "14:15"),
        ("18:00", "18:05", "18:10", "18:15"),
    ]


def test_main_drops_stop_with_no_scheduled_times(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    # Blank S2's times on every R1 trip, as GTFS allows at non-timepoint stops.
    feed = tmp_path / "gtfs"
    shutil.copytree(GTFS_BASIC, feed)
    stop_times = pd.read_csv(feed / "stop_times.txt", dtype=str)
    blank = stop_times["trip_id"].isin(["T1", "T2", "T3"]) & (stop_times["stop_id"] == "S2")
    stop_times.loc[blank, ["arrival_time", "departure_time"]] = ""
    stop_times.to_csv(feed / "stop_times.txt", index=False)

    monkeypatch.setattr(mod, "GTFS_FOLDER_PATH", str(feed))
    monkeypatch.setattr(mod, "BASE_OUTPUT_PATH", str(tmp_path / "out"))
    monkeypatch.setattr(mod, "FILTER_SERVICE_IDS", ["WKDY"])
    monkeypatch.setattr(mod, "FILTER_IN_ROUTES", ["R1"])

    assert mod.main() == 0

    folder = tmp_path / "out" / "weekday_sid_WKDY"
    wb = openpyxl.load_workbook(folder / "route_R1_timetable_weekday_sid_WKDY.xlsx")
    header, *rows = wb["Direction_0"].iter_rows(values_only=True)
    # S2 (Main St & Robin Hood Rd) has no time on any trip, so its column is dropped.
    assert header[3:] == (
        "Main St & Mt Vernon Ln Schedule",
        "Main St & Doe St Schedule",
        "Main St & Lockheed Blvd Schedule",
        "Main St & Ox Rd Schedule",
        "Main St & Beacon Hill Rd Schedule",
    )
    assert [row[3:] for row in rows] == [
        ("07:00", "07:10", "07:15", "07:20", "07:25"),
        ("12:00", "12:10", "12:15", "12:20", "12:25"),
        ("17:00", "17:10", "17:15", "17:20", "17:25"),
    ]


# ---------------------------------------------------------------------------
# Time parsing and display
# ---------------------------------------------------------------------------


def test_time_to_minutes_accepts_gtfs_seconds() -> None:
    assert time_to_minutes("08:20:00") == 500
    assert time_to_minutes("25:15:00") == 1515


def test_adjust_time_12h_wraps_past_midnight() -> None:
    assert adjust_time("25:15:00", "12") == "1:15 AM"
    assert adjust_time("24:00:00", "12") == "12:00 AM"


# ---------------------------------------------------------------------------
# prepare_timepoints — blank timepoint column, final-stop flag
# ---------------------------------------------------------------------------


def test_prepare_timepoints_blank_column_keeps_all_rows() -> None:
    df = _stop_times_with_timepoint().assign(timepoint=[None, None, None])
    assert len(prepare_timepoints(df)) == 3


def test_prepare_timepoints_flags_final_stop_before_filtering() -> None:
    df = _stop_times_with_timepoint().assign(timepoint=["1", "1", "0"])
    result = prepare_timepoints(df)
    # S3 is the trip's final stop but not a timepoint, so no kept row is final.
    assert result["stop_id"].tolist() == ["S1", "S2"]
    assert not result["is_final_stop"].any()


# ---------------------------------------------------------------------------
# Column layout: pattern merging, alignment, headers
# ---------------------------------------------------------------------------


def test_merge_stop_patterns_keeps_both_branches() -> None:
    assert merge_stop_patterns(["A", "B", "D"], ["A", "C", "D"]) == ["A", "B", "C", "D"]


def test_merge_stop_patterns_adds_short_turn_extension() -> None:
    assert merge_stop_patterns(["B", "C"], ["A", "B", "C", "D"]) == ["A", "B", "C", "D"]


def test_build_column_layouts_splits_conflicting_stop_order() -> None:
    layouts, assignment = build_column_layouts({("A", "B", "C"): 3, ("C", "B", "A"): 1})
    assert layouts == [["A", "B", "C"], ["C", "B", "A"]]
    assert assignment == {("A", "B", "C"): 0, ("C", "B", "A"): 1}


def test_build_column_layouts_keeps_loop_in_one_layout() -> None:
    layouts, _ = build_column_layouts({("A", "B", "C", "A"): 2, ("C", "A"): 1})
    assert layouts == [["A", "B", "C", "A"]]


def test_embed_pattern_places_loop_tail_on_later_visit() -> None:
    assert embed_pattern(["C", "A"], ["A", "B", "C", "A"]) == [2, 3]
    assert embed_pattern(["A", "B"], ["A", "B", "C", "A"]) == [0, 1]


def test_build_column_headers_disambiguates_shared_names() -> None:
    stops = pd.DataFrame(
        {
            "stop_id": ["A", "B", "C"],
            "stop_code": ["101", "102", "103"],
            "stop_name": ["Central Station", "Main St", "Central Station"],
        }
    )
    assert build_column_headers(["A", "B", "C", "A"], stops) == [
        "Central Station [101] Schedule",
        "Main St Schedule",
        "Central Station [103] Schedule",
        "Central Station [101] (2) Schedule",
    ]


def test_build_column_headers_falls_back_to_stop_id_without_codes() -> None:
    stops = pd.DataFrame({"stop_id": ["A", "C"], "stop_name": ["Central", "Central"]})
    assert build_column_headers(["A", "C"], stops) == [
        "Central [A] Schedule",
        "Central [C] Schedule",
    ]


# ---------------------------------------------------------------------------
# check_schedule_order — reports pass/fail
# ---------------------------------------------------------------------------


def test_check_schedule_order_counts_violations(caplog: pytest.LogCaptureFixture) -> None:
    minutes = pd.DataFrame([[510.0, 490.0]], index=["T1"], columns=["A Schedule", "B Schedule"])
    with caplog.at_level(logging.INFO):
        assert check_schedule_order(minutes, "101", "Weekday", "0") == 1
    assert "passed" not in caplog.text
    assert "Trip 'T1'" in caplog.text


def test_check_schedule_order_records_each_violation() -> None:
    minutes = pd.DataFrame([[510.0, 490.0]], index=["T1"], columns=["A Schedule", "B Schedule"])
    issues: list[str] = []
    assert check_schedule_order(minutes, "101", "Weekday", "0", issues) == 1
    assert len(issues) == 1
    assert "Trip 'T1': 'B Schedule' time 08:10" in issues[0]


def test_check_schedule_order_passes_clean_grid(caplog: pytest.LogCaptureFixture) -> None:
    minutes = pd.DataFrame([[480.0, None], [490.0, 500.0]], index=["T1", "T2"], columns=["A", "B"])
    with caplog.at_level(logging.INFO):
        assert check_schedule_order(minutes, "101", "Weekday", "0") == 0
    assert "passed" in caplog.text


# ---------------------------------------------------------------------------
# Names and calendars
# ---------------------------------------------------------------------------


def test_format_workbook_name_includes_service_id_and_sanitizes() -> None:
    assert format_workbook_name("101", "Weekday", "1") == "route_101_timetable_weekday_sid_1.xlsx"
    assert format_workbook_name("10/11", "Weekday", "A:B") == (
        "route_10_11_timetable_weekday_sid_A_B.xlsx"
    )
    assert format_workbook_name("101", "Weekday", "1", flagged=True) == (
        "route_101_timetable_weekday_sid_1_warning.xlsx"
    )


def test_build_service_calendar_adds_calendar_dates_only_services() -> None:
    calendar = _cal_row({"monday": "1"}).to_frame().T.assign(service_id="WK")
    calendar_dates = pd.DataFrame(
        {
            # 2026-01-03 and 2026-01-10 are Saturdays; WK's removal is ignored.
            "service_id": ["SAT", "SAT", "WK"],
            "date": ["20260103", "20260110", "20260105"],
            "exception_type": ["1", "1", "2"],
        }
    )
    result = build_service_calendar(calendar, calendar_dates).set_index("service_id")
    assert sorted(result.index) == ["SAT", "WK"]
    assert map_service_id_to_schedule(result.loc["SAT"]) == "Saturday"
    assert map_service_id_to_schedule(result.loc["WK"]) == "Special"


# ---------------------------------------------------------------------------
# main — synthetic feeds for alignment, branches, and failure reporting
# ---------------------------------------------------------------------------

_CALENDAR = (
    "service_id,monday,tuesday,wednesday,thursday,friday,saturday,sunday,start_date,end_date\n"
    "1,1,1,1,1,1,0,0,20260101,20261231\n"
)


def _write_feed(
    folder: Path,
    stops: str,
    stop_times: str,
    trips: str = "route_id,service_id,trip_id,trip_headsign,direction_id\n"
    "R,1,T1,East,0\nR,1,T2,East,0\n",
    calendar: str | None = _CALENDAR,
    calendar_dates: str | None = None,
) -> Path:
    folder.mkdir(parents=True)
    files = {
        "agency.txt": "agency_id,agency_name,agency_url,agency_timezone\n"
        "A,Agency,http://x,America/New_York\n",
        "routes.txt": "route_id,agency_id,route_short_name,route_type\nR,A,101,3\n",
        "stops.txt": "stop_id,stop_name,stop_lat,stop_lon\n" + stops,
        "trips.txt": trips,
        "stop_times.txt": "trip_id,arrival_time,departure_time,stop_id,stop_sequence,timepoint\n"
        + stop_times,
        "calendar.txt": calendar,
        "calendar_dates.txt": calendar_dates,
    }
    for name, text in files.items():
        if text is not None:
            (folder / name).write_text(text)
    return folder


_WORKBOOK = "route_101_timetable_weekday_sid_1.xlsx"
_FLAGGED_WORKBOOK = "route_101_timetable_weekday_sid_1_warning.xlsx"
_RUN_LOG = "gtfs_timetable_exporter_runlog.txt"


def _run_main(
    monkeypatch: pytest.MonkeyPatch, feed: Path, out: Path, workbook_name: str = _WORKBOOK
) -> tuple[int, list[tuple]]:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    monkeypatch.setattr(mod, "GTFS_FOLDER_PATH", str(feed))
    monkeypatch.setattr(mod, "BASE_OUTPUT_PATH", str(out))
    monkeypatch.setattr(mod, "FILTER_SERVICE_IDS", ["1"])
    monkeypatch.setattr(mod, "FILTER_IN_ROUTES", [])
    monkeypatch.setattr(mod, "SERVICE_LABEL_OVERRIDES", {"1": "Weekday"})
    monkeypatch.setattr(mod, "TIME_FORMAT_OPTION", "24")
    rc = mod.main()
    workbook = out / "weekday_sid_1" / workbook_name
    if not workbook.is_file():
        return rc, []
    wb = openpyxl.load_workbook(workbook)
    return rc, [(ws.title, list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]


_ABCD = "A,Alpha,0,0\nB,Bravo,0,0\nC,Charlie,0,0\nD,Delta,0,0\nX,Detour,0,0\n"


def test_main_aligns_trips_by_stop_order_not_sequence_numbers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # T2 renumbers its stops and adds a non-timepoint detour, shifting B's and C's sequence.
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\nT1,08:20:00,08:20:00,C,3,1\n"
        "T2,09:00:00,09:00:00,A,10,1\nT2,,,X,15,0\nT2,09:12:00,09:12:00,B,20,1\n"
        "T2,09:22:00,09:22:00,C,30,1\n",
    )
    rc, sheets = _run_main(monkeypatch, feed, tmp_path / "out")
    assert rc == 0
    ((title, (header, *rows)),) = sheets
    assert header[3:] == ("Alpha Schedule", "Bravo Schedule", "Charlie Schedule")
    assert [row[3:] for row in rows] == [
        ("08:00", "08:10", "08:20"),
        ("09:00", "09:12", "09:22"),
    ]


def test_main_keeps_timepoints_on_every_branch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\nT1,08:20:00,08:20:00,D,3,1\n"
        "T2,09:00:00,09:00:00,A,1,1\nT2,09:15:00,09:15:00,C,2,1\nT2,09:25:00,09:25:00,D,3,1\n",
    )
    rc, sheets = _run_main(monkeypatch, feed, tmp_path / "out")
    assert rc == 0
    ((_, (header, *rows)),) = sheets
    assert header[3:] == ("Alpha Schedule", "Bravo Schedule", "Charlie Schedule", "Delta Schedule")
    assert [row[3:] for row in rows] == [
        ("08:00", "08:10", MISSING_TIME, "08:20"),
        ("09:00", MISSING_TIME, "09:15", "09:25"),
    ]


def test_main_shared_stop_names_get_distinct_columns(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(
        tmp_path / "gtfs",
        "A,Central Station,0,0\nB,Main St,0,0\nC,Central Station,0,0\n",
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\nT1,08:20:00,08:20:00,C,3,1\n",
        trips="route_id,service_id,trip_id,trip_headsign,direction_id\nR,1,T1,East,0\n",
    )
    rc, sheets = _run_main(monkeypatch, feed, tmp_path / "out")
    assert rc == 0
    ((_, (header, row)),) = sheets
    assert header[3:] == (
        "Central Station [A] Schedule",
        "Main St Schedule",
        "Central Station [C] Schedule",
    )
    assert row[3:] == ("08:00", "08:10", "08:20")


def test_main_shows_departure_and_final_arrival_ordered_by_shared_stops(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # T1 dwells at B and at its final stop D; T2 is a short trip that starts later
    # but serves B and C before T1, so it sorts above T1.
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:15:00,B,2,1\n"
        "T1,08:25:00,08:25:00,C,3,1\nT1,08:40:00,08:45:00,D,4,1\n"
        "T2,08:10:00,08:12:00,B,1,1\nT2,08:20:00,08:22:00,C,2,1\n",
    )
    rc, sheets = _run_main(monkeypatch, feed, tmp_path / "out")
    assert rc == 0
    ((_, (_, *rows)),) = sheets
    assert [row[3:] for row in rows] == [
        (MISSING_TIME, "08:12", "08:20", MISSING_TIME),
        ("08:00", "08:15", "08:25", "08:40"),
    ]


def test_main_exports_trips_without_direction_id(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n",
        trips="route_id,service_id,trip_id,trip_headsign\nR,1,T1,East\n",
    )
    # Trips without a direction_id are flagged for review.
    rc, sheets = _run_main(monkeypatch, feed, tmp_path / "out", _FLAGGED_WORKBOOK)
    assert rc == 0
    assert [title for title, _ in sheets] == ["No_Direction"]


def test_main_reads_calendar_dates_only_zip(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    folder = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n",
        trips="route_id,service_id,trip_id,trip_headsign,direction_id\nR,1,T1,East,0\n",
        calendar=None,
        # 2026-01-05 through 2026-01-09 are Monday-Friday.
        calendar_dates="service_id,date,exception_type\n"
        + "".join(f"1,2026010{d},1\n" for d in range(5, 10)),
    )
    feed_zip = tmp_path / "feed.zip"
    with zipfile.ZipFile(feed_zip, "w") as archive:
        for path in folder.iterdir():
            archive.write(path, f"feed/{path.name}")
    rc, sheets = _run_main(monkeypatch, feed_zip, tmp_path / "out")
    assert rc == 0
    ((_, (_, row)),) = sheets
    assert row[3:] == ("08:00", "08:10")


def test_main_returns_1_when_a_workbook_cannot_be_written(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n",
        trips="route_id,service_id,trip_id,trip_headsign,direction_id\nR,1,T1,East,0\n",
    )
    out = tmp_path / "out"
    # A directory where the workbook should go makes the write fail, like a locked file.
    (out / "weekday_sid_1" / "route_101_timetable_weekday_sid_1.xlsx").mkdir(parents=True)
    rc, _ = _run_main(monkeypatch, feed, out)
    assert rc == 1
    assert "FAILED EXPORTS (1)\n" in (out / _RUN_LOG).read_text(encoding="utf-8")
    assert "route '101', service_id '1'" in (out / _RUN_LOG).read_text(encoding="utf-8")


# T2 starts after T1 at A but reaches C first, so no row order is chronological at both.
_OVERTAKING = (
    "T1,08:00:00,08:00:00,A,1,1\nT1,08:30:00,08:30:00,C,2,1\n"
    "T2,08:05:00,08:05:00,A,1,1\nT2,08:20:00,08:20:00,C,2,1\n"
)


def test_main_flags_questionable_workbook_and_lists_it_in_run_log(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(tmp_path / "gtfs", _ABCD, _OVERTAKING)
    out = tmp_path / "out"
    rc, sheets = _run_main(monkeypatch, feed, out, _FLAGGED_WORKBOOK)
    assert rc == 0
    ((_, (_, *rows)),) = sheets
    # The first displayed time order is kept for review.
    assert [row[3:] for row in rows] == [("08:00", "08:30"), ("08:05", "08:20")]
    assert not (out / "weekday_sid_1" / _WORKBOOK).exists()

    log = (out / _RUN_LOG).read_text(encoding="utf-8")
    assert "FLAGGED WORKBOOKS (1)" in log
    assert str(Path("weekday_sid_1") / _FLAGGED_WORKBOOK) in log
    assert "no single trip row order" in log
    assert "Stop 'Charlie Schedule': trip 'T2' at 08:20" in log
    assert "FAILED EXPORTS (0)" in log
    # The CONFIGURATION block is recorded verbatim.
    assert 'WARNING_FILENAME_SUFFIX = "_warning"' in log


def test_main_clean_workbook_has_plain_name_and_empty_flag_list(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n",
        trips="route_id,service_id,trip_id,trip_headsign,direction_id\nR,1,T1,East,0\n",
    )
    out = tmp_path / "out"
    rc, sheets = _run_main(monkeypatch, feed, out)
    assert rc == 0
    assert sheets
    log = (out / _RUN_LOG).read_text(encoding="utf-8")
    assert "FLAGGED WORKBOOKS (0)" in log
    assert "Service labels:       1 → Weekday" in log


@pytest.mark.parametrize(
    ("stop_times", "stale_name", "kept_name"),
    [
        # Fixed since the last run: the old flagged copy is removed.
        ("T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n", _FLAGGED_WORKBOOK, _WORKBOOK),
        # Newly flagged: the old plain copy is removed.
        (_OVERTAKING, _WORKBOOK, _FLAGGED_WORKBOOK),
    ],
)
def test_main_removes_previous_copy_under_the_other_name(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    stop_times: str,
    stale_name: str,
    kept_name: str,
) -> None:
    feed = _write_feed(tmp_path / "gtfs", _ABCD, stop_times)
    out = tmp_path / "out"
    folder = out / "weekday_sid_1"
    folder.mkdir(parents=True)
    (folder / stale_name).write_text("previous run")
    rc, sheets = _run_main(monkeypatch, feed, out, kept_name)
    assert rc == 0
    assert sheets
    assert sorted(p.name for p in folder.iterdir()) == [kept_name]


def test_main_returns_1_when_required_run_log_cannot_be_written(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import scripts.gtfs_exports.gtfs_timetable_exporter as mod

    feed = _write_feed(
        tmp_path / "gtfs",
        _ABCD,
        "T1,08:00:00,08:00:00,A,1,1\nT1,08:10:00,08:10:00,B,2,1\n",
        trips="route_id,service_id,trip_id,trip_headsign,direction_id\nR,1,T1,East,0\n",
    )
    out = tmp_path / "out"
    # A directory where the run log should go makes the write fail.
    (out / _RUN_LOG).mkdir(parents=True)
    rc, _ = _run_main(monkeypatch, feed, out)
    assert rc == 1

    monkeypatch.setattr(mod, "REQUIRE_RUN_LOG", False)
    rc, _ = _run_main(monkeypatch, feed, out)
    assert rc == 0
